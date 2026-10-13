#!/usr/bin/env python3
"""Adversarial MC3 corpus + gold generator (v1: traps T1-T12, T13-T24 + canaries; opt-in v2: T25-T28).

    python3 eval/gen_corpus.py --seed 1 --out eval/suites/s1 [--n-files 40] [--traps T1,T2,...]
    python3 eval/gen_corpus.py --seed 201 --out eval/suites/p1a --traps T25,T26,T27,T28      # probe (v2 traps are opt-in)
    python3 eval/gen_corpus.py --seed 301 --out eval/suites/s5 --n-files 300 --n-images 60 \\
        --traps T1,T6,T7,T20,T23 --max-refusals 6                                           # scale suite

Writes <out>/corpus/ (the files), <out>/questions.json (EXACTLY the format of
sample-questions.json), <out>/traps.json ({query n: trap id}) and <out>/MANIFEST.txt.

Everything is drawn from a fictional universe (company, products, part numbers, tickets,
error codes, versions, quarters) that depends only on --seed, so two seeds never share
answers. Generation is deterministic: same seed, same questions and same file contents.

Adding a trap = writing `trap_T13(ctx)` (register files with ctx.text/pdf/docx/xlsx/...,
questions with ctx.q(...)) and appending "T13" to TRAPS. Canaries (questions that must be
refused) live in `canaries(ctx)`; each declares which traps it needs.

Gold convention (same as the sample): the value only; numbers without unit, identifiers
verbatim, quarters with FY, versions "7.2.1", pins "B14", revisions "REV-D1"; aliases like
the sample ("87 C", "87°C", "87 degrees"; "<Firmware> 7.2.1"; "pin B14").
"""

from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import os
import random
import re
import shutil
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

# Opt-in traps (v2): only run via --traps, never part of the default list, so existing seeds stay reproducible.
EXTRA_TRAPS = ["T25", "T26", "T27", "T28"]
# T28 re-uses the attribute names of the T1/T5/T8 datasheets for the product family (with its own values)
TRAP_CONFLICTS = {"T28": ("T1", "T2", "T5", "T8", "T14")}

TRAPS = ["T1", "T2", "T3", "T4", "T5", "T6", "T7", "T8", "T9", "T10", "T11", "T12",
         "T13", "T14b", "T15b", "T17b", "T18b", "T19b", "T20", "T21", "T22", "T23", "T24"]

ANSWER_RULE = ("Answer with the value only — a number, a part number, a version or a quarter. "
               "Cite the corpus-relative path of every file the answer came from. If the corpus "
               "does not contain the answer, return an empty answer and no citations; do not guess.")

FIXED_DT = dt.datetime(2026, 9, 1, 9, 0, 0)

# ------------------------------------------------------------------ word lists

FIRST = ["Halcyon", "Vantor", "Kestrel", "Solstice", "Argent", "Tessera", "Borealis", "Lumen",
         "Quasar", "Cobalt", "Ferrant", "Obsidian", "Zenith", "Ardent", "Pallas", "Veridian",
         "Marlow", "Corvid", "Helix", "Sundial", "Tamarind", "Noctua", "Brindle", "Caldera"]
SECOND = ["Dynamics", "Systems", "Microsystems", "Devices", "Semiconductor", "Computing",
          "Photonics", "Labs", "Silicon", "Instruments", "Logic", "Networks"]
PREFIXES = ["KX", "VR", "NB", "ZT", "RQ", "PL", "DX", "MV", "SK", "HN", "GW", "JT"]
NAMES = ["Aurora", "Zephyr", "Ember", "Tundra", "Basalt", "Cinder", "Mistral", "Sirocco", "Talon",
         "Quartz", "Nimbus", "Cirrus", "Stratus", "Monsoon", "Glacier", "Harrier", "Willow",
         "Juniper", "Lynx", "Onyx"]
SITES = ["Dublin", "Osaka", "Austin", "Lyon", "Tallinn", "Pune", "Porto", "Haifa", "Gdansk", "Reno"]
SUPPLIERS = ["Brightwell", "Norvik", "Castellan", "Ostrava", "Pemberton", "Rhodes"]

# (log text, noun phrase used in questions, ticket summary that shares NO words with the noun,
#  component)
SYMPTOMS = [
    ("link flap detected on fabric port 3, retraining", "fabric link flap",
     "Intermittent signal loss on interconnect lane under burst traffic", "interconnect"),
    ("HBM uncorrectable error on stack 2, page retired", "memory ECC",
     "Row-remap table exhausts early on repeated bit flips", "memory"),
    ("command queue stalled, watchdog expired after 30s", "queue stall",
     "Scheduler ring pointer fails to advance after doorbell race", "scheduler"),
    ("replay timeout on slot 4, link downtrained to x8", "PCIe replay timeout",
     "Lane equalisation is not re-run after a hot reset", "pcie"),
    ("rotor B reports 0 RPM, tachometer fault", "fan tachometer fault",
     "Tach input debounce window too short on second rotor", "thermal"),
    ("rail VDDQ below threshold for 4 ms", "voltage droop",
     "Power-good comparator trips on brief load step", "power"),
    ("reference clock lost lock, switched to backup", "PLL lock loss",
     "Spread-spectrum setting leaves oscillator near edge of capture range", "clocking"),
    ("controller 1 reset by supervisor", "watchdog reset",
     "Heartbeat task starved while firmware update is staged", "firmware"),
    ("die 0 at 96C, clocks reduced to 60%", "thermal throttle",
     "Die temperature sensor reports high under sustained small-batch load", "power"),
    ("DMA descriptor ring overrun on engine 5", "DMA ring overrun",
     "Producer index wraps early when completion batching is enabled", "dma"),
]
# (log text, noun phrase used in questions, ticket summary that shares NO words with the noun, component)
# T26/T27 events: none of them is a coolant-pump / coolant-loop / fan-tray event (T6 canary, T23) or one of SYMPTOMS.
EVENTS = [
    ("chiller compressor tripped on overcurrent", "chiller compressor trip",
     "Inrush limiter threshold is read from the wrong register", "hvac"),
    ("rack door sensor reports open during operation", "rack door alarm",
     "Debounce counter is shared between two contacts", "facilities"),
    ("firmware watchdog expired on port 2, link reset", "NIC watchdog expiry",
     "Keepalive timer restarts on every received frame", "network"),
    ("wear counter above 90 percent on slot 1", "flash wear alert",
     "Erase-cycle tally saturates at an 8-bit boundary", "storage"),
    ("manifold flow sensor reads zero", "flow sensor dropout",
     "Pulse counter overflows at high rotor speed", "thermal"),
    ("self-test failed on battery backup unit 2", "backup battery failure",
     "Load-test window ends before the voltage settles", "power"),
    ("bus transaction timed out, bus recovered", "SMBus timeout",
     "Clock-stretch limit is not honoured by the secondary device", "firmware"),
    ("lane group 2 skew above tolerance after retrain", "clock skew warning",
     "Phase detector gain is mis-scaled after link recovery", "clocking"),
    ("scrub pass aborted by supervisor", "memory scrub abort",
     "Patrol engine yields to a higher priority task indefinitely", "memory"),
    ("inlet filter differential pressure above limit", "filter clog warning",
     "Transducer offset drifts with ambient humidity", "facilities"),
]
GENERIC_BUGS = [
    ("firmware", "Release notes omit ticket list for {v}"),
    ("thermal", "Fan curve too aggressive at idle"),
    ("telemetry", "Per-die occupancy counter drifts after 18h uptime"),
    ("installer", "Update tool fails on first attempt after cold boot"),
    ("docs", "Pin table in the integration guide lists a swapped pair"),
    ("cli", "Status command prints stale firmware string"),
    ("scheduler", "Fairness weights ignored for jobs below 4 requests"),
    ("power", "Idle power reading rounds down to nearest 5 W"),
    ("logging", "Timestamps lose milliseconds after rotation"),
    ("telemetry", "Exporter drops samples when the buffer is exactly full"),
    ("storage", "Checkpoint upload retries without backoff"),
    ("security", "Debug port stays enabled after provisioning"),
]
RELEASE_LINES = [
    "Scheduler fairness rework. New telemetry counters for per-die occupancy.",
    "Fixes a rare hang in the PCIe doorbell path under sustained small-batch load.",
    "Packaging only; no functional change.",
    "Improved boot-time memory training; adds the status subcommand to the CLI.",
    "Updated thermal curves for the passive heatsink SKU.",
    "Logging: millisecond timestamps in rotated files.",
    "Security: debug port is locked after provisioning.",
    "Telemetry exporter batching tuned for large fleets.",
    "Installer now verifies the image signature before staging.",
    "Documentation fixes only.",
]
SENTENCES = [
    "The team reviewed the supplier onboarding checklist and agreed to close the remaining items.",
    "Action: confirm the shipping window with logistics before the next planning call.",
    "Customer feedback on the integration guide was mostly positive, with two requests for diagrams.",
    "Packaging samples for the backplane cable arrived and were sent to quality for inspection.",
    "The rack validation plan was updated; owners will report progress every second Tuesday.",
    "Support volume was flat week over week, with no new severity-one escalations.",
    "We discussed moving the weekly sync to a shorter format and rotating the facilitator.",
    "Finance asked for a refreshed forecast; the draft is due after the quarter-end review.",
    "Training material for field engineers is being translated; review copies go out soon.",
    "The lab will be closed for a facilities inspection; testing resumes the following day.",
    "A new intern joins the validation group and will shadow the thermal characterisation work.",
    "Documentation for the status subcommand needs one more pass before publication.",
    "The shared test rack was reserved for the upcoming interoperability event.",
    "Procurement is comparing two alternative connector vendors on lead time and tolerance.",
]


# ------------------------------------------------------------------ small helpers

def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower())


def comma(n: int) -> str:
    return f"{n:,}"


def al_temp(v):
    return [f"{v} C", f"{v}°C", f"{v} degrees"]


def al_units(v, *units):
    out = []
    for u in units:
        out.append(f"{v} {u}")
    return out


def money(v: float):
    s = f"{v:.2f}"
    al = [f"${s}", f"{v:g}", f"USD {s}"]
    return s, [a for a in dict.fromkeys(al) if a != s]


def bound_re(token: str) -> re.Pattern:
    """Token with identifier boundaries: not inside a longer word / dotted / dashed run."""
    return re.compile(r"(?<![A-Za-z0-9_.\-])" + re.escape(token) + r"(?![A-Za-z0-9_\-]|\.\d)", re.I)


def is_ident(ans: str) -> bool:
    """Answers specific enough to test for stray occurrences in other files."""
    return len(ans) >= 4 and bool(re.search(r"[A-Za-z]", ans) and re.search(r"\d", ans)) \
        or bool(re.fullmatch(r"\d+\.\d+\.\d+", ans))


# ------------------------------------------------------------------ universe

def make_universe(rng: random.Random, seed: int) -> SimpleNamespace:
    u = SimpleNamespace()
    first = rng.choice(FIRST)
    u.company = f"{first} {rng.choice(SECOND)}"
    u.cslug = slug(first)
    u.code3 = first[:3].upper()
    cons = [c for c in first.upper() if c not in "AEIOU"]
    u.tix = ("".join(cons[:3]) + "XQZ")[:3]
    if u.tix == u.code3:
        u.tix = cons[0] + cons[-1] + "X"
    pool = NAMES[:]
    rng.shuffle(pool)
    u.fw = pool.pop()
    u.fwslug = slug(u.fw)
    u.codename = {k: pool.pop() for k in ("p0", "p2", "p3")}
    u.pre = rng.choice(PREFIXES)
    sc = rng.choice([1, 10])
    n0 = rng.randrange(22, 57)
    off = {"p0": 0, "p1": -20, "p2": 20, "p3": 10, "p4": 40}
    u.p = {k: f"{u.pre}-{(n0 + d) * sc}" for k, d in off.items()}
    u.nx_product = f"{u.pre}-{99 * sc}"
    u.fy = rng.choice([27, 28, 29])
    u.year = 2026
    used: set = set()

    def uniq(fn):
        while True:
            x = fn()
            if x not in used:
                used.add(x)
                return x

    def part(cat):
        return uniq(lambda: f"{u.code3}-{cat}-{rng.randint(1000, 9999)}-{rng.choice('ABCDEFG')}")

    def ticket():
        return uniq(lambda: f"{u.tix}-{rng.randint(1500, 2999)}")

    def ecode():
        return uniq(lambda: f"E{rng.randint(1000, 9999)}")

    def dev():
        return uniq(lambda: f"DEV-{rng.randint(1000, 9999)}")

    def cost():
        return rng.randrange(20 * 4, 400 * 4) / 4

    # --- datasheet values (T1/T8/T14), p0
    u.jt_old = rng.randint(100, 115)
    u.jt_new = u.jt_old - rng.randint(6, 14)
    u.jt_p1 = u.jt_new + rng.choice([-9, -7, -5, 5, 7, 9])
    u.jt_typ = u.jt_new - rng.randint(18, 28)
    u.pw_new = rng.randrange(300, 800, 5)
    u.pw_old = u.pw_new - rng.randrange(25, 80, 5)
    u.cap = rng.randrange(32, 257, 8)
    u.cap_old = rng.choice([c for c in range(32, 257, 8) if c != u.cap])
    u.bw = f"{rng.uniform(2.0, 8.0):.1f}"
    u.bw_old = f"{rng.uniform(1.5, 7.5):.1f}"
    if u.bw_old == u.bw:
        u.bw_old = f"{float(u.bw) + 0.4:.1f}"
    u.moq = rng.randrange(500, 25000, 250)
    u.storage_max = rng.choice([85, 100, 105, 125])
    u.altitude = rng.choice([2000, 3000, 3500, 4000])
    u.aux = rng.choice(["12V-2x6", "16-pin 12VHPWR", "8-pin EPS"])
    u.noise_p1 = rng.randint(34, 52)
    # --- p3 (errata chain)
    u.idle_x = rng.randint(40, 95)
    u.idle_y = u.idle_x + rng.randint(3, 9)
    u.idle_w = u.idle_x - rng.randint(5, 15)
    u.bw3 = f"{rng.uniform(2.0, 8.0):.1f}"
    u.bw3_old = f"{float(u.bw3) - rng.choice([0.4, 0.6, 0.8]):.1f}"
    # --- parts
    u.fan_p0, u.fan_p1, u.fan_arch = part("FAN"), part("FAN"), part("FAN")
    u.hs_new, u.hs_old, u.hs_p1 = part("HSK"), part("HSK"), part("HSK")
    u.psu_new, u.psu_old, u.psu_arch = part("PSU"), part("PSU"), part("PSU")
    u.cab_p0, u.cab_p1, u.pump_p1 = part("CAB"), part("CAB"), part("PMP")
    u.cost_cab_p0 = cost()
    u.cost_cab_p1 = cost()
    while u.cost_cab_p1 == u.cost_cab_p0:
        u.cost_cab_p1 = cost()
    u.lead = {k: rng.choice([7, 10, 14, 21, 28, 35, 42]) for k in ("fan", "hs", "psu", "cab", "pump")}
    # --- versions, tickets, devices, error codes
    major, minor = rng.randint(3, 9), rng.randint(0, 5)
    pa, pb = rng.randint(2, 3), rng.randint(5, 7)
    if rng.random() < 0.5:
        pa, pb = pb, pa
    u.ver_base = f"{major}.{minor}"
    u.v6, u.v7 = f"{major}.{minor}.{pa}", f"{major}.{minor}.{pb}"
    u.ver_run = f"{major}.{minor}.{pa - 1}"          # what the node was running when it failed
    u.t6, u.t7, u.t_dec = ticket(), ticket(), ticket()
    u.nx_tickets = [uniq(lambda: f"{u.tix}-{rng.randint(3000, 3999)}") for _ in range(2)]
    u.dev7, u.dev_dec = dev(), dev()
    u.nx_dev = dev()
    u.e6, u.e7, u.e_dec = ecode(), ecode(), ecode()
    u.nx_err = ecode()
    syms = SYMPTOMS[:]
    rng.shuffle(syms)
    u.s6, u.s7, u.s_dec = syms[0], syms[1], syms[2]
    u.syms = syms                                   # v1 traps use syms[3:] (never T6/T7's own)
    u.d6 = dt.date(2026, rng.randint(8, 10), rng.randint(1, 28))
    u.d7 = u.d6 + dt.timedelta(days=rng.randint(3, 12))
    # --- python constants
    u.svc = rng.choice([("ingest service", "ingest_service", "batch", "BATCH"),
                        ("export gateway", "export_gateway", "upload", "UPLOAD"),
                        ("sync daemon", "sync_daemon", "lease", "LEASE")])
    u.to_new = rng.randrange(100, 1000, 5)
    u.to_old = rng.choice([x for x in (30, 45, 60, 90) if x < u.to_new])
    u.hb = rng.choice([15, 20, 45, 50, 75, 90])
    u.hb_doc = rng.choice([x for x in (10, 30, 60) if x != u.hb])
    # --- runbook / cooling
    u.floor0 = rng.randrange(25, 61, 5)
    u.floor1 = rng.choice([x for x in range(25, 61, 5) if x != u.floor0])
    while True:
        u.c0, u.c1 = rng.randint(28, 40), rng.randint(28, 40)
        u.c_lab, u.c_comm = rng.randint(20, 26), rng.randint(24, 27)
        if len({u.c0, u.c1, u.c_lab, u.c_comm}) == 4:
            break
    # --- quarters (all distinct)
    qs = [f"Q{q} FY{y}" for y in (u.fy, u.fy + 1) for q in range(1, 5)]
    rng.shuffle(qs)
    u.q_sample_p2, u.q_vol_p2, u.q_sample_p4, u.q_vol_p4 = qs[:4]
    u.q_tapeout, u.q_pdk, u.q_bringup, u.q_es = qs[4:8]
    # --- vendor secrets (encrypted / unreadable)
    u.supplier = rng.choice(SUPPLIERS)
    u.vol = rng.choice([5000, 10000, 25000, 50000])
    u.price = rng.randint(2500, 9000)
    u.pay_days = rng.choice([30, 45, 60, 90])
    u.liab = rng.choice([500000, 1000000, 2500000])
    u.penalty = rng.choice([1.5, 2.0, 2.5, 3.0, 5.0])
    u.nfind = rng.randint(3, 12)
    # --- image-only facts
    u.rev = f"REV-{rng.choice('ABCDEF')}{rng.randint(1, 9)}"
    up = "ABCDEFGHJKLMNPQRSTUVWXYZ"
    u.serial = f"{u.code3[0]}{rng.choice('QRSTWXZ')}{rng.randint(1, 9)}-{rng.randint(100000, 999999)}-{rng.randint(10, 99)}"
    u.fcc = f"{rng.randint(2, 9)}{''.join(rng.choice(up) for _ in range(4))}-{rng.choice(up)}{rng.randint(100, 999)}"
    u.made_in = rng.choice(["MALAYSIA", "TAIWAN", "VIETNAM", "MEXICO", "POLAND"])
    letter, n = rng.choice("ABCDJ"), rng.randint(1, 40)
    u.pin_row = [f"{letter}{n + i}" for i in range(6)]
    u.pin_sig = rng.choice(["THERM_ALERT#", "PWR_GOOD", "FAULT#", "WAKE#", "PERST#", "SMB_ALERT#"])
    others = [s for s in ["GND", "SDA", "SCL", "PRSNT#", "CLK_P", "CLK_N", "3V3_AUX", "RESET#", "TX_EN",
                          "RX_LOS"] if s != u.pin_sig]
    rng.shuffle(others)
    u.pin_idx = rng.randrange(6)
    sigs = others[:5]
    sigs.insert(u.pin_idx, u.pin_sig)
    u.pin_sigs = sigs
    u.pin = u.pin_row[u.pin_idx]
    u.sheet_no = rng.randint(2, 7)
    return u


# ------------------------------------------------------------------ documents

class Doc:
    def __init__(self, path: str, kind: str):
        self.path, self.kind = path, kind
        self.lines: list[str] = []      # text-like files
        self.blocks: list[tuple] = []   # pdf: (style, text); docx: ("h"|"p"|"table"|"img", ...)
        self.sheets: list[dict] = []    # xlsx
        self.image = None               # PIL image (png/jpg)
        self.raw: bytes | None = None   # bin
        self.password: str | None = None
        self.hidden_text: str = ""      # content that is deliberately unreadable (encrypted / 000)
        self.chmod: int | None = None
        self.expect = "ok"              # expected parse status
        self.note = ""
        self.encoding = "utf-8"         # text files: "utf-16" writes a BOM
        self.newline = "\n"             # text files: "\r\n" for Windows line endings

    def text(self) -> str:
        if self.kind == "text":
            return "\n".join(self.lines)
        if self.kind == "pdf":
            return "\n".join(t for _, t in self.blocks) if not self.password else ""
        if self.kind == "docx":
            out = []
            for b in self.blocks:
                if b[0] in ("h", "p"):
                    out.append(b[1])
                elif b[0] == "table":
                    out += [" | ".join(r) for r in b[1]]
            return "\n".join(out)
        if self.kind == "xlsx":
            out = []
            for sh in self.sheets:
                if sh.get("title"):
                    out.append(sh["title"])
                out += [" | ".join(str(c) for c in r) for r in sh.get("pre", [])]
                out.append(" | ".join(sh["header"]))
                # "shown": what the sheet displays (formula results, number formats) when it differs from `rows`
                out += [" | ".join(str(c) for c in r) for r in sh.get("shown", sh["rows"])]
            return "\n".join(out)
        return ""


def fixed_zip(path: Path):
    """Rewrite an OOXML zip with fixed timestamps so the bytes depend only on content."""
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in sorted(zin.infolist(), key=lambda i: i.filename):
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 9, 1, 9, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":    # openpyxl stamps `modified` with now() on save
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-09-01T09:00:00Z", data)
            zout.writestr(zi, data)
    tmp.replace(path)


def _font(size: int):
    from PIL import ImageFont
    return ImageFont.load_default(size=size)


def write_pdf(doc: Doc, dest: Path):
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos
    pdf = FPDF()
    pdf.set_creation_date(FIXED_DT)
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    pdf.set_margins(18, 18, 18)
    pdf.set_y(18)
    for style, text in doc.blocks:
        text = text.encode("latin-1", "replace").decode("latin-1")
        if style == "page":
            pdf.add_page()
            continue
        if style == "title":
            pdf.set_font("Helvetica", "B", 14)
            h = 8
        elif style == "h":
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 12)
            h = 7
        elif style == "mono":       # fixed-pitch line: padded columns stay aligned (spec tables)
            pdf.set_font("Courier", "", 9)
            h = 5.5
        else:
            pdf.set_font("Helvetica", "", 11)
            h = 6.5
        pdf.multi_cell(0, h, text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    data = bytes(pdf.output())
    if doc.password:
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
        w.encrypt(doc.password)
        buf = io.BytesIO()
        w.write(buf)
        data = buf.getvalue()
    dest.write_bytes(data)


def write_docx(doc: Doc, dest: Path):
    import docx
    from docx.shared import Inches
    d = docx.Document()
    d.core_properties.author = "docs"
    d.core_properties.created = d.core_properties.modified = FIXED_DT
    d.core_properties.last_modified_by = "docs"
    for b in doc.blocks:
        if b[0] == "h":
            d.add_heading(b[1], level=b[2] if len(b) > 2 else 1)
        elif b[0] == "p":
            d.add_paragraph(b[1])
        elif b[0] == "table":
            rows = b[1]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c
        elif b[0] == "img":
            buf = io.BytesIO()
            b[1].save(buf, "PNG")
            buf.seek(0)
            d.add_picture(buf, width=Inches(5.5))
    d.save(dest)
    fixed_zip(dest)


def write_xlsx(doc: Doc, dest: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "docs"
    wb.properties.created = wb.properties.modified = FIXED_DT
    for sh in doc.sheets:
        ws = wb.create_sheet(sh["name"])
        pre = ([[sh["title"]]] if sh.get("title") else []) + [list(x) for x in sh.get("pre", [])]
        for row in pre:
            ws.append(row)
        if pre:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(sh["header"]))
            ws["A1"].font = Font(bold=True, size=13)
        r0 = len(pre) + 1
        ws.append(sh["header"])
        for c in ws[r0]:
            c.font = Font(bold=True)
        for row in sh["rows"]:
            ws.append(list(row))          # a str starting with "=" is written as a formula (no cached value)
        for ci, col in enumerate(sh["header"], start=1):
            nf = "0.00" if col in sh.get("money", ()) else sh.get("fmt", {}).get(col)
            if nf:
                for r in range(r0 + 1, r0 + 1 + len(sh["rows"])):
                    ws.cell(r, ci).number_format = nf
            ws.column_dimensions[ws.cell(r0, ci).column_letter].width = 26
        if sh.get("hidden"):
            ws.sheet_state = "hidden"
    wb.save(dest)
    fixed_zip(dest)


def write_doc(doc: Doc, dest: Path):
    dest.parent.mkdir(parents=True, exist_ok=True)
    k = doc.kind
    if k == "text":
        dest.write_bytes((doc.newline.join(doc.lines) + doc.newline).encode(doc.encoding))
    elif k == "pdf":
        write_pdf(doc, dest)
    elif k == "docx":
        write_docx(doc, dest)
    elif k == "xlsx":
        write_xlsx(doc, dest)
    elif k == "png":
        doc.image.save(dest, "PNG")
    elif k == "jpg":
        doc.image.convert("RGB").save(dest, "JPEG", quality=90)
    elif k == "bin":
        dest.write_bytes(doc.raw)
    else:
        raise ValueError(k)


# ------------------------------------------------------------------ context

class Ctx:
    def __init__(self, seed: int, attempt: int, traps: list[str], n_files: int, n_images: int = 0):
        self.seed, self.attempt, self.traps, self.n_files = seed, attempt, traps, n_files
        self.n_images = n_images
        self.rng = random.Random(f"mc3:{seed}:{attempt}")
        self.rng_fill = random.Random(f"mc3fill:{seed}:{attempt}")
        self.u = make_universe(self.rng, seed)
        self.docs: dict[str, Doc] = {}
        self.queries: list[dict] = []
        self.empty_dirs: list[str] = []
        self.cache: dict = {}            # values shared between traps / canaries (per trap, drawn from xr())
        self.forbid: list[str] = []      # tokens fillers must never contain (image-only facts, ...)
        self.absent: list[str] = []      # well-formed ids that must appear in NO file (false-premise questions)
        self.secrets: dict[str, set] = {}  # token -> the only files whose text may contain it

    def xr(self, name: str) -> random.Random:
        """Per-trap RNG: v1 traps draw from here so their values never depend on which other traps ran."""
        return random.Random(f"mc3x:{self.seed}:{self.attempt}:{name}")

    def doc(self, path: str, kind: str) -> Doc:
        if path not in self.docs:
            self.docs[path] = Doc(path, kind)
        return self.docs[path]

    def text(self, path):
        return self.doc(path, "text")

    def q(self, trap, query, answer, cites, category, note="", aliases=None, also=(), unanswerable=False,
          imgonly=False):
        if unanswerable:
            answer, cites = "", []
        self.queries.append({"trap": trap, "n": 0, "query": query, "expected_answer": answer,
                             "answer_aliases": list(aliases or []), "expected_citations": list(cites),
                             "category": category, "note": note, "unanswerable": unanswerable,
                             "also": set(also), "imgonly": imgonly})


# ------------------------------------------------------------------ shared documents

def ensure_p0_datasheets(ctx: Ctx):
    """p0 r1 (WITHDRAWN) and r2 datasheets; every trap that needs them calls this (idempotent)."""
    u = ctx.u
    s = slug(u.p["p0"])
    r2p = f"specs/{s}_datasheet_r2.pdf"
    if r2p in ctx.docs:
        return r2p, f"specs/{s}_datasheet_r1_WITHDRAWN.pdf"
    r1p = f"specs/{s}_datasheet_r1_WITHDRAWN.pdf"
    hdr = f'{u.company} {u.p["p0"]} "{u.codename["p0"]}" Accelerator'
    d1 = ctx.doc(r1p, "pdf")
    d1.blocks = [("title", hdr),
                 ("p", "Datasheet, revision 1 - WITHDRAWN, SUPERSEDED BY REVISION 2"),
                 ("p", "Preliminary silicon. Do not design to these numbers."),
                 ("h", "Electrical and thermal"),
                 ("l", f"Maximum junction temperature .......... {u.jt_old} C"),
                 ("l", f"Board power (TBP) ..................... {u.pw_old} W"),
                 ("l", f"Maximum storage temperature ........... {u.storage_max} C"),
                 ("l", f"Maximum operating altitude ........... {comma(u.altitude)} m"),
                 ("h", "Memory"),
                 ("l", f"Capacity .............................. {u.cap_old} GiB HBM3e"),
                 ("l", f"Peak bandwidth ........................ {u.bw_old} TB/s")]
    d2 = ctx.doc(r2p, "pdf")
    d2.blocks = [("title", hdr),
                 ("p", "Datasheet, revision 2 - supersedes revision 1"),
                 ("h", "Electrical and thermal"),
                 ("l", f"Maximum junction temperature .......... {u.jt_new} C"),
                 ("l", f"Board power (TBP) ..................... {u.pw_new} W"),
                 ("l", f"Auxiliary power connector ............. {u.aux}"),
                 ("h", "Memory"),
                 ("l", f"Capacity .............................. {u.cap} GiB HBM3e"),
                 ("l", f"Peak bandwidth ........................ {u.bw} TB/s"),
                 ("h", "Ordering"),
                 ("l", f"Minimum order quantity ................. {comma(u.moq)} units"),
                 ("h", "Revision history"),
                 ("l", f"r2  Junction temperature corrected to {u.jt_new} C after"),
                 ("l", "    characterisation of production silicon. Revision 1 is withdrawn."),
                 ("l", "r1  Initial release, preliminary silicon.")]
    return r2p, r1p


def ensure_bug_db(ctx: Ctx):
    path = "support/bug_database.csv"
    if path in ctx.docs:
        return path
    u, rng = ctx.u, ctx.rng_fill
    d = ctx.text(path)
    rows = [(u.t6, u.s6[3], "S2", u.s6[2], "closed", u.v6),
            (u.t7, u.s7[3], "S2", u.s7[2], "closed", u.v7),
            (u.t_dec, u.s_dec[3], "S3", u.s_dec[2], "closed", f"{u.ver_base.split('.')[0]}.{u.ver_base.split('.')[1]}.9")]
    used = {u.t6, u.t7, u.t_dec, *u.nx_tickets}
    bugs = GENERIC_BUGS[:]
    rng.shuffle(bugs)
    u.fleet_tickets = []
    for comp, summ in bugs[:9]:
        while True:
            t = f"{u.tix}-{rng.randint(1500, 2999)}"
            if t not in used:
                used.add(t)
                break
        closed = rng.random() < 0.6
        major, minor = u.ver_base.split(".")
        fixed = f"{int(major) - 1}.{rng.randint(0, 5)}.{rng.randint(0, 4)}" if closed else ""
        if "{v}" in summ:
            summ = summ.format(v=u.v6)
        rows.append((t, comp, rng.choice(["S3", "S4"]), summ, "closed" if closed else "open", fixed))
        u.fleet_tickets.append(t)
    rng.shuffle(rows)
    d.lines = ["ticket,component,severity,summary,status,fixed_in"] + [",".join(r) for r in rows]
    return path


def ensure_release_notes(ctx: Ctx):
    path = f"engineering/{ctx.u.fwslug}_release_notes.txt"
    if path in ctx.docs:
        return path
    u, rng = ctx.u, ctx.rng_fill
    d = ctx.text(path)
    skip = {u.v6, u.v7}
    ents = []
    for p in range(0, 10):
        v = f"{u.ver_base}.{p}"
        if v in skip:
            continue
        ents.append(f"{v}  {rng.choice(RELEASE_LINES)}")
    d.lines = [f"{u.fw} firmware - release notes", ""] + ents + [
        "", "Entries are published on release. Fixes are tracked per ticket in the",
        "bug database, which is authoritative for which release closed a ticket."]
    return path


def ensure_parts(ctx: Ctx):
    path = f"support/spare_parts_{slug(ctx.u.pre)}_series.xlsx"
    if path in ctx.docs:
        return path
    u, rng = ctx.u, ctx.rng_fill
    p0, p1 = u.p["p0"], u.p["p1"]
    cst = lambda: rng.randrange(20 * 4, 400 * 4) / 4
    hdr = ["part_number", "description", "compatible_with", "unit_cost_usd", "lead_time_days", "status"]
    rows = [
        (u.fan_p0, "Fan assembly, dual-rotor, field replaceable", p0, cst(), u.lead["fan"], "active"),
        (u.fan_p1, "Fan assembly, single-rotor (legacy)", p1, cst(), 14, "active"),
        (u.psu_old, "Power supply module, 12V-2x6 input", p0, cst(), u.lead["psu"], "superseded"),
        (u.psu_new, "Power supply module, 12V-2x6 input", p0, cst(), u.lead["psu"], "active"),
        (u.hs_old, "Heatsink, vapour chamber (legacy)", p0, cst(), u.lead["hs"], "last-time-buy"),
        (u.hs_new, "Heatsink, vapour chamber", p0, cst(), u.lead["hs"], "active"),
        (u.hs_p1, "Heatsink, vapour chamber", p1, cst(), 28, "active"),
        (u.cab_p0, "Backplane cable, 400mm", p0, u.cost_cab_p0, u.lead["cab"], "active"),
        (u.cab_p1, "Backplane cable, 400mm", p1, u.cost_cab_p1, 7, "active"),
        (u.pump_p1, "Coolant pump assembly", p1, cst(), u.lead["pump"], "active"),
    ]
    rng.shuffle(rows)
    d = ctx.doc(path, "xlsx")
    d.sheets = [
        {"name": "RMA Parts", "title": f"{u.company} - RMA spare parts, {u.pre} series (FY{u.fy})",
         "header": hdr, "rows": rows, "money": ["unit_cost_usd"]},
        {"name": "Suppliers", "header": ["supplier", "category", "contact", "last_audit_score"],
         "rows": [(s, rng.choice(["fans", "thermal", "power", "cabling"]),
                   f"{slug(s)}@example.com", rng.randint(61, 98)) for s in SUPPLIERS[:4]]},
        {"name": "Notes", "header": ["note"],
         "rows": [("Lead times are supplier-quoted and exclude customs.",)]},
        {"name": "Archive", "header": hdr, "hidden": True,
         "rows": [(u.psu_arch, "Power supply module, 8-pin input", p0, cst(), 21, "retired"),
                  (u.fan_arch, "Fan assembly, dual-rotor, first revision", p0, cst(), 21, "retired")]},
    ]
    return path


# ------------------------------------------------------------------ traps

def trap_T1(ctx: Ctx):
    """Withdrawn revision next to the current one."""
    u = ctx.u
    r2, r1 = ensure_p0_datasheets(ctx)
    p0 = u.p["p0"]
    ctx.q("T1", f"What is the maximum junction temperature of the {p0}?", str(u.jt_new), [r2],
          "pdf, superseded-document trap", aliases=al_temp(u.jt_new),
          note=f"The WITHDRAWN r1 datasheet says {u.jt_old} C. Citing r1 is wrong even though it is a real file with a real number in it.")
    ctx.q("T1", f"What is the memory capacity of the {p0}, in GiB?", str(u.cap), [r2],
          "pdf, superseded-document trap", aliases=al_units(u.cap, "GiB"),
          note=f"The WITHDRAWN r1 datasheet says {u.cap_old} GiB.")
    ctx.q("T14", f"What is the minimum order quantity of the {p0}?", comma(u.moq), [r2],
          "pdf, number format with thousands separator",
          aliases=sorted({str(u.moq), f"{comma(u.moq)} units"} - {comma(u.moq)}),
          note="Only r2 lists the minimum order quantity; keep the comma as printed when the value is >= 1,000.")


def trap_T2(ctx: Ctx):
    """r1 withdrawn -> r2 value -> errata that overrides the r2 figure."""
    u = ctx.u
    p3 = u.p["p3"]
    s = slug(p3)
    r1, r2, er = (f"specs/{s}_datasheet_r1_WITHDRAWN.pdf", f"specs/{s}_datasheet_r2.pdf",
                  f"specs/errata_{s}.txt")
    hdr = f'{u.company} {p3} "{u.codename["p3"]}" Accelerator'
    ctx.doc(r1, "pdf").blocks = [
        ("title", hdr), ("p", "Datasheet, revision 1 - WITHDRAWN, SUPERSEDED BY REVISION 2"),
        ("p", "Preliminary silicon. Do not design to these numbers."),
        ("h", "Table 3 - Power"),
        ("l", f"Idle board power ........................ {u.idle_w} W"),
        ("h", "Memory"),
        ("l", f"Peak bandwidth ........................ {u.bw3_old} TB/s")]
    ctx.doc(r2, "pdf").blocks = [
        ("title", hdr), ("p", "Datasheet, revision 2 - supersedes revision 1"),
        ("h", "Table 3 - Power"),
        ("l", f"Idle board power ........................ {u.idle_x} W"),
        ("h", "Memory"),
        ("l", f"Peak bandwidth ........................ {u.bw3} TB/s"),
        ("h", "Revision history"),
        ("l", "r2  Characterised on production silicon. Revision 1 is withdrawn.")]
    ctx.text(er).lines = [
        f"Errata sheet - {u.company} {p3} datasheet revision 2", "Issued after the r2 release.", "",
        f'E1. Table 3, "Idle board power": value corrected to {u.idle_y} W (supersedes datasheet r2 figure).',
        "E2. No other changes to revision 2."]
    ctx.q("T2", f"What is the idle board power of the {p3}?", str(u.idle_y), [er],
          "three-level chain, errata overrides datasheet", aliases=al_units(u.idle_y, "W", "watts"),
          note=f"r1 (withdrawn) says {u.idle_w} W, r2 says {u.idle_x} W, the errata corrects Table 3 to {u.idle_y} W. "
               "By citation-by-necessity the errata alone holds the value and names the field, so only the errata is cited "
               "(the datasheet is only the figure that was superseded).",
          also=[r2])
    ctx.q("T2", f"What is the peak memory bandwidth of the {p3}?", u.bw3, [r2],
          "pdf, superseded-document trap next to an errata sheet", aliases=al_units(u.bw3, "TB/s"),
          note=f"The errata does not touch bandwidth, so r2 is the answer; the WITHDRAWN r1 says {u.bw3_old}.")


def trap_T3(ctx: Ctx):
    """Stale record inside an xlsx: status=superseded row and a (legacy) description."""
    u = ctx.u
    path = ensure_parts(ctx)
    p0 = u.p["p0"]
    ctx.q("T3", f"What is the part number of the power supply module for the {p0}?", u.psu_new, [path],
          "xlsx, stale-row trap (status column)",
          note=f"The same sheet also lists {u.psu_old} with status=superseded, and the hidden Archive sheet lists {u.psu_arch}.")
    ctx.q("T3", f"What is the part number of the vapour chamber heatsink for the {p0}?", u.hs_new, [path],
          "xlsx, stale-row trap ((legacy) suffix)",
          note=f"The row marked '(legacy)' holds {u.hs_old}; the {u.p['p1']} heatsink {u.hs_p1} is the near-miss.")


def trap_T4(ctx: Ctx):
    """Comment / docstring names an old value; the literal is the answer."""
    u = ctx.u
    svc_name, svc_mod, kind, K = u.svc
    path = f"engineering/{svc_mod}.py"
    ctx.text(path).lines = [
        f'"""{svc_name.capitalize()} for {u.company} telemetry batches."""', "",
        f"# Seconds a {kind} may sit in the staging queue before it is abandoned.",
        f"# Raised from {u.to_old} after the 2026-08 backlog incident.",
        f"DEFAULT_{K}_TIMEOUT_S = {u.to_new}", "",
        "MAX_INFLIGHT_BATCHES = 12", f'STAGING_ROOT = "/var/lib/{u.cslug}/staging"', "", "",
        f"def accept(batch, timeout_s=DEFAULT_{K}_TIMEOUT_S):",
        '    """Accept a batch for ingestion, or raise if the queue is saturated."""',
        "    if batch.size_bytes > 2 << 30:", '        raise ValueError("batch exceeds 2 GiB")',
        "    return batch.enqueue(timeout_s=timeout_s)"]
    ctx.q("T4", f"What is the default {kind} timeout, in seconds, in the {svc_name}?", str(u.to_new), [path],
          "python source, comment-vs-literal trap",
          aliases=[f"{u.to_new} seconds", f"{u.to_new}s"],
          note=f"A comment names the old value of {u.to_old}. The constant is the answer.")
    path2 = "engineering/fleet_agent.py"
    ctx.text(path2).lines = [
        '"""Fleet agent: reports node health to the fleet controller.', "",
        f"Sends a heartbeat every {u.hb_doc} seconds.", '"""', "",
        "# Heartbeat period; see the controller documentation.", f"HEARTBEAT_INTERVAL_S = {u.hb}", "",
        "RETRY_LIMIT = 5", "", "",
        "def run(controller, interval_s=HEARTBEAT_INTERVAL_S):",
        '    """Send heartbeats until the controller closes the session."""',
        "    while controller.open():", "        controller.beat()", "        controller.sleep(interval_s)"]
    ctx.q("T4", "What is the default heartbeat interval, in seconds, in the fleet agent?", str(u.hb), [path2],
          "python source, docstring-vs-literal trap", aliases=[f"{u.hb} seconds", f"{u.hb}s"],
          note=f"The module docstring still says {u.hb_doc} seconds. The constant (and the function default) is the answer.")


def trap_T5(ctx: Ctx):
    """Near-miss row for the sibling product."""
    u = ctx.u
    path = ensure_parts(ctx)
    p0, p1 = u.p["p0"], u.p["p1"]
    ctx.q("T5", f"What is the part number of the field-replaceable fan assembly for the {p0}?", u.fan_p0, [path],
          "xlsx, near-miss row", note=f"A legacy single-rotor fan for the {p1} ({u.fan_p1}) is the near-miss row; the hidden Archive sheet holds {u.fan_arch}.")
    s, al = money(u.cost_cab_p0)
    ctx.q("T5", f"What is the unit cost in USD of the backplane cable for the {p0}?", s, [path],
          "xlsx, near-miss row, price format", aliases=al,
          note=f"The {p1} cable ({u.cab_p1}) costs {u.cost_cab_p1:.2f}; only the {p0} row counts.")
    # p1's own datasheet: a sibling-only attribute (acoustic noise) lives here
    s1 = slug(p1)
    ctx.doc(f"specs/{s1}_datasheet_r1.pdf", "pdf").blocks = [
        ("title", f"{u.company} {p1} Accelerator"), ("p", "Datasheet, revision 1"),
        ("h", "Electrical and thermal"),
        ("l", f"Maximum junction temperature .......... {u.jt_p1} C"),
        ("l", f"Typical acoustic noise at 50% fan ...... {u.noise_p1} dB(A)")]


def trap_T6(ctx: Ctx):
    """Log -> ticket -> bug database (2 links), plus the direct and single-hop forms."""
    u = ctx.u
    db, rn = ensure_bug_db(ctx), ensure_release_notes(ctx)
    path = f"logs/prod_inference_{u.d6.isoformat()}.log"
    rng = ctx.rng_fill
    node = f"inference-node-{rng.randint(1, 12):02d}"
    proc = f"{u.fwslug}[{rng.randint(1000, 4999)}]"
    t = dt.datetime.combine(u.d6, dt.time(3, rng.randint(0, 30), rng.randint(0, 59)))
    lines = []

    def log(level, msg, step=None):
        nonlocal t
        t += dt.timedelta(seconds=step if step is not None else rng.randint(2, 40))
        lines.append(f"{t.strftime('%Y-%m-%dT%H:%M:%SZ')} {node} {proc}: {level:<5} {msg}")

    log("INFO", f"firmware {u.ver_run} started, {rng.choice([4, 8])} dies online", 0)
    log("INFO", f"scheduler: batch {rng.randint(10000, 99999)} accepted, {rng.randint(8, 128)} requests")
    log("INFO", f"thermal: die 0 at {rng.randint(62, 74)}C, within envelope")
    log("WARN", f"thermal: die 0 at {rng.randint(84, 90)}C, approaching limit")
    log("ERROR", f"{u.e6}: {u.s6[0]}")
    log("INFO", f"scheduler: batch latency {rng.randint(2, 6)}.{rng.randint(0, 9)}x baseline", 0)
    log("INFO", f"thermal: die 0 at {rng.randint(60, 70)}C, recovered", rng.randint(120, 400))
    log("INFO", f"incident logged against {u.t6}", 0)
    ctx.text(path).lines = lines
    ctx.q("T6", f"Which firmware version fixed ticket {u.t6}?", u.v6, [db], "csv, single-hop",
          aliases=[f"{u.fw} {u.v6}"],
          note="The release notes are the near-miss: topically the most release-ish file in the corpus. "
               f"They list {u.ver_base}.{int(u.v6.split('.')[-1]) - 1} and {u.ver_base}.{int(u.v6.split('.')[-1]) + 1} but not {u.v6}.")
    ctx.q("T6", f"The production log shows a {u.s6[1]} incident. Which firmware release fixed the underlying defect?",
          u.v6, [path, db], "multi-hop, two citations required", aliases=[f"{u.fw} {u.v6}"],
          note=f"The log names the ticket ({u.t6}); only the bug database carries the fixed-in version. "
               "The ticket summary does not share the question's words, so the ticket id in the log is the only route.")
    ctx.q("T6", f"What error code is logged on a {u.s6[1]}?", u.e6, [path], "text log, single-hop",
          note="ERROR line in the production log.")
    return rn


def trap_T7(ctx: Ctx):
    """Fleet log (device id) -> fleet.csv (device -> ticket) -> bug database (ticket -> fixed_in)."""
    u = ctx.u
    db = ensure_bug_db(ctx)
    rng = ctx.rng_fill
    path = f"logs/fleet_controller_{u.d7.isoformat()}.log"
    fleet = "support/fleet.csv"
    devs = {u.dev7: u.t7, u.dev_dec: u.t_dec}
    while len(devs) < 9:
        d = f"DEV-{rng.randint(1000, 9999)}"
        if d not in devs and d != u.nx_dev:
            devs[d] = rng.choice(u.fleet_tickets)
    t = dt.datetime.combine(u.d7, dt.time(14, rng.randint(0, 20), rng.randint(0, 59)))
    lines = []
    ids = list(devs)

    def log(level, msg):
        nonlocal t
        t += dt.timedelta(seconds=rng.randint(3, 50))
        lines.append(f"{t.strftime('%Y-%m-%dT%H:%M:%SZ')} fleet-controller[{rng.randint(300, 900)}]: {level:<5} {msg}")

    for i in range(10):
        if i == 3:
            log("ERROR", f"{u.e7}: {u.dev7} {u.s7[0]}")
        elif i == 6:
            log("ERROR", f"{u.e_dec}: {u.dev_dec} {u.s_dec[0]}")
        else:
            log("INFO", f"{rng.choice(ids)} heartbeat ok, uptime {rng.randint(2, 400)}h")
    ctx.text(path).lines = lines
    rows = ["device_id,site,model,deployed,open_ticket"]
    for d, tk in devs.items():
        rows.append(f"{d},{rng.choice(SITES)},{rng.choice([u.p['p0'], u.p['p0'], u.p['p1']])},"
                    f"2026-0{rng.randint(1, 7)}-{rng.randint(10, 28)},{tk}")
    ctx.text(fleet).lines = rows
    ctx.q("T7", f"The fleet controller log reports a {u.s7[1]} on one device. Which firmware release fixed the underlying defect?",
          u.v7, [path, fleet, db], "multi-hop, three citations required", aliases=[f"{u.fw} {u.v7}"],
          note=f"Chain: log ({u.dev7}) -> support/fleet.csv ({u.dev7} -> {u.t7}) -> bug database ({u.t7} -> {u.v7}). "
               f"A second device ({u.dev_dec}) fails with a different symptom and maps to {u.t_dec}.")


def trap_T8(ctx: Ctx):
    """Same value in the datasheet and in a marketing text."""
    u = ctx.u
    r2, _ = ensure_p0_datasheets(ctx)
    p0 = u.p["p0"]
    mk = f"marketing/{slug(p0)}_product_brief.txt"
    ctx.text(mk).lines = [
        f"{u.company} {p0} - product brief", "",
        f"The {p0} \"{u.codename['p0']}\" is built for dense inference racks.",
        f"Up to {u.pw_new} W board power, with a field-replaceable dual-rotor fan.",
        "Design-in values are in the datasheet; this brief is for orientation only."]
    note = (f"The marketing brief repeats {u.pw_new} W, but the datasheet is the authoritative source; "
            f"the WITHDRAWN r1 says {u.pw_old} W.")
    ctx.q("T8", f"According to the datasheet, what is the board power of the {p0}?", str(u.pw_new), [r2],
          "pdf, value duplicated in a marketing file", aliases=al_units(u.pw_new, "W", "watts"), note=note, also=[mk])
    ctx.q("T8", f"What is the board power of the {p0}?", str(u.pw_new), [r2],
          "pdf, value duplicated in a marketing file (no source named)", aliases=al_units(u.pw_new, "W", "watts"),
          note="AMBIGUITY BY DESIGN: the question names no source. Gold = the datasheet (authoritative); the marketing brief "
               "holds the same number but is not a specification. " + note, also=[mk])


def trap_T9(ctx: Ctx):
    """Short txt holds the answer; a long md on the same topic and an FAQ carry other numbers."""
    u = ctx.u
    rng = ctx.rng_fill
    p0, p1 = u.p["p0"], u.p["p1"]
    short = "support/cooling_quickref.txt"
    ctx.text(short).lines = [
        f"{p0} cooling quick reference", "",
        f"Maximum coolant inlet temperature: {u.c0} C",
        f"Minimum coolant flow: {rng.randint(6, 14)} L/min", "Coolant: 25% propylene glycol mix"]
    md = "engineering/liquid_cooling_whitepaper.md"
    secs = ["Why coolant temperature matters", "Choosing a coolant", "Flow and pressure budgeting",
            "Commissioning a rack", "Monitoring and alarms", "Failure modes"]
    lines = [f"# Liquid cooling in dense accelerator racks", "",
             "This paper summarises how coolant inlet temperature, flow and plumbing interact in racks built from "
             f"{u.company} accelerators. It is a general guide, not a specification."]
    for i, sname in enumerate(secs):
        lines += ["", f"## {sname}", ""]
        for _ in range(7):
            lines.append(rng.choice(SENTENCES))
        if i == 0:
            lines.append(f"For the {p1}, the maximum coolant inlet temperature is {u.c1} C.")
            lines.append(f"Our lab test rigs run the inlet at {u.c_lab} C to leave headroom for sensor tolerance.")
        if i == 3:
            lines.append(f"Racks of {p0} are commissioned with the inlet at {u.c_comm} C and raised gradually.")
    ctx.text(md).lines = lines
    faq = "support/faq.txt"
    ctx.text(faq).lines = [
        f"{u.company} accelerators - frequently asked questions", "",
        f"Q: How hot does the {p0} run under load?",
        f"A: Typical junction temperature under sustained load is {u.jt_typ} C. The maximum is in the datasheet.", "",
        "Q: Where do I find firmware fixes?", "A: In the bug database; release notes are a summary only."]
    ctx.q("T9", f"What is the maximum coolant inlet temperature of the {p0}?", str(u.c0), [short],
          "txt, thematic decoy (long markdown with other numbers)", aliases=al_temp(u.c0),
          note=f"The long white paper (md) gives {u.c1} C for the {p1}, {u.c_lab} C for lab rigs and {u.c_comm} C for commissioning; none is the {p0} maximum.")
    rb = "engineering/operations_runbook.md"
    ctx.text(rb).lines = [
        f"# {u.company} operations runbook", "", "## Thermal events", "",
        f"1. Confirm the alarm on the fleet controller and note the device id.",
        f"2. For the {p0}, set the fan duty floor to {u.floor0}% (`fan_floor_pct={u.floor0}`) until the node is drained.",
        f"3. For the {p1}, set the fan duty floor to {u.floor1}% (`fan_floor_pct={u.floor1}`).",
        "4. Open a ticket in the bug database and attach the controller log.", "",
        "## Escalation", "", "Page the on-call engineer if two nodes in one rack alarm within an hour."]
    ctx.q("T9", f"What fan duty floor, in percent, does the operations runbook set for the {p0}?", str(u.floor0), [rb],
          "md, near-miss step for the sibling product", aliases=[f"{u.floor0}%", f"{u.floor0} percent"],
          note=f"Step 3 gives {u.floor1}% for the {p1}.")


def trap_T10(ctx: Ctx):
    """Image-only facts: connector pin map (png) and asset label (jpg)."""
    from PIL import Image, ImageDraw
    u = ctx.u
    p0 = u.p["p0"]
    # ---- pin map, 1000x560
    im = Image.new("RGB", (1000, 560), "white")
    dr = ImageDraw.Draw(im)
    dr.text((40, 30), f"{p0} backplane connector - pin assignment", fill="black", font=_font(30))
    dr.rectangle((40, 90, 960, 300), outline="black", width=3)
    for i, (pin, sig) in enumerate(zip(u.pin_row, u.pin_sigs)):
        x = 60 + 150 * i
        dr.rectangle((x, 120, x + 118, 190), outline="black", width=2)
        f = _font(28)
        w = dr.textlength(pin, font=f)
        dr.text((x + 59 - w / 2, 137), pin, fill="black", font=f)
        dr.line((x + 59, 190, x + 59, 212), fill="black", width=2)
        f = _font(18)
        w = dr.textlength(sig, font=f)
        dr.text((x + 59 - w / 2, 220), sig, fill="black", font=f)
    blurb = {"THERM_ALERT#": "is asserted low when the die exceeds the warning threshold.",
             "PWR_GOOD": "is driven high once all rails are within tolerance.",
             "FAULT#": "is asserted low on any latched hardware fault.",
             "WAKE#": "is pulled low by the host to leave the low-power state.",
             "PERST#": "is the fundamental reset, active low.",
             "SMB_ALERT#": "is an open-drain interrupt from the management bus."}[u.pin_sig]
    dr.text((40, 330), f"{u.pin_sig} {blurb}", fill="black", font=_font(22))
    dr.text((40, 370), "Signal levels are 3.3 V unless noted.", fill="black", font=_font(22))
    dr.text((40, 435), f"Sheet {u.sheet_no} of 9 - mechanical drawings are in a separate package", fill="black",
            font=_font(18))
    d = ctx.doc(f"specs/{slug(p0)}_backplane_pinout.png", "png")
    d.image = im
    ctx.q("T10", f"Which backplane pin carries {u.pin_sig} on the {p0}?", u.pin, [d.path],
          "png, multimodal - unanswerable without vision", aliases=[f"pin {u.pin}"],
          note="This value appears nowhere in any text file.")
    # ---- asset label, 760x440
    base = Image.new("RGB", (760, 440), (232, 232, 228))
    lab = Image.new("RGBA", (640, 340), (255, 255, 255, 255))
    ld = ImageDraw.Draw(lab)
    ld.rectangle((0, 0, 639, 339), outline="black", width=3)
    ld.text((28, 22), u.company, fill="black", font=_font(32))
    ld.text((28, 78), f"MODEL   {p0}", fill="black", font=_font(30))
    ld.text((28, 128), f"BOARD REVISION   {u.rev}", fill="black", font=_font(30))
    ld.text((28, 178), f"SN   {u.serial}", fill="black", font=_font(30))
    ld.text((28, 232), f"FCC ID: {u.fcc}", fill="black", font=_font(26))
    ld.text((28, 288), f"MADE IN {u.made_in}", fill="black", font=_font(22))
    x = 420
    for i in range(34):
        w = 2 + (i * 7 + ctx.seed) % 4
        ld.rectangle((x, 236, x + w, 320), fill="black")
        x += w + 2
        if x > 600:
            break
    lab = lab.rotate(2.5, expand=True, resample=Image.BICUBIC)
    base.paste(lab, ((760 - lab.width) // 2, (440 - lab.height) // 2), lab)
    dj = ctx.doc("support/asset_label.jpg", "jpg")
    dj.image = base
    ctx.q("T10", "What board revision is printed on the asset label?", u.rev, [dj.path],
          "jpg, multimodal - unanswerable without vision", note="This value appears nowhere in any text file.")
    ctx.q("T10", "What serial number is printed on the asset label?", u.serial, [dj.path],
          "jpg, multimodal - unanswerable without vision", note="Verbatim identifier read from the picture only.")
    ctx.q("T10", "What FCC ID is printed on the asset label?", u.fcc, [dj.path],
          "jpg, multimodal - unanswerable without vision", note="Verbatim identifier read from the picture only.")


def trap_T11(ctx: Ctx):
    """Value lives only in a picture embedded in a Word document."""
    from PIL import Image, ImageDraw
    u = ctx.u
    p2 = u.p["p2"]
    im = Image.new("RGB", (640, 260), "white")
    dr = ImageDraw.Draw(im)
    rows = [("Milestone", "Target"), ("PDK freeze", u.q_pdk), ("Tape-out", u.q_tapeout),
            ("Bring-up", u.q_bringup), ("Engineering samples", u.q_es)]
    for i, (a, b) in enumerate(rows):
        y = 10 + 50 * i
        dr.rectangle((10, y, 630, y + 48), outline="black", width=2)
        dr.line((330, y, 330, y + 48), fill="black", width=2)
        f = _font(26)
        dr.text((22, y + 9), a, fill="black", font=f)
        dr.text((342, y + 9), b, fill="black", font=f)
    path = f"planning/design_review_{slug(p2)}.docx"
    d = ctx.doc(path, "docx")
    d.blocks = [("h", f"{p2} design review", 1),
                ("p", f"Minutes of the {u.codename['p2']} design review. The milestone schedule agreed in the meeting is shown in Figure 1."),
                ("img", im),
                ("p", "Figure 1 - milestone schedule (as presented)."),
                ("p", "Open items: backplane compatibility study, thermal solution sign-off.")]
    ctx.q("T11", f"In which quarter is the tape-out of the {p2} scheduled?", u.q_tapeout, [path],
          "docx, image inside a Word document (value only in the picture)",
          aliases=[u.q_tapeout.replace(" ", "")],
          note="The schedule is a picture embedded in the docx; the value appears in no text.")


def ensure_roadmap(ctx: Ctx):
    """planning/roadmap_fy<FY>.docx (the real Word table); shared by T12, T17b, T18b."""
    u = ctx.u
    path = f"planning/roadmap_fy{u.fy}.docx"
    if path in ctx.docs:
        return path
    p0, p1, p2, p4 = (u.p[k] for k in ("p0", "p1", "p2", "p4"))
    d = ctx.doc(path, "docx")
    d.blocks = [
        ("h", f"{u.company} - Accelerator Roadmap (Internal)", 1),
        ("p", "This document is internal and forward-looking. Dates are targets, not commitments."),
        ("h", "Current generation", 2),
        ("p", f"{p0} \"{u.codename['p0']}\" is in volume production. Remaining work is firmware only."),
        ("h", "Schedule", 2),
        ("table", [["Product", "Status", "Customer sampling", "Volume production"],
                   [p0, "In volume production", "-", "Shipping"],
                   [p1, "End of sale", "-", "-"],
                   [p2, "In development", u.q_sample_p2, u.q_vol_p2],
                   [p4, "Architecture study", u.q_sample_p4, u.q_vol_p4]]),
        ("h", "Beyond", 2),
        ("p", "Later parts are under study. Nothing is committed and no date should be quoted to customers.")]
    return path


def trap_T12(ctx: Ctx):
    """Roadmap docx with a real Word table."""
    u = ctx.u
    p2, p4 = u.p["p2"], u.p["p4"]
    path = ensure_roadmap(ctx)
    ctx.q("T12", f"In which quarter does the {p2} enter customer sampling?", u.q_sample_p2, [path],
          "docx, table lookup", aliases=[u.q_sample_p2.replace(" ", "")],
          note=f"Value is in the Word table, not in a paragraph. The {p4} row ({u.q_sample_p4}) is the near-miss.")
    ctx.q("T12", f"In which quarter does the {p4} enter volume production?", u.q_vol_p4, [path],
          "docx, table lookup", aliases=[u.q_vol_p4.replace(" ", "")],
          note=f"Different column of the same table; the {p2} row ({u.q_vol_p2}) is the near-miss.")


# ------------------------------------------------------------------ v1 traps (T13 - T24)
# Every v1 trap draws its values from ctx.xr(<name>) so they never depend on which other traps ran.

def xpart(r: random.Random, u, cat: str) -> str:
    return f"{u.code3}-{cat}-{r.randint(1000, 9999)}-{r.choice('ABCDEFG')}"


def al_dedupe(gold: str, aliases) -> list:
    return [a for a in dict.fromkeys(aliases) if a != gold]


def _bg_noise(w: int, h: int, seed: int, sigma: float):
    """Deterministic gaussian noise layer (numpy, seeded) as an 'L' image."""
    import numpy as np
    from PIL import Image
    g = np.random.default_rng(seed)
    a = np.clip(g.normal(128, sigma, (h, w)), 0, 255).astype("uint8")
    return Image.fromarray(a, "L")


def trap_T13(ctx: Ctx):
    """xlsx: formulas without cached values, merged title, number formats, hidden stale sheet, header on row 3."""
    u, r = ctx.u, ctx.xr("T13")
    p2 = u.p["p2"]
    names = ["Rack rail kit", "Manifold block", "Quick-disconnect coupler", "Blind-mate backplane",
             "Power shelf", "Cable management arm"]
    r.shuffle(names)
    items, totals = [], set()
    while len(items) < 5:
        qty, unit = r.randint(6, 48), r.randrange(250 * 4, 900 * 4) / 4
        if qty * unit in totals:
            continue
        totals.add(qty * unit)
        items.append({"name": names[len(items)], "part": xpart(r, u, "BOM"), "qty": qty, "unit": unit,
                      "total": qty * unit})
    tgt = items[r.randrange(5)]
    legacy = {"name": "Legacy air shroud", "part": xpart(r, u, "SHR"), "qty": r.randint(6, 30),
              "unit": r.randrange(60 * 4, 200 * 4) / 4}
    fm = lambda v: f"{v:,.2f}"
    hdr = ["item", "part_number", "qty", "unit_cost_usd", "total_usd"]
    rows = [(i["name"], i["part"], i["qty"], i["unit"], f"=C{3 + k}*D{3 + k}") for k, i in enumerate(items)]
    shown = [(i["name"], i["part"], i["qty"], fm(i["unit"]), fm(i["total"])) for i in items]
    stale = [(i["name"], i["qty"], i["unit"], i["total"] + r.choice([-1, 1]) * r.randrange(100, 900, 25))
             for i in items]
    stale.append((legacy["name"], legacy["qty"], legacy["unit"], legacy["qty"] * legacy["unit"]))
    # attach rate: units a multiple of 2000 and a x.5 percentage keep the accessory counts integral
    quarters = [f"Q{q} FY{u.fy - 1}" for q in range(1, 5)]
    qrows, qshown, pcts = [], [], []
    for k, qn in enumerate(quarters):
        units = r.choice([2000, 4000, 8000])
        pct = r.randint(3, 38) + 0.5
        acc = int(units * pct / 100)
        pcts.append(pct)
        qrows.append((qn, units, acc, f"=C{4 + k}/B{4 + k}"))
        qshown.append((qn, units, acc, f"{pct:.1f}%"))
    d = ctx.doc(f"planning/{slug(p2)}_rack_bom.xlsx", "xlsx")
    d.sheets = [
        {"name": "BOM", "title": f"{u.company} - {p2} rack bill of materials (FY{u.fy})", "header": hdr,
         "rows": rows, "shown": shown, "fmt": {"unit_cost_usd": "#,##0.00", "total_usd": "#,##0.00"}},
        {"name": "Quarterly", "title": f"{p2} accessory attach rate by quarter",
         "pre": [["Figures from the sales operations system; accessories are units sold with an expansion kit."]],
         "header": ["quarter", "units", "accessories", "attach_rate"], "rows": qrows, "shown": qshown,
         "fmt": {"attach_rate": "0.0%"}},
        {"name": "Prior_Estimate", "header": ["item", "qty", "unit_cost_usd", "total_usd"], "rows": stale,
         "shown": [(a, b, fm(c), fm(t)) for a, b, c, t in stale], "hidden": True,
         "fmt": {"unit_cost_usd": "#,##0.00", "total_usd": "#,##0.00"}}]
    ctx.cache["T13"] = {"legacy": legacy, "path": d.path, "tgt": tgt}
    tot = fm(tgt["total"])
    ctx.q("T13", f"What is the total cost in USD of the {tgt['name']} line in the {p2} rack bill of materials?", tot,
          [d.path], "xlsx, formula without cached value + number format",
          aliases=al_dedupe(tot, [f"${tot}", f"{tgt['total']:.2f}", f"{tgt['total']:g}", f"USD {tot}"]),
          note=f"total_usd is the formula =C*D ({tgt['qty']} x {fm(tgt['unit'])}) with no cached value; the sheet displays {tot} "
               f"(#,##0.00). The hidden Prior_Estimate sheet holds a stale hard-coded total for the same item.")
    k = r.randrange(4)
    ctx.q("T13", f"What was the accessory attach rate in {quarters[k]} according to the {p2} rack workbook?",
          f"{pcts[k]:.1f}%", [d.path], "xlsx, formula + percent format, header on row 3 (second sheet)",
          aliases=[f"{pcts[k]:.1f}"],
          note=f"attach_rate is the formula =accessories/units with no cached value, formatted 0.0%; the Quarterly sheet has a merged "
               f"title on row 1, a note on row 2 and its header on row 3.")
    ctx.q("T13", f"What is the unit cost in USD of the {tgt['name']} in the {p2} rack bill of materials?", fm(tgt["unit"]),
          [d.path], "xlsx, number format #,##0.00 (constant cell)",
          aliases=al_dedupe(fm(tgt["unit"]), [f"${fm(tgt['unit'])}", f"{tgt['unit']:.2f}", f"{tgt['unit']:g}"]),
          note="Plain numeric cell displayed with a thousands separator.")


def p2_spec(ctx: Ctx):
    """specs/<p2>_preliminary_spec.docx: a Word table with differently formatted values (T14b, T17b)."""
    if "p2spec" in ctx.cache:
        return ctx.cache["p2spec"]
    u, r = ctx.u, ctx.xr("p2spec")
    p2, p4 = u.p["p2"], u.p["p4"]
    s = SimpleNamespace(path=f"specs/{slug(p2)}_preliminary_spec.docx")
    s.price = r.randint(1000, 3999)
    s.yld = f"{r.randint(5, 94)}.5"
    s.bw = f"{r.uniform(1.2, 3.9):.1f}"
    s.tmin = r.choice([-20, -25, -30, -40, -55])
    s.drv = f"{r.randint(2, 9)}.{r.randint(6, 9)}.{r.randint(1, 9)}"
    s.drv_p4 = f"{r.randint(2, 9)}.{r.randint(6, 9)}.{r.randint(1, 9)}"
    while s.drv_p4 == s.drv:
        s.drv_p4 = f"{s.drv_p4[:-1]}{r.randint(1, 9)}"
    taken = {u.q_sample_p2, u.q_vol_p2}
    while True:
        s.kitq = f"FY{u.fy} Q{r.randint(1, 4)}"
        if f"Q{s.kitq[-1]} FY{u.fy}" not in taken:
            break
    s.lanes, s.aux_w = r.choice([8, 16]), r.choice([450, 600])
    minus = "−" if s.tmin < 0 else ""
    d = ctx.doc(s.path, "docx")
    d.blocks = [
        ("h", f"{u.company} {p2} \"{u.codename['p2']}\" - preliminary specification", 1),
        ("p", "Preliminary and subject to change. Values apply to the developer kit unless noted."),
        ("table", [["Parameter", "Value", "Note"],
                   ["Developer kit list price", f"${s.price:,}.00", "per kit, excluding tax"],
                   ["Memory bandwidth", f"{s.bw} TB/s", "peak"],
                   ["Target production yield", f"{s.yld}%", "first production lot"],
                   ["Minimum operating temperature", f"{minus}{abs(s.tmin)} C", "ambient"],
                   ["Required driver", f"v{s.drv}", "host package"],
                   ["Developer kit availability", s.kitq, "ships to registered partners"],
                   ["Host interface", f"PCIe 5.0, {s.lanes} lanes", "x16 slot, bifurcation supported"],
                   ["Expansion interface", "PCIe 4.0, 4 lanes", "card edge"],
                   ["Auxiliary power", f"12V-2x6 connector, {s.aux_w} W maximum", "single cable"],
                   ["Legacy auxiliary power", "8-pin EPS, 300 W maximum", "adapter cable"]]),
        ("p", f"The {p4} is a separate architecture study and is not covered by this document."),
        ("p", f"Required driver for the {p4} developer kit, once defined, will be v{s.drv_p4}.")]
    ctx.cache["p2spec"] = s
    return s


def trap_T14b(ctx: Ctx):
    """Value formats: currency with separator, percent, unit suffix, negative number, v-prefixed version, FY-first quarter."""
    u = ctx.u
    s = p2_spec(ctx)
    p2 = u.p["p2"]
    price = f"{s.price:,}.00"
    ctx.q("T14b", f"What is the list price in USD of the {p2} developer kit?", price, [s.path],
          "docx table, currency format with thousands separator",
          aliases=[f"${price}", f"{s.price}", f"{s.price}.00"],
          note=f"Printed as ${price}; keep the comma. Aliases cover the bare number forms.")
    ctx.q("T14b", f"What is the target production yield of the {p2}?", f"{s.yld}%", [s.path],
          "docx table, percent", aliases=[s.yld], note=f"Printed as {s.yld}%.")
    ctx.q("T14b", f"What is the peak memory bandwidth of the {p2}, in TB/s?", s.bw, [s.path],
          "docx table, unit suffix", aliases=[f"{s.bw} TB/s"], note=f"Printed as {s.bw} TB/s; the unit is not part of the answer.")
    neg = f"-{abs(s.tmin)}"
    ctx.q("T14b", f"What is the minimum operating temperature of the {p2}, in degrees Celsius?", neg, [s.path],
          "docx table, negative number (unicode minus in the source)",
          aliases=al_dedupe(neg, [f"{neg} C", f"{neg}°C", f"−{abs(s.tmin)}", f"−{abs(s.tmin)} C"]),
          note=f"Printed with a U+2212 minus sign: '−{abs(s.tmin)} C'.")
    ctx.q("T14b", f"Which driver version does the {p2} developer kit require?", s.drv, [s.path],
          "docx table, version with a v prefix", aliases=[f"v{s.drv}"],
          note=f"Printed as v{s.drv}; the sibling {u.p['p4']} line shows v{s.drv_p4}.")
    qn = s.kitq.split()[1]
    ctx.q("T14b", f"In which quarter does the {p2} developer kit become available?", s.kitq, [s.path],
          "docx table, quarter written FY-first",
          aliases=[f"{qn} FY{u.fy}", f"{qn}FY{u.fy}", s.kitq.replace(" ", "")],
          note=f"Printed 'FY{u.fy} {qn}' (not 'Q.. FY..'); the roadmap uses the other order for different quarters.")


def trap_T15b(ctx: Ctx):
    """Token boundary: the short token is a prefix of longer tokens in the same files."""
    u, r = ctx.u, ctx.xr("T15b")
    p2 = u.p["p2"]
    # --- tickets HLX-204 vs HLX-2041 / HLX-2047
    n = r.randint(100, 999)
    short, d1, d2 = f"{u.tix}-{n}", f"{u.tix}-{n}1", f"{u.tix}-{n}7"
    comps = r.sample(["fabric controller", "coolant pump", "power shelf", "BMC board", "fan tray", "backplane",
                      "heatsink retention", "PSU backplane"], 6)
    days = r.sample(range(3, 70), 6)
    owners = ["Okafor", "Lindqvist", "Tanaka", "Moreau", "Ibrahim", "Kowalski"]
    sites = r.sample(SITES, 6)
    ids = [short, d1, d2]
    while len(ids) < 6:
        t = f"{u.tix}-{r.randint(100, 999)}"
        if t not in ids and t[len(u.tix) + 1:] != str(n):
            ids.append(t)
    rows = [(t, sites[i], comps[i], r.choice(owners), days[i], r.choice(["closed", "closed", "open"]))
            for i, t in enumerate(ids)]
    r.shuffle(rows)
    path = "support/escalation_log.csv"
    ctx.text(path).lines = ["ticket,site,component,owner,days_to_resolve,status"] + [",".join(map(str, x)) for x in rows]
    ctx.q("T15b", f"How many days did it take to resolve ticket {short}?", str(days[0]), [path],
          "csv, token boundary (ticket id is a prefix of longer ids)", aliases=[f"{days[0]} days"],
          note=f"{d1} ({days[1]} days) and {d2} ({days[2]} days) contain '{short}' as a prefix; only the row for {short} counts.")
    ctx.q("T15b", f"Which ticket was raised for the {comps[0]} at the {sites[0]} site?", short, [path],
          "csv, token boundary (the answer is a prefix of longer ids)",
          note=f"The answer '{short}' also begins the ids {d1} and {d2} (other components); verbatim, not truncated or extended.")
    ctx.absent.append(f"{u.tix}-{n // 10}")
    ctx.cache["T15b_ticket"] = (short, n)
    # --- bootloader versions 5.0.1 vs 5.0.10
    major = r.choice([m for m in range(2, 10) if str(m) != u.ver_base.split(".")[0]])
    minor = r.randint(0, 3)
    base = f"{major}.{minor}"
    vers = [f"{base}.1", f"{base}.10", f"{base}.11", f"{base}.12", f"{base}.2"]
    mbs = r.sample(range(180, 640), 5)
    dts = sorted(dt.date(2026, r.randint(1, 8), r.randint(1, 28)) for _ in range(5))
    while len(set(dts)) < 5:
        dts = sorted(dt.date(2026, r.randint(1, 8), r.randint(1, 28)) for _ in range(5))
    vrows = [(v, dts[i].isoformat(), mbs[i], "yes") for i, v in enumerate(vers)]
    vpath = "engineering/bootloader_builds.csv"
    ctx.text(vpath).lines = ["version,release_date,image_mb,signed"] + [",".join(map(str, x)) for x in vrows]
    ctx.q("T15b", f"What is the image size, in MB, of bootloader version {vers[0]}?", str(mbs[0]), [vpath],
          "csv, token boundary (version is a prefix of longer versions)", aliases=[f"{mbs[0]} MB"],
          note=f"{vers[1]}, {vers[2]} and {vers[3]} share the prefix {vers[0]} but are different builds.")
    ctx.q("T15b", f"Which bootloader version was released on {dts[0].isoformat()}?", vers[0], [vpath],
          "csv, token boundary (the answer is a prefix of longer versions)",
          note=f"The answer {vers[0]} is the first characters of {vers[1]}/{vers[2]}/{vers[3]}; keep it exact.")
    ctx.absent.append(f"{base}.100")
    ctx.cache["T15b_ver"] = (base, vers)
    # --- connector pins A3 vs A34
    L = r.choice("KLMN")
    k = r.randint(3, 9)
    pins = [f"{L}{k}", f"{L}{k}4", f"{L}{k}0"]
    while len(pins) < 8:
        pn = f"{L}{r.randint(1, 30)}"
        if pn not in pins and not pn.startswith(f"{L}{k}"):
            pins.append(pn)
    sigs = r.sample(["PERST#", "WAKE#", "CLKREQ#", "SMB_CLK", "SMB_DAT", "PRSNT1#", "PRSNT2#", "REFCLK_P",
                     "REFCLK_N", "LED_FAULT", "LED_PWR", "BOOT_SEL", "JTAG_TCK", "JTAG_TMS"], 8)
    prow = [(pn, sigs[i], r.choice(["in", "out", "bidir"]), r.choice(["3.3 V", "1.8 V", "12 V"]))
            for i, pn in enumerate(pins)]
    r.shuffle(prow)
    ppath = f"specs/{slug(p2)}_expansion_connector.csv"
    ctx.text(ppath).lines = ["pin,signal,direction,level"] + [",".join(x) for x in prow]
    ctx.q("T15b", f"Which signal is on pin {pins[0]} of the {p2} expansion connector?", sigs[0], [ppath],
          "csv, token boundary (pin is a prefix of longer pins)",
          note=f"Pins {pins[1]} ({sigs[1]}) and {pins[2]} ({sigs[2]}) start with {pins[0]}.")
    ctx.q("T15b", f"Which pin of the {p2} expansion connector carries {sigs[0]}?", pins[0], [ppath],
          "csv, token boundary (the answer is a prefix of longer pins)", aliases=[f"pin {pins[0]}"],
          note=f"The answer {pins[0]} is also the start of {pins[1]} and {pins[2]}.")
    ctx.absent.append(f"{L}{k}5")
    ctx.cache["T15b_pin"] = (L, k)


def trap_T17b(ctx: Ctx):
    """False-premise variants, and answerable questions that merely look suspicious (must NOT be refused)."""
    u = ctx.u
    p0, p2, p4 = u.p["p0"], u.p["p2"], u.p["p4"]
    r2, _ = ensure_p0_datasheets(ctx)
    road = ensure_roadmap(ctx)
    s = p2_spec(ctx)
    ctx.q("T17b", f"According to roadmap_fy{u.fy}.docx, in which quarter does the {p4} enter customer sampling?",
          u.q_sample_p4, [road], "docx table, question names the file (answerable)",
          aliases=[u.q_sample_p4.replace(" ", "")],
          note="The file name in the question is real; naming a file must not make the question unanswerable. "
               f"The {p2} row ({u.q_sample_p2}) is the near-miss.")
    ctx.q("T17b", f"According to {Path(r2).name}, what is the maximum junction temperature of the {p0}?", str(u.jt_new),
          [r2], "pdf, question names the current datasheet file (answerable)", aliases=al_temp(u.jt_new),
          note="Names the r2 file explicitly; the WITHDRAWN r1 says " + f"{u.jt_old} C.")
    ctx.q("T17b", f"In the {p0} datasheet r2, what is the capacity of the HBM3e memory, in GiB?", str(u.cap), [r2],
          "pdf, unindexable tokens in the question (HBM3e, r2) - answerable", aliases=al_units(u.cap, "GiB"),
          note=f"'HBM3e' and 'r2' are real tokens of the datasheet; r1 (withdrawn) says {u.cap_old} GiB.")
    ctx.q("T17b", f"How many lanes does the PCIe 5.0 host interface of the {p2} use?", str(s.lanes), [s.path],
          "docx table, unindexable token PCIe 5.0 in the question - answerable", aliases=[f"{s.lanes} lanes"],
          note="The expansion interface is PCIe 4.0 with 4 lanes (near-miss row).")
    ctx.q("T17b", f"What is the maximum power, in watts, of the 12V-2x6 auxiliary connector on the {p2}?", str(s.aux_w),
          [s.path], "docx table, unindexable token 12V-2x6 in the question - answerable", aliases=[f"{s.aux_w} W"],
          note="The legacy 8-pin EPS row (300 W) is the near-miss.")
    ctx.q("T17b", f"What is the list price in USD of the {p4} developer kit?", "", [], "refusal - attribute belongs to a sibling only",
          unanswerable=True,
          note=f"The {p4} exists (roadmap, and the {p2} spec says it is a separate study) but only the {p2} kit has a price (${s.price:,}.00).")
    ctx.q("T17b", f"What is the maximum junction temperature of the {p0} according to datasheet r3?", "", [],
          "refusal - false premise (revision r3 does not exist)", unanswerable=True,
          note="Only r1 (withdrawn) and r2 exist; r3 is plausible but absent. Answering with the r2 value is wrong.")
    ctx.q("T17b", f"What is the board power of the {p0} according to {slug(p0)}_datasheet_r3.pdf?", "", [],
          "refusal - false premise (file named in the question does not exist)", unanswerable=True,
          note="The named file is not in the corpus.")


def trap_T18b(ctx: Ctx):
    """Supersession variants: r1 withdrawn / r2 superseded in its title only / r3 current + errata; DRAFT roadmap; cancelled CSV row."""
    u, r = ctx.u, ctx.xr("T18b")
    p0, p2, p4 = u.p["p0"], u.p["p2"], u.p["p4"]
    smc = f"SMC-{r.randint(200, 899)}"
    sm = slug(smc)
    poll = r.sample([50, 100, 200, 250, 500, 1000], 3)
    wd = r.sample([5, 8, 10, 12, 15, 20, 30], 4)
    fans = r.sample([2, 4, 6, 8, 10, 12], 3)
    baud, jit = r.choice([9600, 19200, 57600, 115200]), r.choice([20, 40, 80, 160])
    hdr = f"{u.company} {smc} management controller (fitted to the {p0})"
    p1_, p2_, p3_ = f"specs/{sm}_spec_r1_WITHDRAWN.pdf", f"specs/{sm}_spec_r2.pdf", f"specs/{sm}_spec_r3.pdf"
    er = f"specs/errata_{sm}.txt"
    ctx.doc(p1_, "pdf").blocks = [
        ("title", hdr), ("p", "Specification, revision 1 - WITHDRAWN, SUPERSEDED BY REVISION 2"),
        ("h", "Controller parameters"),
        ("l", f"Sensor poll interval .................. {poll[0]} ms"),
        ("l", f"Watchdog timeout ...................... {wd[0]} s"),
        ("l", f"Maximum fan channels ................... {fans[0]}"),
        ("l", f"Debug UART baud rate .................. {baud} baud")]
    ctx.doc(p2_, "pdf").blocks = [
        ("title", hdr), ("p", "Specification, revision 2 - SUPERSEDED BY REVISION 3. Do not use for new designs."),
        ("h", "Controller parameters"),
        ("l", f"Sensor poll interval .................. {poll[1]} ms"),
        ("l", f"Watchdog timeout ...................... {wd[1]} s"),
        ("l", f"Maximum fan channels ................... {fans[1]}"),
        ("l", f"Heartbeat jitter budget ............... {jit} us")]
    ctx.doc(p3_, "pdf").blocks = [
        ("title", hdr), ("p", "Specification, revision 3 - supersedes revision 2"),
        ("h", "Controller parameters"),
        ("l", f"Sensor poll interval .................. {poll[2]} ms"),
        ("l", f"Watchdog timeout ...................... {wd[2]} s"),
        ("l", f"Maximum fan channels ................... {fans[2]}"),
        ("h", "Revision history"),
        ("l", "r3  Poll interval and watchdog re-characterised; UART and jitter entries removed."),
        ("l", "r2  Added the heartbeat jitter budget."), ("l", "r1  Initial release (withdrawn).")]
    ctx.text(er).lines = [
        f"Errata sheet - {smc} specification revision 3", "Issued after the r3 release.", "",
        f'E1. "Watchdog timeout": value corrected to {wd[3]} s (supersedes the specification r3 figure).',
        "E2. No other changes to revision 3."]
    ctx.q("T18b", f"What is the sensor poll interval of the {smc} management controller, in milliseconds?", str(poll[2]),
          [p3_], "pdf, r1 withdrawn / r2 superseded in title only / r3 current", aliases=[f"{poll[2]} ms"],
          note=f"r1 (WITHDRAWN in name and title) says {poll[0]} ms; r2 (name has no marker, title says SUPERSEDED BY REVISION 3) "
               f"says {poll[1]} ms; r3 is current.")
    ctx.q("T18b", f"What is the watchdog timeout of the {smc} management controller, in seconds?", str(wd[3]), [er],
          "errata overrides one field of the current revision", aliases=[f"{wd[3]} s", f"{wd[3]} seconds"], also=[p3_],
          note=f"r1={wd[0]}, r2={wd[1]}, r3={wd[2]}, errata corrects the field to {wd[3]}. As in T2 the errata alone holds the value "
               "and names the field, so only the errata is cited.")
    ctx.q("T18b", f"What is the maximum number of fan channels of the {smc} management controller?", str(fans[2]), [p3_],
          "current revision, field untouched by the errata", note=f"r1 says {fans[0]}, r2 says {fans[1]}; the errata does not touch this field.")
    # --- DRAFT roadmap next to the final one
    road = ensure_roadmap(ctx)
    dq = [f"Q{q} FY{y}" for y in (u.fy + 2, u.fy + 3) for q in range(1, 5)]
    r.shuffle(dq)
    num4 = int(u.p["p4"].split("-")[1])
    draft_prod = f"{u.pre}-{num4 - (20 if num4 >= 100 else 2)}"       # (n0 + 38) x scale: unique among the line
    ctx.cache["T18b"] = {"draft_prod": draft_prod, "dq": dq}
    dpath = f"planning/roadmap_fy{u.fy}_DRAFT.docx"
    dd = ctx.doc(dpath, "docx")
    dd.blocks = [
        ("h", f"{u.company} - Accelerator Roadmap (DRAFT)", 1),
        ("p", "Working draft for the planning offsite; dates are placeholders under discussion."),
        ("h", "Schedule", 2),
        ("table", [["Product", "Status", "Customer sampling", "Volume production"],
                   [p2, "In development", dq[0], dq[1]],
                   [p4, "Architecture study", dq[2], dq[3]],
                   [draft_prod, "Concept", dq[4], dq[5]]])]
    ctx.q("T18b", f"In which quarter does the {p2} enter volume production?", u.q_vol_p2, [road],
          "docx, DRAFT roadmap next to the approved one", aliases=[u.q_vol_p2.replace(" ", "")],
          note=f"planning/roadmap_fy{u.fy}_DRAFT.docx (a draft: only the name and one heading say so) lists {dq[1]} for the same cell.")
    # --- cancelled CSV row next to the live one
    pr, pcn = xpart(r, u, "RAL"), xpart(r, u, "MNF")
    q_live, q_can = r.randrange(100, 900, 25), r.randrange(100, 900, 25)
    while q_can == q_live:
        q_can = r.randrange(100, 900, 25)
    w_live, w_can = r.randint(28, 44), r.randint(28, 44)
    while w_can == w_live:
        w_can = r.randint(28, 44)
    po = r.sample(range(1000, 9999), 8)
    prow = [(f"PO-{po[0]}", pr, "Rack rail kit", q_live, w_live, "open"),
            (f"PO-{po[1]}", pr, "Rack rail kit", q_can, w_can, "cancelled"),
            (f"PO-{po[2]}", pcn, "Manifold block rev B", r.randrange(100, 900, 25), r.randint(28, 44), "cancelled")]
    for i in range(3, 7):
        prow.append((f"PO-{po[i]}", xpart(r, u, "OTH"), r.choice(["Power shelf", "Cable arm", "Cold plate"]),
                     r.randrange(25, 900, 25), r.randint(28, 44), "open"))
    r.shuffle(prow)
    cpath = "procurement/po_register.csv"
    ctx.text(cpath).lines = ["po,part_number,description,qty,eta_week,status"] + [",".join(map(str, x)) for x in prow]
    ctx.cache["T18b"].update({"cancelled_part": pcn})
    ctx.q("T18b", f"How many units of part {pr} are on order?", str(q_live), [cpath],
          "csv, status=cancelled row next to the live row",
          note=f"The cancelled order PO-{po[1]} for the same part says {q_can}. Only the open row counts.")
    ctx.q("T18b", f"What is the delivery ETA week for part {pr}?", str(w_live), [cpath],
          "csv, status=cancelled row next to the live row", note=f"The cancelled row says week {w_can}.")
    ctx.cache["T18b"].update({"smc": smc, "baud": baud, "jit": jit, "p1": p1_, "p2": p2_})


def trap_T19b(ctx: Ctx):
    """Unit asked differs from the unit in the source (refuse), and a unit that only is written differently (answer)."""
    u, r = ctx.u, ctx.xr("T19b")
    p0 = u.p["p0"]
    r2, _ = ensure_p0_datasheets(ctx)
    tr, amb = r.choice([120, 250, 400, 800]), r.choice([30, 35, 40])
    idle, boot = r.choice([300, 600, 900, 1200]), r.choice([35, 45, 55, 70])
    fa, fb = r.choice([6, 8, 10]), r.choice([14, 16, 18])
    ra, rb = r.choice([2400, 3000]), r.choice([6000, 7200])
    path = f"engineering/{slug(p0)}_operating_limits.md"
    ctx.text(path).lines = [
        f"# {p0} operating limits", "",
        f"- Thermal trip response time: {tr} ms",
        f"- Maximum ambient temperature: {amb} C",
        f"- Idle shutdown timeout: {idle} seconds",
        f"- Cold boot time: {boot} secs",
        f"- Coolant flow: between {fa} and {fb} L/min depending on the fan profile",
        f"- Fan speed: between {ra} and {rb} RPM", "",
        "Values are for a rack at sea level with the reference firmware."]
    ctx.cache["T19b"] = {"fa": fa, "fb": fb, "ra": ra, "rb": rb, "path": path}
    ctx.q("T19b", f"What is the thermal trip response time of the {p0}, in milliseconds?", str(tr), [path],
          "md, unit matches the question (control)", aliases=[f"{tr} ms"], note="Same unit as the source: answerable.")
    ctx.q("T19b", f"How long is the cold boot time of the {p0}, in seconds?", str(boot), [path],
          "md, unit written differently ('secs')", aliases=[f"{boot} secs", f"{boot} seconds", f"{boot}s"],
          note="The source writes 'secs'; the unit is the same, so this must be answered.")
    ctx.q("T19b", f"What is the thermal trip response time of the {p0}, in seconds?", "", [], "refusal - unit differs (ms in the source, s asked)",
          unanswerable=True, note=f"The source gives {tr} ms; converting is not allowed by the contract (no conversion text in the corpus).")
    ctx.q("T19b", f"What is the maximum ambient temperature of the {p0}, in degrees Fahrenheit?", "", [],
          "refusal - unit differs (C in the source, F asked)", unanswerable=True, note=f"The source gives {amb} C only.")
    ctx.q("T19b", f"What is the idle shutdown timeout of the {p0}, in minutes?", "", [],
          "refusal - unit differs (seconds in the source, minutes asked)", unanswerable=True,
          note=f"The source gives {idle} seconds only.")
    ctx.q("T19b", f"What is the memory capacity of the {p0}, in GB?", "", [], "refusal - unit differs (GiB in the source, GB asked)",
          unanswerable=True, note=f"The r2 datasheet gives {u.cap} GiB; GiB != GB and no conversion is stated.")


def trap_T20(ctx: Ctx):
    """Large files: 3,000-line log (answer near line 2,700, 40 decoy errors), 150-page PDF (value on page 120), 5,000-row CSV."""
    u, r = ctx.u, ctx.xr("T20")
    p0, p1 = u.p["p0"], u.p["p1"]
    taken = {u.e6, u.e7, u.e_dec, u.nx_err}

    def code():
        while True:
            c = f"E{r.randint(1000, 9999)}"
            if c not in taken:
                taken.add(c)
                return c
    # ---- log
    sym, dec_syms = u.syms[3], u.syms[5:]
    generic = [("checksum mismatch on firmware slot 2", "firmware checksum mismatch"),
               ("NVMe write latency above 50 ms for 10 s", "storage latency"),
               ("license server unreachable, grace period started", "license check"),
               ("TLS handshake failed with the fleet controller", "TLS handshake failure")]
    pool = [(s[0], s[1]) for s in dec_syms] + generic
    n_lines, tgt = 3000, r.randint(2680, 2720)
    dec_at = set()
    while len(dec_at) < 40:
        i = r.randint(30, n_lines - 1)
        if abs(i - tgt) > 3:
            dec_at.add(i)
    date = u.d6 + dt.timedelta(days=21)
    t = dt.datetime.combine(date, dt.time(0, 5, 0))
    node = f"stress-node-{r.randint(1, 12):02d}"
    lines = []
    for i in range(n_lines):
        t += dt.timedelta(seconds=r.randint(1, 4))
        ts = t.strftime("%Y-%m-%dT%H:%M:%SZ")
        if i == tgt:
            code_t = code()
            lines.append(f"{ts} {node} stressd[{r.randint(1000, 4999)}]: ERROR {code_t}: {sym[1]} ({sym[0]})")
        elif i in dec_at:
            txt, noun = pool[len(lines) % len(pool)]
            lines.append(f"{ts} {node} stressd[{r.randint(1000, 4999)}]: ERROR {code()}: {noun} ({txt})")
        else:
            k = r.randrange(5)
            msg = [f"INFO  scheduler: batch {r.randint(10000, 99999)} accepted, {r.randint(8, 128)} requests",
                   f"INFO  thermal: die {r.randint(0, 3)} at {r.randint(60, 78)}C, within envelope",
                   f"INFO  heartbeat ok, uptime {r.randint(2, 400)}h",
                   f"INFO  telemetry: exported {r.randint(100, 9000)} samples",
                   f"WARN  thermal: die {r.randint(0, 3)} at {r.randint(82, 90)}C, approaching limit"][k]
            lines.append(f"{ts} {node} stressd[{r.randint(1000, 4999)}]: {msg}")
    lpath = f"logs/stress_run_{date.isoformat()}.log"
    ctx.text(lpath).lines = lines
    ctx.q("T20", f"What error code does the stress run log for a {sym[1]}?", code_t, [lpath],
          "log, 3,000 lines (answer at line ~%d) with 40 decoy ERROR lines" % (tgt + 1),
          note=f"The matching ERROR line is line {tgt + 1} of {n_lines}; 40 other ERROR lines carry other codes and symptoms.")
    absent_sym = ("disk full on scratch volume 3", "scratch disk full")
    ctx.cache["T20"] = {"absent": absent_sym, "lpath": lpath}
    ctx.absent.append(absent_sym[1])
    # ---- 150-page PDF
    cycles = r.randrange(2100, 9900, 50)       # a number unlikely to be anywhere else in the corpus
    other = r.sample([c for c in range(2100, 9900, 50) if c != cycles], 2)
    pages = 150
    blocks = [("title", f"{u.company} {p0} integration handbook")]
    for pg in range(1, pages + 1):
        if pg > 1:
            blocks.append(("page", ""))
        blocks.append(("h", f"Chapter {(pg - 1) // 10 + 1}, section {(pg - 1) % 10 + 1}"))
        for _ in range(11):
            blocks.append(("p", r.choice(SENTENCES)))
        if pg == 120:
            blocks.append(("l", f"Backplane connector rated mating cycles for the {p0}: {cycles} cycles"))
        elif pg == 40:
            blocks.append(("l", f"Backplane connector rated mating cycles for the {p1}: {other[0]} cycles"))
        elif pg == 141:
            blocks.append(("l", f"Fan tray rated insertion cycles for the {p0}: {other[1]} cycles"))
        blocks.append(("l", f"Handbook page {pg} of {pages}"))
    hpath = f"specs/{slug(p0)}_integration_handbook.pdf"
    hd = ctx.doc(hpath, "pdf")
    hd.blocks = blocks
    hd.note = "150 pages; the fact is on page 120"
    ctx.q("T20", f"How many mating cycles is the backplane connector of the {p0} rated for?", str(cycles), [hpath],
          "pdf, 150 pages (value on page 120)", aliases=[f"{cycles} cycles"],
          note=f"Page 40 gives {other[0]} for the {p1}; page 141 gives {other[1]} for the fan tray of the {p0}.")
    # ---- 5,000-row CSV
    n_rows = 5000
    seen, ids = set(), []
    while len(ids) < n_rows:
        x = f"RMA-{r.randint(10000, 99999)}"
        if x not in seen:
            seen.add(x)
            ids.append(x)
    part_pool = [xpart(r, u, "SPR") for _ in range(30)]
    tgt_i = r.randint(4800, 4950)
    tgt_part = xpart(r, u, "SPR")
    while tgt_part in part_pool:
        tgt_part = xpart(r, u, "SPR")
    rows = ["rma_id,site,product,reason,replacement_part,days_open"]
    for i, rid in enumerate(ids):
        rows.append(f"{rid},{r.choice(SITES)},{r.choice([p0, p1, u.p['p3']])},"
                    f"{r.choice(['dead on arrival', 'fan failure', 'link errors', 'coolant leak', 'no boot'])},"
                    f"{tgt_part if i == tgt_i else r.choice(part_pool)},{r.randint(1, 60)}")
    cpath = "support/field_returns.csv"
    ctx.text(cpath).lines = rows
    ctx.q("T20", f"Which replacement part was shipped for {ids[tgt_i]}?", tgt_part, [cpath],
          "csv, 5,000 rows (ticket near the end)",
          note=f"The row is number {tgt_i + 2} of {n_rows + 1} lines in the file.")
    while True:
        x = f"RMA-{r.randint(10000, 99999)}"
        if x not in seen:
            break
    ctx.absent.append(x)
    ctx.cache["T20"].update({"absent_rma": x})


def trap_T21(ctx: Ctx):
    """Encodings and formats: UTF-16+BOM, ';' CSV with CRLF and decimal commas, markdown table, CRLF .log, YAML, JSON."""
    u, r = ctx.u, ctx.xr("T21")
    p0 = u.p["p0"]
    # ---- UTF-16 with BOM
    site = r.choice(SITES)
    racks, fl = r.randint(12, 96), r.randrange(800, 1600, 50)
    ctx.cache["T21"] = {"site": site}
    sp = f"facilities/{slug(site)}_site_survey.txt"
    d = ctx.text(sp)
    d.encoding = "utf-16"
    d.lines = [f"Site survey - {site} lab", "", "Contact: André Müller, facilities",
               f"Maximum racks permitted: {racks}", f"Floor loading: {fl} kg per square metre",
               "Chilled water supply: available on level 2"]
    ctx.q("T21", f"What is the maximum number of racks permitted at the {site} lab?", str(racks), [sp],
          "txt, UTF-16 with BOM", aliases=[f"{racks} racks"], note=f"The file is UTF-16 (BOM); floor loading ({fl}) is the near-miss number.")
    # ---- CSV ';' + CRLF + decimal commas
    sens = r.sample(range(100, 999), 5)
    gains = r.sample([f"{x / 100:.2f}".replace(".", ",") for x in range(50, 300, 7)], 10)
    rows = ["sensor_id;channel;gain;offset_mv;range_v"]
    k = 0
    for sid in sens:
        for ch in (1, 2):
            rows.append(f"TS-{sid};{ch};{gains[k]};{r.randint(-20, 20)};{r.choice(['2,5', '5', '10'])}")
            k += 1
    cp = "support/sensor_calibration.csv"
    c = ctx.text(cp)
    c.newline, c.lines = "\r\n", rows
    g = gains[3]
    ctx.q("T21", f"What is the gain of sensor TS-{sens[1]} on channel 2?", g, [cp],
          "csv, ';' delimiter + CRLF + decimal comma", aliases=al_dedupe(g, [g.replace(",", ".")]),
          note=f"Decimal comma as printed ('{g}'); the same sensor's channel 1 gain ({gains[2]}) is the near-miss.")
    # ---- markdown table
    rails = ["VDD_CORE", "VDD_SRAM", "VDDQ_HBM", "VDD_IO", "VDD_PLL"]
    amps = r.sample(range(8, 480), 5)
    noms = r.sample(["0.75", "0.80", "0.85", "0.90", "1.10", "1.20", "1.80"], 5)
    mp = "engineering/power_rails.md"
    ctx.text(mp).lines = [f"# {p0} power rails", "", "| Rail | Nominal (V) | Max current (A) |", "|---|---|---|"] + \
        [f"| {rl} | {noms[i]} | {amps[i]} |" for i, rl in enumerate(rails)] + ["", "Currents are per module."]
    ctx.q("T21", f"What is the maximum current, in amps, of the {rails[2]} rail of the {p0}?", str(amps[2]), [mp],
          "md, table row (header on another line)", aliases=[f"{amps[2]} A"], note="Markdown table: the header is two lines above the row.")
    # ---- CRLF log
    sym = u.syms[4]
    ecode = f"E{r.randint(1000, 9999)}"
    while ecode in {u.e6, u.e7, u.e_dec, u.nx_err}:
        ecode = f"E{r.randint(1000, 9999)}"
    t = dt.datetime.combine(u.d7 + dt.timedelta(days=9), dt.time(6, 12, 0))
    ll = []
    for i in range(14):
        t += dt.timedelta(seconds=r.randint(5, 90))
        lvl, msg = ("ERROR", f"{ecode}: {sym[1]} ({sym[0]})") if i == 8 else \
            ("INFO", r.choice([f"ups-{r.randint(1, 4)} load {r.randint(30, 80)} percent", "mains voltage nominal",
                               f"battery charge {r.randint(88, 100)} percent", "self-test passed"]))
        ll.append(f"{t.strftime('%Y-%m-%d %H:%M:%S')} facility-monitor {lvl:<5} {msg}")
    lp = f"logs/facility_events_{(u.d7 + dt.timedelta(days=9)).isoformat()}.log"
    lg = ctx.text(lp)
    lg.newline, lg.lines = "\r\n", ll
    ctx.q("T21", f"What error code does the facility monitor log for a {sym[1]}?", ecode, [lp],
          "log, Windows line endings (CRLF)", note="Every line ends with CRLF.")
    # ---- YAML
    tel, logf = r.choice([30, 45, 60, 90]), r.choice([5, 10, 15])
    batch = r.choice([256, 512, 1024])
    ya = "engineering/agent_config.yaml"
    ctx.text(ya).lines = [
        "# Fleet agent configuration", "telemetry:", "  exporter:", f"    flush_interval_s: {tel}   # raised from 20",
        f"    batch_size: {batch}", "  retry:", "    max_attempts: 5", "logging:", "  level: info",
        f"  flush_interval_s: {logf}"]
    ctx.q("T21", "What is the flush interval, in seconds, of the telemetry exporter in the agent config?", str(tel), [ya],
          "yaml, nested key (the same key exists under logging)", aliases=[f"{tel} seconds", f"{tel}s"],
          note=f"logging.flush_interval_s is {logf}; the comment names the old value 20.")
    # ---- JSON
    deep, shallow = r.choice([4096, 8192, 16384]), r.choice([128, 256, 512])
    ob = {"queue": {"max_depth": deep, "retry": {"max_attempts": 7}},
          "dead_letter": {"max_depth": shallow, "ttl_s": 3600}, "name": "ingest-limits"}
    jp = "engineering/limits.json"
    ctx.text(jp).lines = json.dumps(ob, indent=2).splitlines()
    ctx.q("T21", "What is the maximum queue depth in limits.json?", str(deep), [jp],
          "json, nested constant (the same key exists under dead_letter)",
          note=f"dead_letter.max_depth is {shallow}; queue.max_depth is the answer.")


def trap_T22(ctx: Ctx):
    """Odd file names; citations must equal the exact relative path."""
    u, r = ctx.u, ctx.xr("T22")
    p0 = u.p["p0"]
    # 1 spaces (docx)
    a = r.randrange(300, 900, 10)
    pth = "planning/Board Review Notes.docx"
    ctx.doc(pth, "docx").blocks = [("h", "Board review notes", 1),
                                    ("p", f"The board approved a thermal lab upgrade budget of {a} thousand USD."),
                                    ("p", "Capital requests above that level go to the next cycle.")]
    ctx.q("T22", "What budget, in thousand USD, did the board approve for the thermal lab upgrade?", str(a), [pth],
          "docx, file name with spaces", aliases=[f"{a} thousand USD"], note="Path contains spaces.")
    # 2 unicode u-umlaut (pdf)
    b = r.randint(6, 18)
    pth = "specs/Übersicht_Kühlung.pdf"
    ctx.doc(pth, "pdf").blocks = [("title", "Uebersicht Kuehlung - cooling overview"),
                                   ("p", f"Design coolant flow per accelerator: {b} L/min."),
                                   ("p", "Higher flow is possible but not required.")]
    ctx.q("T22", "What is the design coolant flow per accelerator in the cooling overview, in L/min?", str(b), [pth],
          "pdf, unicode file name (Ü, ü)", aliases=[f"{b} L/min"], note="Path contains U+00DC and U+00FC (NFC).")
    # 3 unicode e-acute (csv)
    so, qty = f"SO-{r.randint(10000, 99999)}", r.randint(2, 90)
    pth = "support/café_orders.csv"
    ctx.text(pth).lines = ["order_id,site,item,qty"] + \
        [f"SO-{r.randint(10000, 99999)},{r.choice(SITES)},{r.choice(['fan tray', 'cold plate', 'power shelf'])},{r.randint(2, 90)}"
         for _ in range(6)] + [f"{so},{r.choice(SITES)},manifold block,{qty}"]
    ctx.q("T22", f"How many units are on order {so}?", str(qty), [pth], "csv, unicode file name (é)",
          note="Path contains U+00E9 (NFC).")
    # 4 uppercase extension PDF
    c = r.randint(2, 9)
    pth = "specs/CONNECTOR_SHEET.PDF"
    ctx.doc(pth, "pdf").blocks = [("title", f"{p0} connector sheet"), ("p", f"Contact resistance of the power contacts: {c} mOhm."),
                                  ("p", "Measured at 20 C on new contacts.")]
    ctx.q("T22", f"What is the contact resistance, in milliohms, of the {p0} power contacts?", str(c), [pth],
          "pdf, uppercase extension (.PDF)", aliases=[f"{c} mOhm"], note="Extension is upper case.")
    # 5 uppercase extension XLSX
    part = xpart(r, u, "INV")
    oh = r.randint(10, 400)
    pth = "support/INVENTORY_SUMMARY.XLSX"
    ctx.doc(pth, "xlsx").sheets = [{"name": "Stock", "header": ["part_number", "description", "on_hand"],
                                    "rows": [(xpart(r, u, "INV"), "Fan tray", r.randint(10, 400)) for _ in range(5)] +
                                            [(part, "Cold plate", oh)]}]
    ctx.q("T22", f"How many units of part {part} are on hand in the inventory summary?", str(oh), [pth],
          "xlsx, uppercase extension (.XLSX)", note="Extension is upper case.")
    # 6 double extension
    h = r.choice([48, 72, 96, 120])
    pth = "specs/report.final.pdf"
    ctx.doc(pth, "pdf").blocks = [("title", "Acceptance report - final"),
                                   ("p", f"Final acceptance test duration: {h} hours of burn-in."),
                                   ("p", "No failures were recorded.")]
    ctx.q("T22", "According to report.final.pdf, how many hours did the final acceptance burn-in run?", str(h), [pth],
          "pdf, double extension in the name (and in the question)", aliases=[f"{h} hours"], note="Path is specs/report.final.pdf.")
    # 7 directory with a dot
    bid = f"BLD-{r.randint(1000, 9999)}-{r.choice('ABCDEFG')}"
    pth = "v2.0/release_manifest.txt"
    ctx.text(pth).lines = ["Release 2.0 manifest", "", f"Build id: {bid}", "Signed: yes", "Channel: stable"]
    ctx.q("T22", "What is the build id in the 2.0 release manifest?", bid, [pth], "txt, directory name with a dot (v2.0/)",
          note="Directory is 'v2.0'.")
    # 8 very long name
    rpm = r.randrange(2000, 6000, 100)
    pth = "engineering/" + ("very_long_descriptive_file_name_" * 6)[:196] + ".txt"
    ctx.text(pth).lines = ["Pump characterisation", "", f"Maximum pump speed: {rpm} RPM", "Minimum pump speed: 800 RPM"]
    ctx.q("T22", "What is the maximum pump speed, in RPM, in the pump characterisation notes?", str(rpm), [pth],
          "txt, very long file name (%d chars)" % len(pth.split("/")[-1]), aliases=[f"{rpm} RPM"],
          note="The file name is 200 characters long.")


def trap_T23(ctx: Ctx):
    """Multi-hop variants: bridging id in three files (cite the two needed); 3-hop through a serial printed only on a label."""
    from PIL import Image, ImageDraw
    u, r = ctx.u, ctx.xr("T23")
    # ---- (a) work order bridging ids: WO-A in four files (log, work_orders, labor_hours, ops notes), WO-B in three
    wo, wo2 = f"WO-{r.randint(4000, 9999)}", f"WO-{r.randint(1000, 3999)}"
    part, part2 = xpart(r, u, "WOP"), xpart(r, u, "WOP")
    date = u.d6 + dt.timedelta(days=30)
    lp = f"logs/maintenance_{date.isoformat()}.log"
    t = dt.datetime.combine(date, dt.time(9, 0, 0))
    rack = f"R{r.randint(1, 30)}"
    ll = []
    for i in range(10):
        t += dt.timedelta(seconds=r.randint(30, 400))
        if i == 3:
            msg = f"WARN  coolant loop pressure dropped below 1.2 bar on rack {rack}; work order {wo} raised"
        elif i == 7:
            msg = f"WARN  fan tray vibration above limit on rack R{r.randint(31, 40)}; work order {wo2} raised"
        else:
            msg = f"INFO  rack R{r.randint(1, 40)} walk-through complete, {r.randint(0, 3)} items noted"
        ll.append(f"{t.strftime('%Y-%m-%dT%H:%M:%SZ')} facility-ops: {msg}")
    ctx.text(lp).lines = ll
    wp = "support/work_orders.csv"
    wrows = [f"{wo},{date.isoformat()},cooling,{part},closed", f"{wo2},{date.isoformat()},mechanical,{part2},closed"]
    for _ in range(5):
        wrows.append(f"WO-{r.randint(1000, 9999)},{date.isoformat()},{r.choice(['cooling', 'power', 'network'])},"
                     f"{xpart(r, u, 'WOP')},{r.choice(['closed', 'open'])}")
    r.shuffle(wrows)
    ctx.text(wp).lines = ["wo_id,raised,team,replaced_part,status"] + wrows
    hrs = r.randint(2, 19)
    lh = "support/labor_hours.csv"
    ctx.text(lh).lines = ["wo_id,technician,hours"] + [f"{wo},{r.choice(['Okafor', 'Tanaka', 'Moreau'])},{hrs}"] + \
        [f"WO-{r.randint(1000, 9999)},{r.choice(['Okafor', 'Tanaka', 'Moreau'])},{r.randint(1, 20)}" for _ in range(6)]
    nt = "planning/ops_review_notes.txt"
    ctx.text(nt).lines = ["Ops review notes", "",
                          f"Overtime for the night shift was approved against {wo}; no other staffing changes.",
                          f"The {wo2} vibration work was reviewed; the supplier was not charged.",
                          r.choice(SENTENCES), r.choice(SENTENCES)]
    ctx.q("T23", "The maintenance log shows a coolant loop pressure incident. Which part number was replaced under the work "
                 "order raised for it?", part, [lp, wp], "multi-hop, bridging id appears in four files (df = 4)",
          note=f"{wo} appears in the maintenance log, work_orders.csv (the replaced part), labor_hours.csv and ops_review_notes.txt; "
               "only the log and the work-order table are needed. Rare-id hopping with a df cap of 3 cannot follow this id.")
    ctx.q("T23", "The maintenance log shows a fan tray vibration incident. Which part number was replaced under the work "
                 "order raised for it?", part2, [lp, wp], "multi-hop, bridging id appears in three files (df = 3)",
          note=f"{wo2} appears in the maintenance log, work_orders.csv (the replaced part) and ops_review_notes.txt (a decoy "
               "mention); only the log and the work-order table are needed.")
    ctx.q("T23", f"How many labor hours were booked against {wo}?", str(hrs), [lh], "csv, id shared with other files",
          aliases=[f"{hrs} hours"], note=f"{wo} also appears in three other files that do not carry hours.")
    ctx.cache["T23"] = {}
    # ---- (b) serial printed only on a label (image -> csv -> csv)
    ser = f"CS-{r.randint(100000, 999999)}-{r.choice('KMNPQRT')}"
    near = [ser[:-1] + c for c in "ABCD"] + [f"CS-{int(ser[3:9]) + 1}-{ser[-1]}"]
    chassis = r.sample(range(1000, 9999), 7)
    ch = f"CH-{chassis[0]}"
    ctx.forbid.append(ser)
    ctx.secrets[ser] = {"support/chassis_registry.csv"}
    reg = [(ser, ch, r.choice(SITES), "2026-02-11")] + \
          [(near[i], f"CH-{chassis[1 + i]}", r.choice(SITES), f"2026-0{r.randint(1, 6)}-{r.randint(10, 28)}") for i in range(4)]
    rp = "support/chassis_registry.csv"
    r.shuffle(reg)
    ctx.text(rp).lines = ["serial,chassis_id,site,installed"] + [",".join(x) for x in reg]
    new_part, old_part = xpart(r, u, "SVC"), xpart(r, u, "SVC")
    srows = [(ch, "2026-04-02", old_part, "Okafor"), (ch, "2026-07-19", new_part, "Tanaka")]
    for i in range(1, 6):
        srows.append((f"CH-{chassis[i]}", f"2026-0{r.randint(1, 8)}-{r.randint(10, 28)}", xpart(r, u, "SVC"), "Moreau"))
    r.shuffle(srows)
    sp = "support/chassis_service.csv"
    ctx.text(sp).lines = ["chassis_id,service_date,replaced_part,technician"] + [",".join(x) for x in srows]
    im = Image.new("RGB", (800, 420), (222, 224, 228))
    lab = Image.new("RGBA", (680, 340), (255, 255, 255, 255))
    ld = ImageDraw.Draw(lab)
    ld.rectangle((0, 0, 679, 339), outline="black", width=3)
    ld.text((28, 24), f"{u.company}", fill="black", font=_font(30))
    ld.text((28, 82), "CHASSIS ASSEMBLY", fill="black", font=_font(30))
    ld.text((28, 140), f"MODEL   {u.p['p0']}-CH", fill="black", font=_font(28))
    ld.text((28, 196), f"SERIAL   {ser}", fill="black", font=_font(34))
    ld.text((28, 262), "Property of the rack operations team", fill="black", font=_font(20))
    lab = lab.rotate(-1.5, expand=True, resample=Image.BICUBIC)
    im.paste(lab, ((800 - lab.width) // 2, (420 - lab.height) // 2), lab)
    ctx.doc("support/chassis_label.jpg", "jpg").image = im
    ctx.q("T23", "Which part number was last replaced in the chassis whose serial number is printed on the chassis label?",
          new_part, ["support/chassis_label.jpg", rp, sp], "multi-hop, three hops: jpg -> csv -> csv (serial only on the label)",
          note=f"Label serial {ser} -> registry -> {ch} -> service log, whose latest entry (2026-07-19) is {new_part}; "
               f"the older entry is {old_part}. The registry also lists near-miss serials. The serial appears in no file except the registry.")
    ctx.cache["T23"].update({"ch_absent": f"CH-{max(chassis) + 1}", "path_label": "support/chassis_label.jpg"})
    ctx.absent.append(f"CH-{max(chassis) + 1}")


def trap_T24(ctx: Ctx):
    """Image traps: rotated label, near-duplicate signals in a pinout, low contrast, photo-like noise, PNG table inside a docx."""
    import numpy as np
    from PIL import Image, ImageDraw, ImageFilter
    u, r = ctx.u, ctx.xr("T24")
    p3, p4 = u.p["p3"], u.p["p4"]
    # ---- rotated shipping label
    lot = f"L{r.randint(10, 99)}-{r.randint(1000, 9999)}"
    po = f"PO-{r.randint(10000, 99999)}"
    lab = Image.new("RGB", (760, 380), "white")
    ld = ImageDraw.Draw(lab)
    ld.rectangle((0, 0, 759, 379), outline="black", width=4)
    ld.text((30, 24), "SHIPPING LABEL", fill="black", font=_font(34))
    ld.text((30, 96), f"SHIP TO   {r.choice(SITES).upper()} LAB", fill="black", font=_font(30))
    ld.text((30, 156), f"LOT {lot}", fill="black", font=_font(40))
    ld.text((30, 226), f"REF {po}", fill="black", font=_font(30))
    ld.text((30, 290), "HANDLE WITH CARE - ESD SENSITIVE", fill="black", font=_font(22))
    ctx.doc("support/shipping_label_rot90.jpg", "jpg").image = lab.rotate(90, expand=True)
    for tok in (lot, po):
        ctx.forbid.append(tok)
    ctx.q("T24", "What lot code is printed on the shipping label?", lot, ["support/shipping_label_rot90.jpg"],
          "jpg, label rotated 90 degrees", note="The picture is rotated by 90 degrees; the lot code appears in no text file.", imgonly=True)
    # ---- pinout with two near-identical signals
    L = r.choice("EFGH")
    n0 = r.randint(1, 20)
    pins = [f"{L}{n0 + i}" for i in range(8)]
    sigs = r.sample(["GND", "SDA", "SCL", "PRSNT#", "CLK_P", "CLK_N", "3V3_AUX", "RESET#", "TX_EN", "RX_LOS", "WAKE#"], 6)
    a, b = r.sample(range(8), 2)
    sig_all = []
    it = iter(sigs)
    for i in range(8):
        sig_all.append("THERM_ALERT#" if i == a else "THERM_ALERT2#" if i == b else next(it))
    im = Image.new("RGB", (1100, 600), "white")
    dr = ImageDraw.Draw(im)
    dr.text((40, 24), f"{p3} service connector - pin assignment", fill="black", font=_font(30))
    dr.rectangle((30, 80, 1070, 330), outline="black", width=3)
    for i, (pin, sg) in enumerate(zip(pins, sig_all)):
        x = 50 + 128 * i
        dr.rectangle((x, 110, x + 108, 180), outline="black", width=2)
        f = _font(28)
        dr.text((x + 54 - dr.textlength(pin, font=f) / 2, 127), pin, fill="black", font=f)
        dr.line((x + 54, 180, x + 54, 202), fill="black", width=2)
        f = _font(15)
        dr.text((x + 54 - dr.textlength(sg, font=f) / 2, 210 + (i % 2) * 24), sg, fill="black", font=f)
    dr.text((40, 370), "THERM_ALERT# warns at the threshold; THERM_ALERT2# is the critical (shutdown) alert.", fill="black", font=_font(20))
    dr.text((40, 410), "Signal levels are 3.3 V unless noted.", fill="black", font=_font(20))
    pp = f"specs/{slug(p3)}_service_pinout.png"
    ctx.doc(pp, "png").image = im
    ctx.forbid += [pins[a], pins[b]]
    ctx.q("T24", f"Which pin of the {p3} service connector carries THERM_ALERT2#?", pins[b], [pp],
          "png, near-duplicate signal names (THERM_ALERT# vs THERM_ALERT2#)", aliases=[f"pin {pins[b]}"],
          note=f"THERM_ALERT# is on {pins[a]} (the near-miss); the pins appear in no text file.", imgonly=True)
    ctx.q("T24", f"Which pin of the {p3} service connector carries THERM_ALERT#?", pins[a], [pp],
          "png, near-duplicate signal names (THERM_ALERT# vs THERM_ALERT2#)", aliases=[f"pin {pins[a]}"],
          note=f"THERM_ALERT2# is on {pins[b]} (the near-miss); the pins appear in no text file.", imgonly=True)
    # ---- low-contrast sticker
    cal = f"CAL-{r.randint(10000, 99999)}-{r.choice('KMNPRT')}"
    st = Image.new("RGB", (760, 300), (236, 236, 236))
    sd = ImageDraw.Draw(st)
    sd.rectangle((8, 8, 751, 291), outline=(200, 200, 200), width=3)
    sd.text((30, 30), "CALIBRATION", fill=(196, 196, 196), font=_font(32))
    sd.text((30, 100), f"CAL ID {cal}", fill=(196, 196, 196), font=_font(40))
    sd.text((30, 180), "VALID FOR 12 MONTHS FROM ISSUE", fill=(200, 200, 200), font=_font(24))
    ctx.doc("support/calibration_sticker.jpg", "jpg").image = st
    ctx.forbid.append(cal)
    ctx.q("T24", "What calibration id is printed on the calibration sticker?", cal, ["support/calibration_sticker.jpg"],
          "jpg, low-contrast label", note="Light grey text on a slightly lighter grey background.", imgonly=True)
    # ---- photo-like JPEG with noise
    plate = f"AP-{r.randint(1000, 9999)}-{r.choice('WXYZ')}"
    W, H = 1200, 800
    base = np.zeros((H, W, 3), dtype=np.float32)
    grad = np.linspace(70, 120, H, dtype=np.float32)[:, None]
    base += grad[:, :, None]
    for x in range(0, W, 150):
        base[:, x:x + 14, :] *= 0.45
    for y in range(40, H, 70):
        base[y:y + 6, :, :] += 25
    g = np.random.default_rng(ctx.seed * 1000 + 7)
    for _ in range(60):
        x, y = int(g.integers(20, W - 20)), int(g.integers(20, H - 20))
        base[y:y + 6, x:x + 6, :] = g.choice([[40, 220, 60], [230, 70, 40], [60, 120, 240]])
    base += g.normal(0, 14, (H, W, 1)).astype(np.float32)
    ph = Image.fromarray(np.clip(base, 0, 255).astype("uint8"), "RGB")
    pl = Image.new("RGBA", (560, 130), (205, 205, 198, 255))
    pd = ImageDraw.Draw(pl)
    pd.rectangle((0, 0, 559, 129), outline=(60, 60, 60), width=4)
    pd.text((20, 14), "ASSET PLATE", fill=(30, 30, 30), font=_font(30))
    pd.text((20, 62), plate, fill=(20, 20, 20), font=_font(44))
    pl = pl.rotate(-4, expand=True, resample=Image.BICUBIC)
    ph.paste(pl, (330, 300), pl)
    ph = ph.filter(ImageFilter.GaussianBlur(0.9))
    arr = np.asarray(ph).astype(np.float32) + g.normal(0, 9, (H, W, 1)).astype(np.float32)
    ctx.doc("support/rack_photo.jpg", "jpg").image = Image.fromarray(np.clip(arr, 0, 255).astype("uint8"), "RGB")
    ctx.forbid.append(plate)
    ctx.q("T24", "What asset plate id is visible in the rack photo?", plate, ["support/rack_photo.jpg"],
          "jpg, photo-like image with noise, blur and a tilted plate", note="Noisy and slightly blurred; the id appears in no text file.",
          imgonly=True)
    # ---- docx whose only answer is inside an embedded PNG table
    rails = ["12V main", "3.3V aux", "HBM supply", "Core supply", "Fan header"]
    watts = r.sample(range(15, 780, 5), 5)
    ti = Image.new("RGB", (640, 320), "white")
    td = ImageDraw.Draw(ti)
    for i, (a_, b_) in enumerate([("Rail", "Budget (W)")] + list(zip(rails, map(str, watts)))):
        y = 10 + 50 * i
        td.rectangle((10, y, 630, y + 48), outline="black", width=2)
        td.line((380, y, 380, y + 48), fill="black", width=2)
        td.text((22, y + 9), a_, fill="black", font=_font(26))
        td.text((392, y + 9), b_, fill="black", font=_font(26))
    pth = f"planning/{slug(p4)}_power_budget.docx"
    ctx.doc(pth, "docx").blocks = [("h", f"{p4} power budget", 1),
                                    ("p", "The power budget per rail is given in Table 1."), ("img", ti),
                                    ("p", "Table 1 - rail budgets (as presented). Margins are tracked separately.")]
    ctx.q("T24", f"What is the power budget, in watts, of the {rails[2]} rail in the {p4} power budget?", str(watts[2]), [pth],
          "docx, answer only inside an embedded PNG table", aliases=[f"{watts[2]} W"],
          note=f"The table is a picture; the near-miss rows are {watts[1]} ({rails[1]}) and {watts[3]} ({rails[3]}).", imgonly=True)
    ctx.cache["T24"] = {}


# ------------------------------------------------------------------ v2 traps (opt-in: --traps T25,...)

def _events(ctx: Ctx) -> list:
    """EVENTS in a per-seed order shared by T26/T27: T26 takes [0:3] (+[7] as its canary), T27 takes [3:7]."""
    ev = EVENTS[:]
    ctx.xr("EVT").shuffle(ev)
    return ev


def _log_lines(r: random.Random, node: str, svc: str, date: dt.date, specials: list, n: int = 24) -> list:
    """A service log for `date`: n routine INFO/WARN lines with the (level, message) `specials` spread among them."""
    t = dt.datetime.combine(date, dt.time(r.randint(0, 20), r.randint(0, 59), r.randint(0, 59)))
    slots = sorted(r.sample(range(3, n), len(specials)))
    out, k = [], 0
    for i in range(n):
        t += dt.timedelta(seconds=r.randint(5, 240))
        if k < len(slots) and i == slots[k]:
            level, msg = specials[k]
            k += 1
        else:
            level, msg = [("INFO", f"scheduler: batch {r.randint(10000, 99999)} accepted, {r.randint(8, 128)} requests"),
                          ("INFO", f"thermal: die {r.randint(0, 3)} at {r.randint(60, 78)}C, within envelope"),
                          ("INFO", f"heartbeat ok, uptime {r.randint(2, 400)}h"),
                          ("INFO", f"telemetry: exported {r.randint(100, 9000)} samples"),
                          ("WARN", f"thermal: die {r.randint(0, 3)} at {r.randint(82, 90)}C, approaching limit")][r.randrange(5)]
        out.append(f"{t.strftime('%Y-%m-%dT%H:%M:%SZ')} {node} {svc}[{r.randint(1000, 4999)}]: {level:<5} {msg}")
    return out


def trap_T25(ctx: Ctx):
    """xlsx cross-sheet join: Parts (product, component, part_number) + LeadTimes / Costs keyed ONLY by part_number."""
    u, r = ctx.u, ctx.xr("T25")
    lo, mid, hi = sorted(("p1", "p0", "p2"), key=lambda k: int(u.p[k].rsplit("-", 1)[1]))
    prods = [u.p[lo], u.p[mid], u.p[hi]]
    comps = ["Retimer board", "Midplane connector", "Battery backup unit", "Rail kit", "Expansion riser", "Front bezel"]
    r.shuffle(comps)
    only = comps.pop()                                # exists for ONE product only (canary: ask it for another product)
    only_prod = prods[r.randrange(3)]
    taken, lead_used, cost_used = set(), set(), set()

    def pn():
        while True:
            x = xpart(r, u, "BOM")
            if x not in taken:
                taken.add(x)
                return x

    def uniq(draw, used):
        while True:
            v = draw()
            if v not in used:
                used.add(v)
                return v
    rows, lt, costs, info = [], [], [], {}
    for comp in comps + [only]:
        for prod in prods:
            if comp == only and prod != only_prod:
                continue
            part, weeks = pn(), uniq(lambda: r.randint(3, 26), lead_used)
            rows.append((prod, comp, part, f"{comp}, {prod} series", r.choice(["active", "active", "preferred"])))
            lt.append((part, weeks, r.choice(SUPPLIERS)))
            brk = {b: uniq(lambda: r.randrange(20 * 4, 400 * 4) / 4, cost_used) for b in (100, 1000, 5000)}
            costs += [(part, b, c) for b, c in brk.items()]
            info[(prod, comp)] = {"part": part, "weeks": weeks, "cost": brk}
    for lst in (rows, lt, costs):
        r.shuffle(lst)
    path = f"procurement/{slug(u.pre)}_bom_costs.xlsx"
    ctx.doc(path, "xlsx").sheets = [
        {"name": "Parts", "title": f"{u.company} - {u.pre} series bill of parts (FY{u.fy})",
         "header": ["product", "component", "part_number", "description", "status"], "rows": rows},
        {"name": "LeadTimes", "header": ["part_number", "lead_time_weeks", "supplier"], "rows": lt},
        {"name": "Costs", "header": ["part_number", "qty_break", "unit_cost_usd"], "rows": costs,
         "money": ["unit_cost_usd"], "shown": [(a, b, f"{c:.2f}") for a, b, c in costs]},
        {"name": "Notes", "header": ["note"],
         "rows": [("Lead times and costs are keyed by part_number; product and component are on the Parts sheet.",)]}]
    pr = prods[:]
    r.shuffle(pr)
    asked = [(pr[0], comps[0]), (pr[1], comps[0]), (pr[2], comps[1])]       # 1st/2nd: same component, sibling products
    for prod, comp in asked:
        e = info[(prod, comp)]
        sibs = ", ".join(f"{info[(q, comp)]['weeks']} ({q})" for q in prods if q != prod)
        ctx.q("T25", f"What is the lead time, in weeks, of the {comp.lower()} for the {prod}?", str(e["weeks"]), [path],
              "xlsx, cross-sheet join (product on Parts, lead time on LeadTimes via part_number)",
              aliases=[f"{e['weeks']} weeks"],
              note=f"Parts maps {prod} + {comp} to {e['part']}; LeadTimes holds {e['weeks']} for that part number and never "
                   f"names the product. The sibling products' rows are the near-misses: {sibs}.")
    cprod, ccomp = prods[1], comps[2]
    e = info[(cprod, ccomp)]
    s, al = money(e["cost"][1000])
    ctx.q("T25", f"What is the unit cost in USD of the {ccomp.lower()} for the {cprod} at 1000 unit volume?", s, [path],
          "xlsx, cross-sheet join + quantity break (Costs has one row per part and break)", aliases=al,
          note=f"{e['part']}: {e['cost'][100]:.2f} at 100, {e['cost'][1000]:.2f} at 1000, {e['cost'][5000]:.2f} at 5000.")
    ctx.cache["T25"] = {"path": path, "only": only, "other_prod": [q for q in prods if q != only_prod][0],
                        "only_prod": only_prod, "cprod": cprod, "ccomp": ccomp}


def trap_T26(ctx: Ctx):
    """Chain with a topical decoy: log -> ticket id -> csv fixed_in, while the csv also holds a DECOY ticket whose
    summary uses the question's own words (different ticket, different fixed_in)."""
    u, r = ctx.u, ctx.xr("T26")
    ev = _events(ctx)
    major, minor = u.ver_base.split(".")
    taken_v, taken_t, taken_e = set(), set(), {u.e6, u.e7, u.e_dec, u.nx_err}

    def ver():
        while True:
            v = f"{major}.{minor}.{r.randint(10, 39)}"          # patch >= 10: never in the release notes / bug database
            if v not in taken_v:
                taken_v.add(v)
                return v

    def tick():
        while True:
            t = f"{u.tix}-{r.randint(4000, 4999)}"
            if t not in taken_t:
                taken_t.add(t)
                return t

    def code():
        while True:
            c = f"E{r.randint(1000, 9999)}"
            if c not in taken_e:
                taken_e.add(c)
                return c
    decoy_tpl = ["{N} counter is not exposed through the status subcommand",
                 "Runbook section on {n} recovery is outdated",
                 "{N} events are missing from the exported telemetry",
                 "{N} threshold wording in the integration guide is ambiguous"]
    cap = lambda s_: s_[0].upper() + s_[1:]
    rows, chains = [], []
    base = dt.date(2026, 9, 1) + dt.timedelta(days=r.randint(0, 6))
    for k in range(3):
        text, noun, summ, comp = ev[k]
        real, decoy, v_real, v_dec = tick(), tick(), ver(), ver()
        date = base + dt.timedelta(days=4 * k + r.randint(0, 3))
        ecode = code()
        lp = f"logs/prod_serving_{date.isoformat()}.log"
        ctx.text(lp).lines = _log_lines(r, f"serving-node-{r.randint(1, 12):02d}", u.fwslug, date, [
            ("WARN", f"supervisor: {text.split(',')[0]}, retrying"), ("ERROR", f"{ecode}: {text}"),
            ("INFO", f"incident logged against {real}")])
        rows.append((real, comp, "S2", summ, "closed", v_real))
        rows.append((decoy, comp, "S3", r.choice(decoy_tpl).format(N=cap(noun), n=noun), "closed", v_dec))
        chains.append({"log": lp, "noun": noun, "real": real, "decoy": decoy, "v_real": v_real, "v_dec": v_dec})
    text, noun, summ, comp = ev[7]                      # decoy-only event: a ticket exists, no log shows the incident
    d_only, v_only = tick(), ver()
    rows.append((d_only, comp, "S3", r.choice(decoy_tpl).format(N=cap(noun), n=noun), "closed", v_only))
    bugs = GENERIC_BUGS[:]
    r.shuffle(bugs)
    for comp_g, summ_g in bugs[:6]:
        rows.append((tick(), comp_g, r.choice(["S3", "S4"]), summ_g.format(v=u.ver_base), "closed",
                     f"{int(major) - 1}.{r.randint(0, 5)}.{r.randint(0, 4)}"))
    r.shuffle(rows)
    cp = "support/incident_tickets.csv"
    ctx.text(cp).lines = ["ticket,component,severity,summary,status,fixed_in"] + [",".join(x) for x in rows]
    for c in chains:
        ctx.q("T26", f"The production log shows a {c['noun']} incident. Which firmware release fixed the underlying defect?",
              c["v_real"], [c["log"], cp], "multi-hop with topical decoy ticket, two citations required",
              aliases=[f"{u.fw} {c['v_real']}"],
              note=f"Chain: log -> {c['real']} -> incident_tickets.csv ({c['v_real']}). The csv also holds {c['decoy']}, whose "
                   f"summary repeats the words '{c['noun']}' but is a different ticket (fixed in {c['v_dec']}); answering from it "
                   "leaves the log out of the citations.")
    ctx.q("T26", f"Which firmware version fixed ticket {chains[0]['decoy']}?", chains[0]["v_dec"], [cp],
          "csv, single-hop (decoy ticket asked by id)", aliases=[f"{u.fw} {chains[0]['v_dec']}"],
          note="Control: the decoy ticket is a real ticket and has its own fixed_in.")
    ctx.cache["T26"] = {"noun_absent": noun, "csv": cp, "ticket_absent": d_only, "v_absent": v_only}


def trap_T27(ctx: Ctx):
    """Several incidents across several logs (different dates and services); one event recurs on two dates with
    different codes, so the question must carry the date. Each log also has noise ERROR lines with other codes."""
    u, r = ctx.u, ctx.xr("T27")
    ev = _events(ctx)
    a, b, c, never = ev[3], ev[4], ev[5], ev[6]
    taken = {u.e6, u.e7, u.e_dec, u.nx_err}

    def code():
        while True:
            x = f"E{r.randint(1000, 9999)}"
            if x not in taken:
                taken.add(x)
                return x
    generic = ["checksum mismatch on firmware slot 2", "NVMe write latency above 50 ms for 10 s",
               "license server unreachable, grace period started", "TLS handshake failed with the fleet controller"]
    dates = [dt.date(2026, 9, 8) + dt.timedelta(days=r.randint(0, 5))]
    for _ in range(3):
        dates.append(dates[-1] + dt.timedelta(days=r.randint(2, 5)))
    svcs = ["chassisd", "powermgr", "fabricd", "railmon"]
    got = []
    for event, date, svc in zip((a, a, b, c), dates, svcs):
        ecode = code()
        specials = [("ERROR", f"{code()}: {g}") for g in r.sample(generic, 2)] + [("ERROR", f"{ecode}: {event[0]}")]
        r.shuffle(specials)
        path = f"logs/{svc}_{date.isoformat()}.log"
        ctx.text(path).lines = _log_lines(r, f"{svc}-{r.randint(1, 9):02d}", svc, date, specials, n=20)
        got.append((event, date, path, ecode))
    for i, (event, date, path, ecode) in enumerate(got):
        dated = i != 2                                   # the b event occurs once: no date needed
        query = f"What error code is logged when the {event[1]} occurs" + (f" on {date.isoformat()}?" if dated else "?")
        twin = dates[1 - i] if i < 2 else None
        ctx.q("T27", query, ecode, [path],
              "text log, same event in two service logs (date picks the file)" if i < 2 else "text log, one incident per service log",
              note=f"One ERROR line in {path}; noise ERROR lines carry other codes." +
                   (f" The same event is logged on {twin.isoformat()} with a different code." if twin else ""))
    ctx.cache["T27"] = {"never": never[1], "b": b[1], "first": dates[0].isoformat(),
                        "no_log": (dates[3] + dt.timedelta(days=3)).isoformat(), "a": a[1]}


def trap_T28(ctx: Ctx):
    """Datasheet PDF with the product family side by side (Parameter | lo | mid | hi) in padded fixed-pitch lines;
    r1 (WITHDRAWN) carries different numbers in every numeric row. One cell is '-' (not applicable)."""
    u, r = ctx.u, ctx.xr("T28")
    lo, mid, hi = sorted(("p1", "p0", "p2"), key=lambda k: int(u.p[k].rsplit("-", 1)[1]))
    prods = [u.p[lo], u.p[mid], u.p[hi]]
    asc = lambda a, b, step=1: sorted(r.sample(range(a, b, step), 3))
    num = [("Maximum junction temperature", asc(80, 112), lambda x: f"{x} C"),
           ("Board power (TBP)", asc(250, 800, 5), lambda x: f"{x} W"),
           ("Memory capacity", asc(32, 257, 8), lambda x: f"{x} GiB"),
           ("Peak memory bandwidth", asc(15, 85), lambda x: f"{x / 10:.1f} TB/s"),
           ("Typical acoustic noise", asc(30, 56), lambda x: f"{x} dB(A)")]
    inlet = asc(30, 46)[:2]
    aux = r.sample(["12V-2x6", "16-pin 12VHPWR", "8-pin EPS", "6-pin PCIe"], 3)

    def cells_of(rev):
        out = []
        for name, vals, fmt in num:
            if rev == 1:
                vals = [x + r.choice([-9, -6, 6, 9]) if name != "Peak memory bandwidth" else x + r.choice([-6, 6]) for x in vals]
            out.append((name, [fmt(x) for x in vals]))
        out.append(("Maximum coolant inlet", ["-"] + [f"{x} C" for x in inlet]))
        if rev == 2:
            out.append(("Auxiliary power connector", aux))
        return out

    rule = r.random() < 0.5         # half of the seeds print a '-----' rule between the header and the rows

    def table(rows):
        head = f"{'Parameter':32}" + "".join(f"{p:<16}" for p in prods)
        return [head.rstrip()] + (["-" * 80] if rule else []) + \
            [(f"{n:32}" + "".join(f"{c:<16}" for c in cs)).rstrip() for n, cs in rows]
    rows1, rows2 = cells_of(1), cells_of(2)
    s = u.pre.lower()
    r2p, r1p = f"specs/{s}_family_datasheet_r2.pdf", f"specs/{s}_family_datasheet_r1_WITHDRAWN.pdf"
    title = f"{u.company} {u.pre}-series accelerators - family datasheet"
    ctx.doc(r1p, "pdf").blocks = [("title", title), ("p", "Datasheet, revision 1 - WITHDRAWN, SUPERSEDED BY REVISION 2"),
                                  ("p", "Preliminary silicon. Do not design to these numbers."), ("h", "Specifications")] + \
        [("mono", ln) for ln in table(rows1)]
    ctx.doc(r2p, "pdf").blocks = [("title", title), ("p", "Datasheet, revision 2 - supersedes revision 1"),
                                  ("h", "Specifications")] + [("mono", ln) for ln in table(rows2)] + \
        [("h", "Revision history"),
         ("p", "r2  Re-characterised on production silicon; connector row added. Revision 1 is withdrawn."),
         ("p", "r1  Initial release, preliminary silicon.")]
    qnames = {"Maximum junction temperature": "maximum junction temperature", "Board power (TBP)": "board power",
              "Memory capacity": "memory capacity, in GiB", "Peak memory bandwidth": "peak memory bandwidth, in TB/s",
              "Auxiliary power connector": "auxiliary power connector"}
    for ri, ci in [(0, 2), (1, 0), (2, 1), (3, 2), (6, 1)]:       # spread over rows and columns
        name, cells = rows2[ri]
        val = cells[ci]
        number, _, unit = val.partition(" ")
        if name == "Auxiliary power connector":
            ans, al = val, []
        else:
            ans, al = number, ([f"{number} C", f"{number}°C", f"{number} degrees"] if unit == "C" else [val])
        what, _, in_unit = qnames[name].partition(", ")
        ctx.q("T28", f"What is the {what} of the {prods[ci]}" + (f", {in_unit}?" if in_unit else "?"), ans, [r2p],
              "pdf, family table with the products side by side (pick the right column)", aliases=al,
              note=f"Row '{name}': " + " | ".join(f"{p} = {c}" for p, c in zip(prods, cells))
                   + (f". r1 (withdrawn) says {rows1[ri][1][ci]}." if ri < 5 else "."))
    ctx.cache["T28"] = {"path": r2p, "lo": prods[0], "mid": prods[1], "hi": prods[2]}


# ------------------------------------------------------------------ extra images (--n-images)

_SIGNAL_POOL = ["GND", "SDA", "SCL", "PRSNT#", "CLK_P", "CLK_N", "3V3_AUX", "RESET#", "TX_EN", "RX_LOS", "WAKE#", "PERST#"]
_COMPONENTS = ["power shelf", "midplane", "fan tray", "backplane", "management card", "switch module", "pump module"]


def _scene(w: int, h: int, seed: int, base: int = 118, sigma: float = 11.0):
    """Photo-like background: seeded gaussian noise (the `_bg_noise` layer) tinted to a bench/rack surface."""
    from PIL import Image
    n = _bg_noise(w, h, seed, sigma)
    return Image.merge("RGB", (n.point(lambda v: min(255, int(v * base / 128) + 8)),
                               n.point(lambda v: min(255, int(v * base / 128))),
                               n.point(lambda v: min(255, int(v * (base - 6) / 128)))))


def _card(w: int, h: int, lines: list, title: str, r: random.Random, size: int = 30):
    """White card with a border, a title and (text, font size) lines."""
    from PIL import Image, ImageDraw
    c = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(c)
    d.rectangle((0, 0, w - 1, h - 1), outline="black", width=4)
    d.text((26, 20), title, fill="black", font=_font(size + 2))
    y = 20 + size + 34
    for text, sz in lines:
        d.text((26, y), text, fill="black", font=_font(sz))
        y += sz + 22
    return c


def _on_scene(card, ctx_seed: int, angle: float, r: random.Random, blur: float = 0.0, noise: float = 0.0):
    from PIL import Image, ImageFilter
    import numpy as np
    cr = card.convert("RGBA").rotate(angle, expand=True, resample=Image.BICUBIC)
    W, H = cr.width + 160, cr.height + 140
    bg = _scene(W, H, ctx_seed)
    bg.paste(cr, ((W - cr.width) // 2, (H - cr.height) // 2), cr)
    if blur:
        bg = bg.filter(ImageFilter.GaussianBlur(blur))
    if noise:
        g = np.random.default_rng(ctx_seed + 1)
        arr = np.asarray(bg).astype(np.float32) + g.normal(0, noise, (H, W, 1)).astype(np.float32)
        bg = Image.fromarray(np.clip(arr, 0, 255).astype("uint8"), "RGB")
    return bg


def _pinout_img(r, title: str, pins: list, sigs: list, note: str):
    from PIL import Image, ImageDraw
    n = len(pins)
    im = Image.new("RGB", (60 + 128 * n, 470), "white")
    dr = ImageDraw.Draw(im)
    dr.text((40, 24), title, fill="black", font=_font(30))
    dr.rectangle((30, 80, im.width - 30, 330), outline="black", width=3)
    for i, (pin, sg) in enumerate(zip(pins, sigs)):
        x = 50 + 128 * i
        dr.rectangle((x, 110, x + 108, 180), outline="black", width=2)
        f = _font(28)
        dr.text((x + 54 - dr.textlength(pin, font=f) / 2, 127), pin, fill="black", font=f)
        dr.line((x + 54, 180, x + 54, 202), fill="black", width=2)
        f = _font(15)
        dr.text((x + 54 - dr.textlength(sg, font=f) / 2, 210 + (i % 2) * 24), sg, fill="black", font=f)
    dr.text((40, 380), note, fill="black", font=_font(20))
    return im


def images_extra(ctx: Ctx, n: int):
    """--n-images N: N png/jpg files that carry printed text only. The first 9 are 'answer' images (10 questions:
    two refusals ask for a field the picture does not have); the rest are look-alikes of the same kinds with other
    ids, sites, products and components, so nothing is retrievable by file name or by kind alone."""
    u, r = ctx.u, ctx.xr("IMG")
    prods = list(u.p.values())
    used = set()

    def uniq(draw):
        while True:
            x = draw()
            if x not in used:
                used.add(x)
                return x
    dirs = ["support/photos", "engineering/lab_photos", "planning/scans", "marketing/assets", "facilities/photos"]
    counter = [0]

    def put(img, ext, descriptive=None):
        counter[0] += 1
        name = descriptive or f"IMG_{uniq(lambda: r.randint(1000, 9999))}"
        path = f"{dirs[counter[0] % len(dirs)]}/{name}.{ext}"
        ctx.doc(path, ext).image = img
        return path

    def tok(prefix, lo=1000, hi=9999, suffix=""):
        return uniq(lambda: f"{prefix}{r.randint(lo, hi)}{suffix}")
    sites = SITES[:]
    r.shuffle(sites)
    comps = _COMPONENTS[:]
    r.shuffle(comps)
    p_print, p_plate = prods[2], prods[3]
    # 1. shipping label (lot + PO)
    site, lot, po = sites[0], f"L{r.randint(10, 99)}-{uniq(lambda: r.randint(1000, 9999))}", f"PO-{uniq(lambda: r.randint(10000, 99999))}"
    card = _card(760, 400, [(f"SHIP TO   {site.upper()} LAB", 30), (f"LOT {lot}", 40), (f"REF {po}", 30),
                            ("HANDLE WITH CARE - ESD SENSITIVE", 22)], "SHIPPING LABEL", r)
    pth = put(_on_scene(card, ctx.seed + 11, r.choice([-3, 2, 90]), r, blur=0.5, noise=5), "jpg")
    ctx.forbid += [lot, po]
    ctx.q("IMG", f"What lot code is printed on the shipping label addressed to the {site} lab?", lot, [pth],
          "jpg, shipping label among many look-alike photos", note="Lot appears in no text file.", imgonly=True)
    ctx.q("IMG", f"Which purchase order reference is printed on the shipping label addressed to the {site} lab?", po, [pth],
          "jpg, shipping label among many look-alike photos", note="Same picture as the lot question.", imgonly=True)
    # 2. rack asset plate (photo-like)
    rack = f"R{uniq(lambda: r.randint(1, 60))}"
    plate = f"AP-{r.randint(1000, 9999)}-{r.choice('WXYZ')}"
    card = _card(560, 230, [(f"RACK {rack}", 30), (plate, 44)], "ASSET PLATE", r, size=28)
    pth = put(_on_scene(card, ctx.seed + 12, r.choice([-4, 3]), r, blur=0.9, noise=9), "jpg")
    ctx.forbid.append(plate)
    ctx.q("IMG", f"What asset plate id is on rack {rack} in the rack photos?", plate, [pth],
          "jpg, noisy rack photo with a tilted plate", note="One of several rack-plate photos; the rack number selects it.",
          imgonly=True)
    # 3.+4. pinouts with near-duplicate signals
    for k, (a_sig, b_sig, hdr) in enumerate([("FAN_TACH1", "FAN_TACH2", "management header"), ("PWR_OK_A", "PWR_OK_B", "power header")]):
        jid = uniq(lambda: f"J{r.randint(10, 40)}")
        L = r.choice("ABCDEFGH")
        n0 = r.randint(1, 20)
        pins = [f"{L}{n0 + i}" for i in range(8)]
        others = r.sample([s for s in _SIGNAL_POOL], 6)
        ia, ib = r.sample(range(8), 2)
        it = iter(others)
        sigs = [a_sig if i == ia else b_sig if i == ib else next(it) for i in range(8)]
        im = _pinout_img(r, f"{prods[1 + k]} {jid} {hdr} - pin assignment", pins, sigs,
                         f"{a_sig} and {b_sig} are separate signals. Levels are 3.3 V unless noted.")
        pth = put(im, "png")
        ctx.forbid += [pins[ia], pins[ib]]
        ctx.q("IMG", f"Which pin of the {jid} {hdr} of the {prods[1 + k]} carries {b_sig}?", pins[ib], [pth],
              "png, pinout with near-duplicate signal names", aliases=[f"pin {pins[ib]}"],
              note=f"{a_sig} is on {pins[ia]} (the near-miss).", imgonly=True)
        if k == 0:
            absent = uniq(lambda: f"SPI_{r.choice('ABCD')}_CS#")
            ctx.cache["IMG_pin"] = {"jid": jid, "hdr": hdr, "prod": prods[1], "absent": absent}
    # 5. component nameplate (serial)
    comp = comps[0]
    ser = f"{u.code3[0]}{r.choice('QRSTWXZ')}{r.randint(1, 9)}-{r.randint(100000, 999999)}-{r.randint(10, 99)}"
    card = _card(700, 330, [(f"MODEL   {p_plate}-{comp.split()[0][:2].upper()}", 28), (f"SERIAL   {ser}", 34),
                            (f"MADE IN {u.made_in}", 24)], f"{comp.upper()} NAMEPLATE", r)
    pth = put(_on_scene(card, ctx.seed + 13, -2, r, blur=0.4, noise=4), "jpg")
    ctx.forbid.append(ser)
    ctx.q("IMG", f"What serial number is printed on the nameplate of the {p_plate} {comp}?", ser, [pth],
          "jpg, component nameplate", note="Look-alike nameplates exist for other components and products.", imgonly=True)
    ctx.cache["IMG_name"] = {"prod": p_plate, "comp": comp}
    # 6. BMC settings screenshot (numeric value)
    floor = r.randrange(25, 66, 5)
    img = _card(900, 420, [(f"Fan floor (percent)          {floor}", 30), (f"Power cap (W)                {r.randrange(300, 700, 25)}", 30),
                           ("Boot order                   NVMe, PXE", 30), ("Telemetry export             enabled", 30)],
                f"{prods[4]} BMC - thermal settings", r)
    pth = put(img, "png")
    ctx.q("IMG", f"What fan floor percentage is shown in the thermal settings screenshot of the {prods[4]} BMC?", str(floor), [pth],
          "png, settings screenshot (numeric value only in the picture)", aliases=[f"{floor} percent", f"{floor}%"],
          note="Short number: only the cited file's own text must not carry it.", imgonly=True)
    # 7. test printout (measured value)
    il = f"{r.randint(5, 39) / 10:.1f}"
    img = _card(860, 380, [(f"Loopback test  {p_print}", 28), (f"Insertion loss (dB)      {il}", 32), ("Return loss (dB)         18.4", 28),
                           ("Result                  PASS", 28)], "TEST PRINTOUT", r)
    pth = put(_on_scene(img, ctx.seed + 14, 1.5, r, blur=0.3, noise=3), "jpg")
    ctx.q("IMG", f"What insertion loss, in dB, is printed on the loopback test printout of the {p_print}?", il, [pth],
          "jpg, test printout (numeric value only in the picture)", aliases=[f"{il} dB"], note="Return loss is 18.4 on every printout.",
          imgonly=True)
    # 8. spare-parts bin tag (location)
    part = xpart(r, u, "SPR")
    loc = f"{r.choice('ABCDEF')}-{r.randint(10, 29)}-{r.randint(1, 4)}"
    img = _card(700, 300, [(part, 40), (f"BIN {loc}", 40), (f"QTY {r.randint(2, 40)}", 26)], "SPARES BIN TAG", r)
    pth = put(_on_scene(img, ctx.seed + 15, 2, r, noise=4), "jpg")
    ctx.forbid += [part, loc]
    ctx.q("IMG", f"Which bin location is printed on the spares bin tag for part {part}?", loc, [pth],
          "jpg, bin tag (part number and location only in the picture)", note="Look-alike bin tags carry other parts.",
          imgonly=True)
    # 9. console boot screen (version)
    ver = f"{u.ver_base}.{r.randint(40, 79)}"
    site_b = sites[1]
    img = _card(900, 360, [(f"{u.fw} firmware {ver}", 34), (f"Node  {site_b.lower()}-ctl-{r.randint(1, 9):02d}", 28), ("Boot OK  all dies online", 28)],
                "CONSOLE", r)
    pth = put(_on_scene(img, ctx.seed + 16, -1, r, blur=0.4, noise=5), "jpg")
    ctx.forbid.append(ver)
    ctx.q("IMG", f"What firmware version does the boot screen photo of the {site_b} console show?", ver, [pth],
          "jpg, console photo (version only in the picture)", aliases=[f"{u.fw} {ver}"], imgonly=True,
          note="Look-alike boot screens show other sites.")
    # refusals: the asked field is not in the (existing) picture
    pc = ctx.cache["IMG_pin"]
    ctx.q("IMG", f"Which pin of the {pc['jid']} {pc['hdr']} of the {pc['prod']} carries {pc['absent']}?", "", [],
          "image refusal - the signal is not on the pinout", unanswerable=True,
          note="The pinout lists eight signals; this one is not among them.")
    ctx.q("IMG", f"What purchase order reference is printed on the nameplate of the {p_plate} {comp}?", "", [],
          "image refusal - the nameplate has no such field", unanswerable=True,
          note="The nameplate carries model, serial and origin only.")
    ctx.q("IMG", f"What firmware version does the boot screen photo of the {sites[2]} console show?", "", [],
          "image refusal - no boot screen photo for this site", unanswerable=True,
          note="Boot screen photos exist for other sites only.")
    # ---- look-alike fillers (no questions); never reuse an answered rack, site, jid, product-kind pair, part or component
    kinds = ["return", "plate", "pinout", "nameplate", "printout", "bin", "boot", "photo"]
    for i in range(max(0, n - counter[0])):
        k = kinds[i % len(kinds)]
        if k == "return":
            card = _card(700, 340, [(f"FROM   {r.choice(sites[3:]).upper()} LAB", 28), (f"RMA {tok('RMA-', 10000, 99999)}", 36),
                                    (f"CARTON {r.randint(1, 12)} OF {r.randint(12, 20)}", 26)], "RETURN TAG", r)
            put(_on_scene(card, ctx.seed + 100 + i, r.choice([-3, 2, 90]), r, blur=0.5, noise=5), "jpg")
        elif k == "plate":
            card = _card(560, 230, [(f"RACK {tok('R', 61, 99)}", 30), (f"AP-{r.randint(1000, 9999)}-{r.choice('ABCD')}", 44)],
                         "ASSET PLATE", r, size=28)
            put(_on_scene(card, ctx.seed + 100 + i, r.choice([-4, 3]), r, blur=0.9, noise=9), "jpg")
        elif k == "pinout":
            L = r.choice("ABCDEFGH")
            n0 = r.randint(1, 20)
            sg = r.sample(_SIGNAL_POOL, 8)
            put(_pinout_img(r, f"{r.choice(prods)} {tok('J', 41, 99)} debug header - pin assignment", [f"{L}{n0 + j}" for j in range(8)], sg,
                            "Levels are 1.8 V unless noted."), "png")
        elif k == "nameplate":
            comp_f = comps[1 + i % (len(comps) - 1)]
            pr = r.choice(prods)
            card = _card(700, 330, [(f"MODEL   {pr}-{comp_f.split()[0][:2].upper()}", 28),
                                    (f"SERIAL   {u.code3[0]}{r.choice('ABCDEFGH')}{r.randint(1, 9)}-{r.randint(100000, 999999)}-{r.randint(10, 99)}", 34),
                                    (f"MADE IN {r.choice(['MALAYSIA', 'TAIWAN', 'VIETNAM', 'MEXICO'])}", 24)],
                         f"{comp_f.upper()} NAMEPLATE", r)
            put(_on_scene(card, ctx.seed + 100 + i, -2, r, blur=0.4, noise=4), "jpg")
        elif k == "printout":
            pr = r.choice([x for x in prods if x != p_print])
            card = _card(860, 380, [(f"Loopback test  {pr}", 28), (f"Insertion loss (dB)      {r.randint(5, 39) / 10:.1f}", 32),
                                    ("Return loss (dB)         18.4", 28), ("Result                  PASS", 28)], "TEST PRINTOUT", r)
            put(_on_scene(card, ctx.seed + 100 + i, 1.5, r, blur=0.3, noise=3), "jpg")
        elif k == "bin":
            card = _card(700, 300, [(xpart(r, u, "SPR"), 40), (f"BIN {r.choice('ABCDEF')}-{r.randint(30, 49)}-{r.randint(1, 4)}", 40),
                                    (f"QTY {r.randint(2, 40)}", 26)], "SPARES BIN TAG", r)
            put(_on_scene(card, ctx.seed + 100 + i, 2, r, noise=4), "jpg")
        elif k == "boot":
            card = _card(900, 360, [(f"{u.fw} firmware {u.ver_base}.{r.randint(80, 99)}", 34),
                                    (f"Node  {r.choice(sites[3:]).lower()}-ctl-{r.randint(1, 9):02d}", 28), ("Boot OK  all dies online", 28)],
                         "CONSOLE", r)
            put(_on_scene(card, ctx.seed + 100 + i, -1, r, blur=0.4, noise=5), "jpg")
        else:                                   # a bare bench photo with a handwritten-style note
            card = _card(520, 200, [(f"{r.choice(['checked', 'moved', 'swapped'])} {r.choice(['sled', 'cable', 'psu'])} {r.randint(1, 9)}", 30)],
                         "BENCH NOTE", r, size=26)
            put(_on_scene(card, ctx.seed + 100 + i, r.choice([-6, 5]), r, blur=0.8, noise=8), "jpg")


# ------------------------------------------------------------------ v1 canaries

def canaries_x(ctx: Ctx):
    """Refusals for the v1 traps. Tag = source trap id + 'c'. Each lambda runs only when its traps were generated."""
    u, c = ctx.u, ctx.cache
    p0, p2 = u.p["p0"], u.p["p2"]
    enc = f"vendor/{slug(u.supplier)}_agreement_ENCRYPTED.pdf"
    items = [
        (("T14b",), "T14bc", lambda: (f"What is the total list price in USD of 3 {p2} developer kits?",
         "refusal - needs arithmetic the corpus does not contain",
         f"One kit is ${c['p2spec'].price:,}.00; no document states a three-kit total.")),
        (("T13",), "T13c", lambda: (f"What is the total cost in USD of 3 {c['T13']['tgt']['name']} units at the unit cost in the {p2} rack bill of materials?",
         "refusal - needs arithmetic (3 x unit cost) that is not in the workbook",
         "The BOM line total is for a different quantity; no cell holds 3 x unit cost.")),
        (("T13",), "T13c", lambda: (f"What is the unit cost in USD of the {c['T13']['legacy']['name']} in the {p2} rack bill of materials?",
         "refusal - the value exists only in the HIDDEN (stale) sheet",
         f"Prior_Estimate (hidden) lists {c['T13']['legacy']['unit']:.2f}; the visible BOM no longer has this item.")),
        (("T13",), "T13c", lambda: (f"How many {c['T13']['legacy']['name']} units does the {p2} rack bill of materials list?",
         "refusal - the value exists only in the HIDDEN (stale) sheet",
         f"Prior_Estimate (hidden) lists {c['T13']['legacy']['qty']}; the visible BOM no longer has this item.")),
        (("T5",), "T19c", lambda: (f"What is the total cost in USD of 5 backplane cables for the {p0}?",
         "refusal - needs arithmetic the corpus does not contain",
         f"The parts list gives a unit cost ({u.cost_cab_p0:.2f}) only.")),
        (("T19b",), "T19bc", lambda: (f"What is the coolant flow of the {p0}, in L/min?",
         "refusal - the source gives a range, not a value",
         f"The limits sheet says 'between {c['T19b']['fa']} and {c['T19b']['fb']} L/min depending on the fan profile'.")),
        (("T19b",), "T19bc", lambda: (f"What is the fan speed of the {p0}, in RPM?",
         "refusal - the source gives a range, not a value",
         f"The limits sheet says 'between {c['T19b']['ra']} and {c['T19b']['rb']} RPM'.")),
        ((), "T16", lambda: ("According to vendor/internal_audit.txt, what is the late-delivery penalty, in percent?",
         "refusal - the file named in the question is unreadable (chmod 000)",
         f"{u.penalty} percent appears only in that file, which cannot be read.")),
        ((), "T16", lambda: (f"According to {enc}, what is the unit price of the {p0} at {comma(u.vol)} unit volume?",
         "refusal - the file named in the question is ENCRYPTED", f"{u.price} USD appears only inside the encrypted PDF.")),
        (("T18b",), "T18bc", lambda: (f"In which quarter does the {c['T18b']['draft_prod']} enter customer sampling?",
         "refusal - the product exists only in the DRAFT roadmap",
         f"{c['T18b']['draft_prod']} appears only in roadmap_fy{u.fy}_DRAFT.docx ({c['T18b']['dq'][4]}); the approved roadmap has no such row.")),
        (("T18b",), "T18bc", lambda: (f"In which quarter does the {c['T18b']['draft_prod']} enter volume production?",
         "refusal - the product exists only in the DRAFT roadmap",
         f"Only the DRAFT roadmap lists it ({c['T18b']['dq'][5]}).")),
        (("T18b",), "T18bc", lambda: (f"How many units of part {c['T18b']['cancelled_part']} are on order?",
         "refusal - the only row for this part has status=cancelled", "The CSV row exists but its order is cancelled; nothing is on order.")),
        (("T18b",), "T18bc", lambda: (f"What is the debug UART baud rate of the {c['T18b']['smc']} management controller?",
         "refusal - the value existed only in the WITHDRAWN r1 specification",
         f"r1 lists {c['T18b']['baud']} baud; r2 and r3 drop the line.")),
        (("T18b",), "T18bc", lambda: (f"What is the heartbeat jitter budget of the {c['T18b']['smc']} management controller, in microseconds?",
         "refusal - the value existed only in r2, which is superseded (title says so; the name does not)",
         f"r2 lists {c['T18b']['jit']} us; r3 (current) removed the entry.")),
        (("T20",), "T20c", lambda: (f"What error code does the stress run log for a {c['T20']['absent'][1]}?",
         "refusal - the event is never logged in the 3,000-line run", "No ERROR line mentions it; 40 decoy errors have other symptoms.")),
        (("T20",), "T20c", lambda: (f"Which replacement part was shipped for {c['T20']['absent_rma']}?",
         "refusal - false premise (the RMA id is not among the 5,000 rows)", "Well-formed id, absent from the CSV.")),
        (("T23",), "T23c", lambda: (f"Which part number was last replaced in chassis {c['T23']['ch_absent']}?",
         "refusal - false premise (the chassis id does not exist)", "Neither the registry nor the service log has this chassis.")),
        (("T15b",), "T15bc", lambda: (f"How many days did it take to resolve ticket {u.tix}-{c['T15b_ticket'][1] // 10}?",
         "refusal - the id is only a prefix of real ids (token boundary)",
         f"{u.tix}-{c['T15b_ticket'][1]} and its longer siblings exist; this shorter id does not.")),
        (("T15b",), "T15bc", lambda: (f"What is the image size, in MB, of bootloader version {c['T15b_ver'][0]}.100?",
         "refusal - the version is only a prefix-extension of real ones", f"Builds {c['T15b_ver'][1][0]}..{c['T15b_ver'][1][3]} exist; .100 does not.")),
        (("T15b",), "T15bc", lambda: (f"Which signal is on pin {c['T15b_pin'][0]}{c['T15b_pin'][1]}5 of the {p2} expansion connector?",
         "refusal - the pin does not exist (only a longer/shorter relative does)", "The connector table has no such pin.")),
        (("T25",), "T25c", lambda: (f"What is the lead time, in weeks, of the {c['T25']['only'].lower()} for the {c['T25']['other_prod']}?",
         "refusal - the component exists for a sibling product only (cross-sheet join finds no row)",
         f"Parts lists the {c['T25']['only'].lower()} for the {c['T25']['only_prod']} only.")),
        (("T25",), "T25c", lambda: (f"What is the unit cost in USD of the {c['T25']['ccomp'].lower()} for the {c['T25']['cprod']} at 10,000 unit volume?",
         "refusal - the quantity break is not in the Costs sheet", "Costs has breaks at 100, 1000 and 5000 only.")),
        (("T26",), "T26c", lambda: (f"The production log shows a {c['T26']['noun_absent']} incident. Which firmware release fixed the underlying defect?",
         "refusal - only a topical ticket exists (support/incident_tickets.csv); no log shows the incident",
         f"{c['T26']['ticket_absent']} ({c['T26']['v_absent']}) mentions the topic but is not the defect behind any logged incident.")),
        (("T27",), "T27c", lambda: (f"What error code is logged when the {c['T27']['never']} occurs?",
         "refusal - the event never occurs in any of the service logs", "No ERROR line describes it.")),
        (("T27",), "T27c", lambda: (f"What error code is logged when the {c['T27']['b']} occurs on {c['T27']['first']}?",
         "refusal - the event occurs, but not in the log of that date", "It is logged once, on another date.")),
        (("T27",), "T27c", lambda: (f"What error code is logged when the {c['T27']['a']} occurs on {c['T27']['no_log']}?",
         "refusal - no service log exists for that date", "The logs cover other dates.")),
        (("T28",), "T28c", lambda: (f"What is the maximum coolant inlet temperature of the {c['T28']['lo']}?",
         "refusal - the table cell is '-' (not applicable) for this product", "Only the larger products list an inlet limit.")),
        (("T28",), "T28c", lambda: (f"What is the board power of the {u.nx_product} according to the {u.pre}-series family datasheet?",
         "refusal - false premise (the product does not exist)", f"The family table has no {u.nx_product} column.")),
        (("T21",), "T21c", lambda: (f"What is the maximum number of racks permitted at the {[s for s in SITES if s != c['T21']['site']][0]} lab?",
         "refusal - the site survey is for a different site", f"The only survey in the corpus is for {c['T21']['site']}.")),
    ]
    for needs, trap, fn in items:
        if all(n in ctx.traps for n in needs):
            query, cat, note = fn()
            ctx.q(trap, query, "", [], cat, note=note, unanswerable=True)


# ------------------------------------------------------------------ always-present file traps

def vendor_files(ctx: Ctx):
    u = ctx.u
    ctx.empty_dirs.append("archive")
    a = ctx.text("vendor/internal_audit.txt")
    a.lines = [f"Internal audit {u.year} - supplier compliance", "",
               f"Open findings: {u.nfind}",
               f"Late-delivery penalty: {u.penalty} percent of order value per week"]
    a.chmod, a.expect = 0, "unreadable"
    a.hidden_text = "\n".join(a.lines)
    a.note = "chmod 000"
    p = ctx.doc(f"vendor/{slug(u.supplier)}_agreement_ENCRYPTED.pdf", "pdf")
    p.blocks = [("title", f"Supply agreement - {u.company} / {u.supplier}"),
                ("h", "Schedule B - Pricing"),
                ("l", f"Unit price of the {u.p['p0']} at {comma(u.vol)} unit volume: {u.price} USD"),
                ("h", "Terms"),
                ("l", f"Payment terms: net {u.pay_days} days"),
                ("l", f"Liability cap: {comma(u.liab)} USD")]
    p.password, p.expect = "secret", "encrypted"
    p.hidden_text = p.text() or "\n".join(t for _, t in p.blocks)
    p.note = "encrypted (user password 'secret')"
    n = ctx.text("vendor/supplier_notes.txt")
    n.lines = [f"Supplier notes - {u.supplier}", "",
               "Pricing schedules are held in the supply agreement itself and are not reproduced here.",
               "Contact the procurement desk for the current schedule."]
    b = ctx.doc("vendor/telemetry_capture.dat", "bin")
    b.raw = random.Random(f"mc3dat:{ctx.seed}").randbytes(16384)
    b.expect = "binary"


# ------------------------------------------------------------------ canaries

def canaries(ctx: Ctx):
    u = ctx.u
    p0, p1 = u.p["p0"], u.p["p1"]
    enc = f"vendor/{slug(u.supplier)}_agreement_ENCRYPTED.pdf"
    items = [
        # (needs, trap, query, category, note)
        ((), "T16", f"What is the unit price of the {p0} at {comma(u.vol)} unit volume?",
         "refusal - the answer exists only in the ENCRYPTED file",
         f"The price ({u.price}) is inside {enc}, which cannot be opened. Returning it means cracking the file or hallucinating."),
        ((), "T16", f"What are the payment terms, in days, in the {u.supplier} supply agreement?",
         "refusal - the answer exists only in the ENCRYPTED file", f"Net {u.pay_days} appears only inside the encrypted PDF."),
        ((), "T16", f"What is the liability cap in the {u.supplier} supply agreement?",
         "refusal - the answer exists only in the ENCRYPTED file", f"{comma(u.liab)} USD appears only inside the encrypted PDF."),
        (("T1",), "T17", f"What is the maximum junction temperature of the {u.nx_product}?",
         "refusal - false premise (the product does not exist)", f"No {u.nx_product} anywhere in the corpus."),
        (("T12",), "T17", f"In which quarter does the {u.nx_product} enter customer sampling?",
         "refusal - false premise (the product does not exist)", f"The roadmap lists no {u.nx_product}."),
        (("T5",), "T17", f"What is the part number of the fan assembly for the {u.nx_product}?",
         "refusal - false premise (the product does not exist)", f"The parts list has no {u.nx_product} row."),
        (("T6",), "T17", f"Which firmware version fixed ticket {u.nx_tickets[0]}?",
         "refusal - false premise (the ticket does not exist)", "The id is well-formed but absent from the bug database."),
        (("T6",), "T17", f"What severity is ticket {u.nx_tickets[1]}?",
         "refusal - false premise (the ticket does not exist)", "The id is well-formed but absent from the bug database."),
        (("T7",), "T17", f"Which firmware release fixed the defect reported by device {u.nx_dev}?",
         "refusal - false premise (the device does not exist)", f"{u.nx_dev} appears in no log and not in fleet.csv."),
        (("T1",), "T18", f"What is the maximum storage temperature of the {p0}?",
         "refusal - the value existed only in the WITHDRAWN r1 datasheet",
         f"r1 lists {u.storage_max} C but r1 is withdrawn and r2 removed the line. Answering cites a withdrawn document."),
        (("T1",), "T18", f"What is the maximum operating altitude of the {p0}, in metres?",
         "refusal - the value existed only in the WITHDRAWN r1 datasheet",
         f"r1 lists {comma(u.altitude)} m but r1 is withdrawn and r2 removed the line."),
        ((), "T16", "What late-delivery penalty does the internal audit specify, in percent?",
         "refusal - the answer is in an unreadable (chmod 000) file", f"{u.penalty} percent appears only in vendor/internal_audit.txt, which cannot be read."),
        ((), "T16", "How many open findings does the internal audit report?",
         "refusal - the answer is in an unreadable (chmod 000) file", f"{u.nfind} appears only in vendor/internal_audit.txt, which cannot be read."),
        (("T5",), "T16", f"What is the typical acoustic noise of the {p0}, in dB(A)?",
         "refusal - the attribute exists for a sibling product only",
         f"Only the {p1} datasheet lists acoustic noise ({u.noise_p1} dB(A))."),
        (("T5",), "T16", f"What is the part number of the coolant pump assembly for the {p0}?",
         "refusal - the attribute exists for a sibling product only",
         f"The pump row ({u.pump_p1}) is for the {p1}; no {p0} pump is listed."),
        (("T12",), "T16", f"In which quarter does the {p0} enter customer sampling?",
         "refusal - the table has no value for this product (already in volume production)",
         f"The roadmap row for the {p0} has '-' in the sampling column."),
        (("T1",), "T19", f"What is the maximum junction temperature of the {p0} in degrees Fahrenheit?",
         "refusal - needs a unit conversion that the corpus does not contain",
         f"The datasheet gives {u.jt_new} C only; do not convert."),
        (("T8",), "T19", f"What is the combined board power of four {p0} boards?",
         "refusal - needs arithmetic that the corpus does not contain",
         f"Per-board power is {u.pw_new} W; no document states a four-board total."),
        (("T6",), "T16", "What error code is logged when the coolant pump stalls?",
         "refusal - no log line describes this event", "The production log has a different incident only."),
        (("T6",), "T17", f"Which ticket was opened for error code {u.nx_err}?",
         "refusal - false premise (the error code does not exist)", f"{u.nx_err} appears in no log."),
    ]
    for needs, trap, query, cat, note in items:
        if all(n in ctx.traps for n in needs):
            ctx.q(trap, query, "", [], cat, note=note, unanswerable=True)


# ------------------------------------------------------------------ fillers

def _filler_text(ctx: Ctx, kind: str, i: int, rng: random.Random):
    u = ctx.u
    prods = list(u.p.values())

    def meeting():
        d = f"{u.year}-{rng.randint(1, 9):02d}-{rng.randint(10, 28)}"
        ls = [f"Meeting notes {d}", "", f"Attendees: {rng.randint(4, 12)}", ""]
        for _ in range(rng.randint(5, 9)):
            ls.append(rng.choice(SENTENCES).replace("backplane cable", f"{rng.choice(prods)} backplane cable"))
        ls += ["", f"Next review in {rng.randint(2, 6)} weeks; budget line {rng.randint(100, 999)}."]
        return f"planning/meeting_notes_{d}_{i}.txt", "text", ls

    def shipments():
        ls = ["order_id,customer,product,qty,ship_date"]
        for _ in range(rng.randint(8, 14)):
            ls.append(f"SO-{rng.randint(10000, 99999)},{rng.choice(SITES)} Labs,{rng.choice(prods)},{rng.randint(2, 40)},"
                      f"2026-0{rng.randint(1, 8)}-{rng.randint(10, 28)}")
        return f"support/shipments_{i}.csv", "text", ls

    def cilog():
        t = dt.datetime(2026, rng.randint(6, 9), rng.randint(1, 28), rng.randint(0, 23), 0, 0)
        ls = []
        for _ in range(rng.randint(10, 18)):
            t += dt.timedelta(seconds=rng.randint(20, 400))
            ls.append(f"{t.strftime('%Y-%m-%dT%H:%M:%SZ')} ci-runner-{rng.randint(1, 6)}: "
                      f"{rng.choice(['INFO', 'INFO', 'WARN'])} job {rng.randint(1000, 9999)} "
                      f"{rng.choice(['passed', 'passed', 'retried'])} in {rng.randint(8, 400)}s")
        return f"logs/ci_build_{i}.log", "text", ls

    def blog():
        ls = [f"# {rng.choice(['Notes from the lab', 'Rack deployment tips', 'Field report'])}", ""]
        for _ in range(rng.randint(5, 9)):
            ls += [rng.choice(SENTENCES), ""]
        return f"marketing/blog_{i}.md", "text", ls

    def py():
        n = rng.choice(["retry", "paging", "cache", "export"])
        ls = [f'"""Helpers for {n} handling."""', "", f"PAGE_SIZE = {rng.randint(20, 500)}",
              f"CACHE_TTL_S = {rng.randint(30, 900)}", "", "", f"def {n}_helper(items, size=PAGE_SIZE):",
              "    return [items[i:i + size] for i in range(0, len(items), size)]"]
        return f"engineering/util_{n}_{i}.py", "text", ls

    def pdf():
        blocks = [("title", f"{u.company} application note {i}"), ("p", "Integration guidance for rack builders.")]
        for _ in range(rng.randint(4, 8)):
            blocks.append(("p", rng.choice(SENTENCES)))
        return f"specs/app_note_{i}.pdf", "pdf", blocks

    def docx():
        blocks = [("h", f"Memo {i}", 1)] + [("p", rng.choice(SENTENCES)) for _ in range(rng.randint(4, 8))]
        return f"planning/memo_{i}.docx", "docx", blocks

    def xlsx():
        rows = [(rng.choice(["Validation", "Support", "Field", "Logistics"]), rng.randint(2, 20),
                 rng.randint(100, 900)) for _ in range(rng.randint(5, 9))]
        return f"planning/headcount_{i}.xlsx", "xlsx", [{"name": "Teams", "header": ["team", "people", "budget_k"], "rows": rows}]

    return {"txt": meeting, "csv": shipments, "log": cilog, "md": blog, "py": py, "pdf": pdf,
            "docx": docx, "xlsx": xlsx}[kind]()


def fillers(ctx: Ctx, forbidden: list[re.Pattern]):
    rng = ctx.rng_fill
    order = ["txt", "csv", "log", "md", "py", "pdf", "docx", "xlsx"]
    i = 0
    while len(ctx.docs) < ctx.n_files:
        kind = order[i % len(order)]
        i += 1
        for _ in range(30):
            path, dk, payload = _filler_text(ctx, kind, i, rng)
            d = Doc(path, dk)
            if dk == "text":
                d.lines = payload
            elif dk == "xlsx":
                d.sheets = payload
            else:
                d.blocks = payload
            if not any(f.search(d.text()) for f in forbidden):
                break
        if path not in ctx.docs:
            ctx.docs[path] = d


# ------------------------------------------------------------------ build + validation

def cap_refusals(ctx: Ctx, keep: int):
    """Keep only `keep` of the refusal questions, round-robin over the trap tags (small suites with a few canaries)."""
    rng = ctx.xr("CAP")
    refs = [q for q in ctx.queries if q["unanswerable"]]
    by_tag: dict = {}
    for q in refs:
        by_tag.setdefault(q["trap"], []).append(q)
    for v in by_tag.values():
        rng.shuffle(v)
    tags = sorted(by_tag)
    rng.shuffle(tags)
    chosen = []
    while len(chosen) < keep and any(by_tag.values()):
        for t in tags:
            if by_tag[t] and len(chosen) < keep:
                chosen.append(by_tag[t].pop())
    drop = {id(q) for q in refs} - {id(q) for q in chosen}
    ctx.queries = [q for q in ctx.queries if id(q) not in drop]


def build(seed: int, attempt: int, traps: list[str], n_files: int, n_images: int = 0, max_refusals: int | None = None) -> Ctx:
    ctx = Ctx(seed, attempt, traps, n_files, n_images)
    vendor_files(ctx)
    for t in traps:
        globals()[f"trap_{t}"](ctx)
    canaries(ctx)
    canaries_x(ctx)
    if n_images:
        images_extra(ctx, n_images)
    if max_refusals is not None:
        cap_refusals(ctx, max_refusals)
    u = ctx.u
    forbidden = [bound_re(x) for x in (u.nx_product, *u.nx_tickets, u.nx_dev, u.nx_err, u.v6, u.v7, u.t6, u.t7,
                                       u.e6, u.e7, u.pin, u.rev, u.serial, u.fcc, u.q_tapeout, *ctx.forbid,
                                       *ctx.absent) if x]
    for q in ctx.queries:
        if is_ident(q["expected_answer"]):
            forbidden.append(bound_re(q["expected_answer"]))
    ctx.n_files = max(ctx.n_files, len(ctx.docs) + 8)       # always some background noise, however many traps ran
    fillers(ctx, forbidden)
    return ctx


def validate(ctx: Ctx) -> list[str]:
    """Self-consistency of the gold; returns problems (empty = ok)."""
    bad = []
    texts = {p: d.text() for p, d in ctx.docs.items()}
    image_paths = {p for p, d in ctx.docs.items() if d.kind in ("png", "jpg")}
    for q in ctx.queries:
        a, cites = q["expected_answer"], q["expected_citations"]
        if q["unanswerable"]:
            continue
        for c in cites:
            if c not in ctx.docs:
                bad.append(f"cited file missing: {c}")
        image_only = bool(cites) and all(c in image_paths for c in cites) or \
            (q["trap"] == "T11") or q["imgonly"]
        if image_only:
            # short numeric image-only answers (a wattage, ...) would collide with any number in the corpus: for those,
            # only the cited files' own text must not carry them
            scope = texts if (is_ident(a) or not q["imgonly"]) else {c: texts[c] for c in cites if c in texts}
            hit = [p for p, t in scope.items() if bound_re(a).search(t)]
            if hit:
                bad.append(f"image-only answer {a!r} appears in text of {hit}")
            continue
        blob = "\n".join(texts[c] for c in cites)
        cands = [a] + q["answer_aliases"]
        if not any(re.sub(r"[\s\-.·_]", "", x).upper() in re.sub(r"[\s\-.·_]", "", blob).upper() for x in cands):
            bad.append(f"answer {a!r} not found in cited text for: {q['query']}")
        if is_ident(a):
            stray = [p for p, t in texts.items() if p not in cites and p not in q["also"] and bound_re(a).search(t)]
            if stray:
                bad.append(f"identifier answer {a!r} also appears in uncited {stray}")
    # canary secrets must stay out of plain-text files
    u = ctx.u
    for tok, allowed in ctx.secrets.items():
        hit = [p for p, t in texts.items() if p not in allowed and bound_re(tok).search(t)]
        if hit:
            bad.append(f"image-only token {tok} found in {hit}")
    for sec in (u.nx_product, *u.nx_tickets, u.nx_dev, u.nx_err, *ctx.absent):
        hit = [p for p, t in texts.items() if bound_re(sec).search(t)]
        if hit:
            bad.append(f"nonexistent id {sec} found in {hit}")
    return bad


def render(ctx: Ctx, out: Path) -> dict:
    corpus = out / "corpus"
    if corpus.exists():
        shutil.rmtree(corpus)
    corpus.mkdir(parents=True)
    for p, d in sorted(ctx.docs.items()):
        write_doc(d, corpus / p)
    for e in ctx.empty_dirs:
        (corpus / e).mkdir(parents=True, exist_ok=True)
    for p, d in ctx.docs.items():
        if d.chmod is not None:
            os.chmod(corpus / p, d.chmod)

    qs = ctx.queries[:]
    ctx.rng_fill.shuffle(qs)
    queries, traps = [], {}
    for n, q in enumerate(qs, start=1):
        e = {"n": n, "query": q["query"], "expected_answer": q["expected_answer"],
             "answer_aliases": q["answer_aliases"], "expected_citations": q["expected_citations"],
             "category": q["category"], "note": q["note"]}
        if q["unanswerable"]:
            e["unanswerable"] = True
        queries.append(e)
        traps[str(n)] = q["trap"]
    gold = {"challenge": "mini-challenge-3",
            "description": f"RAG over a mixed-filetype corpus. ADVERSARIAL generated set (seed {ctx.seed}); not an official set.",
            "corpus": "mc3-corpus", "answer_rule": ANSWER_RULE, "queries": queries, "seed": ctx.seed}
    (out / "questions.json").write_text(json.dumps(gold, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (out / "traps.json").write_text(json.dumps(traps, indent=2) + "\n", encoding="utf-8")

    u = ctx.u
    ml = [f"MANIFEST seed={ctx.seed} attempt={ctx.attempt} traps={','.join(ctx.traps)}",
          f"company={u.company} firmware={u.fw} product_line={u.pre} hero={u.p['p0']}",
          f"files={len(ctx.docs)} (+ empty dirs: {', '.join(ctx.empty_dirs)}) questions={len(queries)}", "",
          f"{'path':58} {'bytes':>8}  expected_parse_status  notes"]
    for p, d in sorted(ctx.docs.items()):
        size = (corpus / p).stat().st_size
        notes = []
        if d.chmod is not None:
            notes.append("CHMOD 000 (unreadable to non-root)")
        if d.password:
            notes.append(f"ENCRYPTED PDF (user password '{d.password}'; content is the answer to unanswerable questions)")
        if "WITHDRAWN" in p:
            notes.append("withdrawn revision")
        if d.kind in ("png", "jpg"):
            notes.append("image-only facts")
        if d.kind == "bin":
            notes.append("random bytes")
        if d.note and d.note != "chmod 000" and not d.password:
            notes.append(d.note)
        if d.encoding != "utf-8":
            notes.append(f"encoding {d.encoding}")
        if d.newline != "\n":
            notes.append("CRLF line endings")
        ml.append(f"{p:58} {size:>8}  {d.expect:21}  {'; '.join(notes)}")
    for e in ctx.empty_dirs:
        ml.append(f"{e + '/':58} {'-':>8}  empty directory")
    (out / "MANIFEST.txt").write_text("\n".join(ml) + "\n", encoding="utf-8")
    return gold


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--seed", type=int, required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--n-files", type=int, default=40)
    ap.add_argument("--max-refusals", type=int, default=None,
                    help="keep only N of the refusal (canary) questions, spread over the trap tags; default: all")
    ap.add_argument("--n-images", type=int, default=0,
                    help="add N png/jpg files with printed text (>= 9: 9 answer images with 10 questions + 3 refusals, the "
                         "rest look-alike fillers); counted in --n-files; default 0 = none, seeds stay reproducible")
    ap.add_argument("--traps", default=",".join(TRAPS),
                    help="comma list, default all: " + ",".join(TRAPS) + "; opt-in only (never in the default): "
                    + ",".join(EXTRA_TRAPS))
    a = ap.parse_args(argv)
    canon = {x.upper(): x for x in TRAPS + EXTRA_TRAPS}
    traps = [canon.get(t.strip().upper(), t.strip()) for t in a.traps.split(",") if t.strip()]
    unknown = [t for t in traps if t not in TRAPS + EXTRA_TRAPS]
    if unknown:
        ap.error(f"unknown traps: {unknown}")
    for t, others in TRAP_CONFLICTS.items():
        clash = [o for o in others if t in traps and o in traps]
        if clash:
            ap.error(f"{t} cannot be combined with {clash} (same attributes, different values)")
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    for attempt in range(50):
        ctx = build(a.seed, attempt, traps, a.n_files, a.n_images, a.max_refusals)
        bad = validate(ctx)
        if not bad:
            break
        print(f"attempt {attempt}: {len(bad)} gold problem(s), redrawing: {bad[0]}", file=sys.stderr)
    else:
        print("could not build a self-consistent suite", file=sys.stderr)
        return 1
    gold = render(ctx, out)
    n_un = sum(1 for q in gold["queries"] if q.get("unanswerable"))
    print(f"seed {a.seed}: {len(ctx.docs)} files, {len(gold['queries'])} questions "
          f"({len(gold['queries']) - n_un} answerable, {n_un} refusals) -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
