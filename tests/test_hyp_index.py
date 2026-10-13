"""Hypothesis flags (H1 XLSX_JOIN, H4 RANK, H5 PDF_COLS): derived records, formula refs, ranking order.

Every flag defaults to v5 behaviour; the first tests pin that, the rest turn one flag on via monkeypatch."""

import json

import pytest

import config
import index as ix
import parse
import textnorm
import verify
from conftest import SAMPLE_CORPUS
from schema import FileEntry, Package, Q, ReaderOut, Record


def pf(path, kind, records):
    """Hand-made parse.walk entry; records = [(loc, text) | (loc, text, stale)]."""
    recs = [dict(r) if isinstance(r, dict) else
            {"loc": r[0], "text": r[1], "stale": bool(r[2]) if len(r) > 2 else False} for r in records]
    return {"path": path, "kind": kind, "status": "ok", "records": recs,
            "text": "\n".join(r["text"] for r in recs), "images": [],
            "flags": {"withdrawn": False, "supersedes": [], "revision": None}}


def sheet_row(sheet, n, **kv):
    return (f"{sheet}!R{n}", f"[sheet {sheet}] " + "; ".join(f"{k}={v}" for k, v in kv.items()))


def workbook(extra=()):
    recs = [("Parts!R1", "[sheet Parts] columns: part_number | description | product"),
            sheet_row("Parts", 2, part_number="ORR-PSU-3301-A", description="Power supply unit", product="TQ-40"),
            sheet_row("Parts", 3, part_number="ORR-PSU-3302-A", description="Power supply unit", product="TQ-20"),
            ("LeadTimes!R1", "[sheet LeadTimes] columns: part_number | lead_time_weeks"),
            sheet_row("LeadTimes", 2, part_number="ORR-PSU-3301-A", lead_time_weeks=14),
            sheet_row("LeadTimes", 3, part_number="ORR-PSU-3302-A", lead_time_weeks=9)]
    return pf("support/parts.xlsx", "xlsx", recs + list(extra))


def derived(idx, src):
    return [r for r in idx.records if r.src == src]


# ---------------------------------------------------------------- flags off = v5

def test_flags_default_off():
    assert config.XLSX_JOIN is False and config.PDF_COLS is False and config.RANK == "default"


def test_off_build_has_no_derived_records():
    idx = ix.build([workbook(), pf("specs/fam.pdf", "pdf", PDF_TABLE)])
    assert {r.src for r in idx.records} == {"name", "text"}


def test_off_sample_corpus_unchanged(monkeypatch):
    parsed = parse.walk(SAMPLE_CORPUS)
    off = ix.build(parsed)
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    monkeypatch.setattr(config, "PDF_COLS", True)
    on = ix.build(parsed)
    plain = lambda i: [(r.file, r.loc, r.text, r.src) for r in i.records if r.src in ("name", "text")]
    assert plain(on) == plain(off)           # derived records only ever add
    assert all(r.src in ("name", "text") for r in off.records)


# ---------------------------------------------------------------- H1 join

def test_join_created_only_with_flag(monkeypatch):
    assert not derived(ix.build([workbook()]), "join")
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    idx = ix.build([workbook()])
    joins = derived(idx, "join")
    assert [r.loc for r in joins] == ["LeadTimes!R2~Parts!R2", "LeadTimes!R3~Parts!R3"]
    j = joins[0]
    assert j.file == "support/parts.xlsx" and j.src == "join" and not j.meta and not j.stale
    assert j.text == ("[join] [sheet LeadTimes] part_number=ORR-PSU-3301-A; lead_time_weeks=14 || "
                      "[sheet Parts] part_number=ORR-PSU-3301-A; description=Power supply unit; product=TQ-40")
    # ids, by_file, rid_range and BM25 see the appended records
    fe = idx.file_by_path["support/parts.xlsx"]
    assert idx.by_file["support/parts.xlsx"] == list(range(*fe.rid_range))
    assert [r.rid for r in idx.records] == list(range(len(idx.records)))
    e = idx.ids[textnorm.nid("ORR-PSU-3301-A")]
    assert j.rid in e["rids"] and e["df"] == 1
    assert idx.bm25.search(["lead_time_weeks", "supply"], 5)[0][0] in {r.rid for r in joins}
    assert fe.n_products == 2


def test_join_stale_follows_either_row(monkeypatch):
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    w = workbook()
    w["records"][4]["stale"] = True            # LeadTimes!R2 (a hidden-sheet / legacy row)
    joins = {r.loc: r for r in derived(ix.build([w]), "join")}
    assert joins["LeadTimes!R2~Parts!R2"].stale and not joins["LeadTimes!R3~Parts!R3"].stale
    w = workbook()
    w["records"][2]["stale"] = True            # Parts!R3
    joins = {r.loc: r for r in derived(ix.build([w]), "join")}
    assert joins["LeadTimes!R3~Parts!R3"].stale and not joins["LeadTimes!R2~Parts!R2"].stale


def test_join_not_for_same_sheet_or_non_product_rows(monkeypatch):
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    same = pf("a.xlsx", "xlsx", [sheet_row("Parts", 2, part_number="ORR-PSU-3301-A", product="TQ-40"),
                                 sheet_row("Parts", 3, part_number="ORR-PSU-3301-A", lead_time_weeks=14)])
    assert not derived(ix.build([same]), "join")
    no_product = pf("b.xlsx", "xlsx", [sheet_row("Parts", 2, part_number="ORR-PSU-3301-A", description="PSU"),
                                       sheet_row("LeadTimes", 2, part_number="ORR-PSU-3301-A", lead_time_weeks=14)])
    assert not derived(ix.build([no_product]), "join")
    # R already names the product S names: nothing to add
    both = pf("c.xlsx", "xlsx", [sheet_row("Parts", 2, part_number="ORR-PSU-3301-A", product="TQ-40"),
                                 sheet_row("LeadTimes", 2, part_number="ORR-PSU-3301-A", product="TQ-40", w=14)])
    assert not derived(ix.build([both]), "join")
    # no shared identifier
    none = pf("d.xlsx", "xlsx", [sheet_row("Parts", 2, part_number="ORR-PSU-3301-A", product="TQ-40"),
                                 sheet_row("LeadTimes", 2, part_number="ORR-PSU-9999-A", lead_time_weeks=14)])
    assert not derived(ix.build([none]), "join")


def test_join_ignores_decimals_products_and_header_rows(monkeypatch):
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    # a shared decimal (298.5) and a shared product code (TQ-40) are not join keys
    w = pf("e.xlsx", "xlsx", [sheet_row("Parts", 2, product="TQ-40", price="298.5"),
                              sheet_row("Costs", 2, usd="298.5", model="TQ-40", note="a"),
                              sheet_row("Costs", 3, usd="298.5", model="TQ-60", note="b")])
    assert not derived(ix.build([w]), "join")
    # header rows never join
    h = pf("f.xlsx", "xlsx", [("Parts!R1", "[sheet Parts] columns: part_number | product ORR-PSU-3301-A TQ-40"),
                              sheet_row("LeadTimes", 2, part_number="ORR-PSU-3301-A", lead_time_weeks=14)])
    assert not derived(ix.build([h]), "join")


def test_join_only_for_xlsx(monkeypatch):
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    w = workbook()
    csv_like = pf("x.csv", "csv", w["records"])
    assert not derived(ix.build([csv_like]), "join")


def test_join_cap_per_workbook(monkeypatch):
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    recs = []
    for i in range(700):
        recs.append(sheet_row("Parts", i + 2, part_number=f"PN-{i:04d}-A", product="TQ-40"))
    for i in range(700):
        recs.append(sheet_row("LeadTimes", i + 2, part_number=f"PN-{i:04d}-A", lead_time_weeks=i))
    idx = ix.build([pf("big.xlsx", "xlsx", recs)])
    assert len(derived(idx, "join")) == ix.JOIN_CAP == 500
    # the cap is per workbook
    idx2 = ix.build([pf("big1.xlsx", "xlsx", recs), pf("big2.xlsx", "xlsx", recs)])
    assert len(derived(idx2, "join")) == 1000


def test_join_rows_per_key_limited(monkeypatch):
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    recs = [sheet_row("Bom", i, part_number="ORR-PSU-3301-A", product=f"TQ-{i}0") for i in range(2, 12)]
    recs.append(sheet_row("LeadTimes", 2, part_number="ORR-PSU-3301-A", lead_time_weeks=14))
    assert len(derived(ix.build([pf("bom.xlsx", "xlsx", recs)]), "join")) == 4


def test_join_survives_save_and_load(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    idx = ix.build([workbook()])
    ix.save(tmp_path, idx)
    back = ix.load(tmp_path)
    assert [r.to_json() for r in back.records] == [r.to_json() for r in idx.records]
    assert back.ids == json.loads(json.dumps(idx.ids))


def test_join_record_verifies_entity(monkeypatch):
    """The LeadTimes row alone fails the entity check (n_products >= 2); the join record passes it."""
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    idx = ix.build([workbook()])
    q = Q("What is the lead time, in weeks, of the power supply unit for the TQ-40?", entities=["TQ-40"],
          attr_terms=["lead_time"])
    lead = next(r for r in idx.records if r.loc == "LeadTimes!R2")
    join = next(r for r in idx.records if r.src == "join" and r.loc.startswith("LeadTimes!R2~"))
    other = next(r for r in idx.records if r.src == "join" and r.loc.startswith("LeadTimes!R3~"))
    pkg = Package(rids=[lead.rid, join.rid, other.rid], text="")
    v = verify.verify(q, ReaderOut("14", support=[lead.rid]), pkg, idx)
    assert v.ok and v.support_rid == join.rid and v.answer_span == "14"
    # the join of the other product's part cannot vouch for TQ-40
    v2 = verify.verify(q, ReaderOut("9", support=[other.rid]), pkg, idx)
    assert not v2.ok and v2.reason == "ENTITY_MISMATCH"


# ---------------------------------------------------------------- H1 formulas across sheets

def _formula_book(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parts"
    ws.append(["part_number", "weeks_x2", "rng"])
    ws.append(["ORR-PSU-3301-A", "=LeadTimes!B2*2", "=SUM(LeadTimes!B2:B3)"])
    ws.append(["ORR-PSU-3302-A", "='Lead Times'!B2+1", "=Ghost!B2"])
    lt = wb.create_sheet("LeadTimes")      # defined after Parts: a forward reference
    lt.append(["part_number", "lead_time_weeks"])
    lt.append(["ORR-PSU-3301-A", 7])
    lt.append(["ORR-PSU-3302-A", 5])
    sp = wb.create_sheet("Lead Times")
    sp.append(["part_number", "lead_time_weeks"])
    sp.append(["x", 40])
    wb.save(path)


def _xlsx_texts(path):
    return {r["loc"]: r["text"] for r in parse._xlsx(path)}


def test_formula_cross_sheet_flag_off_is_v5(tmp_path):
    p = tmp_path / "f.xlsx"
    _formula_book(p)
    t = _xlsx_texts(p)
    # v5 reads LeadTimes!B2 from the Parts grid, where B2 is the formula itself (no number): 0*2
    assert "weeks_x2=0" in t["Parts!R2"]


def test_formula_cross_sheet_resolved_with_flag(monkeypatch, tmp_path):
    p = tmp_path / "f.xlsx"
    _formula_book(p)
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    t = _xlsx_texts(p)
    assert "weeks_x2=14" in t["Parts!R2"] and "rng=12" in t["Parts!R2"]
    assert "weeks_x2=41" in t["Parts!R3"]       # quoted sheet name with a space
    assert "rng" not in t["Parts!R3"] or "rng==Ghost!B2" in t["Parts!R3"]   # unknown sheet: no invented value


def test_eval_formula_sheets_argument():
    grids = {"other": {(1, 1): 5, (2, 1): 6}}
    assert parse._eval_formula("=Other!A1+A1", {(1, 1): 1}, grids) == 6
    assert parse._eval_formula("=SUM(Other!A1:A2)", {}, grids) == 11
    assert parse._eval_formula("=Nope!A1", {}, grids) is None
    assert parse._eval_formula("=A1*2", {(1, 1): 4}, grids) == 8
    assert parse._eval_formula("=Other!A1", {(1, 1): 1}) == 1     # sheets=None: v5 reads the current grid


# ---------------------------------------------------------------- H5 PDF family columns

PDF_TABLE = [("p1.L1", "Orrery Systems TQ-series Accelerators - family datasheet, revision 3"),
             ("p1.L2", "Parameter                  TQ-20       TQ-40       TQ-60"),
             ("p1.L3", "Maximum junction temp      85 C        94 C        98 C"),
             ("p1.L4", "Board power (TBP)          300 W       410 W       520 W"),
             ("p1.L5", "Liquid cooling             —           yes-opt     n/a"),
             ("p1.L6", "Memory capacity            32 GiB      56 GiB      96 GiB"),
             ("p1.L7", "Notes follow the table.")]


def test_cols_only_with_flag(monkeypatch):
    assert not derived(ix.build([pf("s.pdf", "pdf", PDF_TABLE)]), "col")
    monkeypatch.setattr(config, "PDF_COLS", True)
    idx = ix.build([pf("s.pdf", "pdf", PDF_TABLE)])
    cols = derived(idx, "col")
    got = {r.loc: r.text for r in cols}
    assert got["p1.L3#TQ-40"] == "[col] TQ-40: Maximum junction temp 94 C"
    assert got["p1.L3#TQ-60"] == "[col] TQ-60: Maximum junction temp 98 C"
    assert got["p1.L4#TQ-20"] == "[col] TQ-20: Board power (TBP) 300 W"
    assert all(r.file == "s.pdf" and not r.meta and not r.stale for r in cols)
    # "—" skipped; the text cell "yes-opt" splits because the row is column-aligned (2+ space gutters)
    assert "p1.L5#TQ-20" not in got and got["p1.L5#TQ-40"] == "[col] TQ-40: Liquid cooling yes-opt"
    assert "p1.L5#TQ-60" not in got      # n/a is an empty cell
    assert got["p1.L6#TQ-60"] == "[col] TQ-60: Memory capacity 96 GiB"
    assert "p1.L7#TQ-20" not in got      # prose ends the table
    fe = idx.file_by_path["s.pdf"]
    assert idx.by_file["s.pdf"] == list(range(*fe.rid_range))


def test_cols_dash_cells_skipped(monkeypatch):
    monkeypatch.setattr(config, "PDF_COLS", True)
    t = [("p1.L1", "Parameter    TQ-20    TQ-40    TQ-60"),
         ("p1.L2", "Fan count    —        4        6"),
         ("p1.L3", "Rails        -        n/a      3")]
    got = {r.loc: r.text for r in derived(ix.build([pf("s.pdf", "pdf", t)]), "col")}
    assert set(got) == {"p1.L2#TQ-40", "p1.L2#TQ-60", "p1.L3#TQ-60"}
    assert got["p1.L2#TQ-40"] == "[col] TQ-40: Fan count 4"


def test_cols_single_space_and_bare_numbers(monkeypatch):
    monkeypatch.setattr(config, "PDF_COLS", True)
    t = [("p2.L4", "Parameter | TQ-20 | TQ-40"),
         ("p2.L5", "Max junction temperature 85 94"),
         ("p2.L6", "Peak bandwidth 2.4 TB/s 4.1 TB/s")]
    got = {r.loc: r.text for r in derived(ix.build([pf("s.pdf", "pdf", t)]), "col")}
    assert got["p2.L5#TQ-40"] == "[col] TQ-40: Max junction temperature 94"
    assert got["p2.L6#TQ-20"] == "[col] TQ-20: Peak bandwidth 2.4 TB/s"


def test_cols_stop_conditions(monkeypatch):
    monkeypatch.setattr(config, "PDF_COLS", True)
    # wrong value count, blank line (line-number gap), new page, prose
    t = [("p1.L1", "Parameter    TQ-20    TQ-40"),
         ("p1.L2", "Power        300 W    410 W"),
         ("p1.L3", "Efficiency   92 %"),           # k-1 values: ends the table
         ("p1.L4", "Weight       3 kg     4 kg"),   # after the end: not split
         ("p1.L6", "Header      TQ-60    TQ-80"),   # gap at L5 (blank), and a new header
         ("p1.L7", "Power        520 W    610 W"),
         ("p3.L1", "Power        1 W      2 W")]    # other page, no header
    got = {r.loc for r in derived(ix.build([pf("s.pdf", "pdf", t)]), "col")}
    assert got == {"p1.L2#TQ-20", "p1.L2#TQ-40", "p1.L7#TQ-60", "p1.L7#TQ-80"}


def test_cols_row_limit(monkeypatch):
    monkeypatch.setattr(config, "PDF_COLS", True)
    t = [("p1.L1", "Parameter    TQ-20    TQ-40")]
    t += [(f"p1.L{i}", f"Item number {i}    {i} W    {i + 1} W") for i in range(2, 80)]
    cols = derived(ix.build([pf("s.pdf", "pdf", t)]), "col")
    assert len(cols) == 2 * ix.COL_MAX_ROWS


@pytest.mark.parametrize("rule", ["--------------------", "=====", "____", "-----|-----|-----", "--- --- ---", "────────"])
def test_cols_rule_line_is_skipped(monkeypatch, rule):
    monkeypatch.setattr(config, "PDF_COLS", True)
    t = [("p1.L1", "Parameter    TQ-20    TQ-40"),
         ("p1.L2", rule),
         ("p1.L3", "Power        300 W    410 W"),
         ("p1.L4", rule),
         ("p1.L5", "Weight       3 kg     4 kg"),
         ("p1.L6", "Revision history")]
    got = {r.loc: r.text for r in derived(ix.build([pf("s.pdf", "pdf", t)]), "col")}
    assert set(got) == {"p1.L3#TQ-20", "p1.L3#TQ-40", "p1.L5#TQ-20", "p1.L5#TQ-40"}


def test_cols_rule_line_gap_still_ends_table(monkeypatch):
    monkeypatch.setattr(config, "PDF_COLS", True)
    t = [("p1.L1", "Parameter    TQ-20    TQ-40"), ("p1.L3", "-----"), ("p1.L4", "Power    300 W    410 W")]
    assert not derived(ix.build([pf("s.pdf", "pdf", t)]), "col")


def test_cols_text_cells_need_aligned_row(monkeypatch):
    monkeypatch.setattr(config, "PDF_COLS", True)
    t = [("p1.L1", "Parameter                SK-22        SK-42          SK-62"),
         ("p1.L2", "Auxiliary power connector  8-pin EPS    6-pin PCIe     16-pin 12VHPWR"),
         ("p1.L3", "Typical acoustic noise     46 dB(A)     48 dB(A)       55 dB(A)"),
         ("p1.L4", "Single spaced text row 6-pin PCIe 8-pin EPS 16-pin"),    # single spaces: strict numeric only
         ("p1.L5", "Cooling   air   liquid   liquid")]
    got = {r.loc: r.text for r in derived(ix.build([pf("s.pdf", "pdf", t)]), "col")}
    assert got["p1.L2#SK-42"] == "[col] SK-42: Auxiliary power connector 6-pin PCIe"
    assert got["p1.L3#SK-62"] == "[col] SK-62: Typical acoustic noise 55 dB(A)"
    assert not any(k.startswith("p1.L4") for k in got)
    assert "p1.L5#SK-22" not in got       # table ended at L4


@pytest.mark.parametrize("line", [
    "The TQ-20 and TQ-40 are available now",     # prose, not a trailing run of product IDs
    "Parameter TQ-20",                           # one product
    "Parameter TQ-20 TQ-20",                     # not distinct
    "Rev 2 board data TQ-20 TQ-40",              # digits in the label, too many words
    "Lanes x8 x16",                              # lower-case single-letter prefixes are not products
])
def test_cols_header_rejects(line):
    assert ix._col_header(line) is None


def test_cols_header_accepts():
    assert ix._col_header("Parameter                  TQ-20       TQ-40       TQ-60") == ["TQ-20", "TQ-40", "TQ-60"]
    assert ix._col_header("TQ‑20 TQ‑40") == ["TQ-20", "TQ-40"]      # non-breaking hyphen folded


def test_cols_record_verifies_entity(monkeypatch):
    monkeypatch.setattr(config, "PDF_COLS", True)
    idx = ix.build([pf("s.pdf", "pdf", PDF_TABLE)])
    q = Q("What is the maximum junction temperature of the TQ-60?", entities=["TQ-60"], attr_terms=["temperature"])
    col = next(r for r in idx.records if r.loc == "p1.L3#TQ-60")
    v = verify.verify(q, ReaderOut("98", support=[col.rid]), Package(rids=[col.rid], text=""), idx)
    assert v.ok and v.support_rid == col.rid


# ---------------------------------------------------------------- H4 rank

def _rank_index():
    files = [FileEntry(file=f, kind="txt", status="ok") for f in ("a.txt", "b.txt")]
    recs = [Record(0, "a.txt", "name", "a txt", meta=True, src="name"),
            Record(1, "a.txt", "L1", "ticket noise value 9"),
            Record(2, "b.txt", "name", "b txt", meta=True, src="name"),
            Record(3, "b.txt", "L1", "widget alpha firmware version 9")]
    return ix.Index(files, recs)


def test_rank_default_prefers_hop_target(monkeypatch):
    idx = _rank_index()
    q = Q("What is the firmware version of widget alpha?")
    pkg = Package(rids=[1, 3], text="", via={1: 3})
    v = verify.verify(q, ReaderOut("9", support=[1, 3]), pkg, idx)
    assert v.ok and v.support_rid == 1          # v5: the hop target beats the better-covered record


def test_rank_cov_first_prefers_coverage(monkeypatch):
    monkeypatch.setattr(config, "RANK", "cov_first")
    idx = _rank_index()
    q = Q("What is the firmware version of widget alpha?")
    pkg = Package(rids=[1, 3], text="", via={1: 3})
    v = verify.verify(q, ReaderOut("9", support=[1, 3]), pkg, idx)
    assert v.ok and v.support_rid == 3
    # attr_score still comes first, hop target still breaks ties of equal coverage
    pkg2 = Package(rids=[1, 3], text="", via={1: 3})
    q2 = Q("What is the value?")
    assert verify.verify(q2, ReaderOut("9", support=[1, 3]), pkg2, idx).support_rid == 1


def test_rank_unknown_value_is_default(monkeypatch):
    monkeypatch.setattr(config, "RANK", "nonsense")
    idx = _rank_index()
    q = Q("What is the firmware version of widget alpha?")
    pkg = Package(rids=[1, 3], text="", via={1: 3})
    assert verify.verify(q, ReaderOut("9", support=[1, 3]), pkg, idx).support_rid == 1


# ---------------------------------------------------------------- H1 type guard (verify, only with XLSX_JOIN)

@pytest.mark.parametrize("text,span,bad", [
    ("What is the unit price of the ZT-220 at 5,000 unit volume?", "TAM-BOM-6460-A", True),
    ("How many fans does the TQ-40 have?", "E7731", True),
    ("What is the lead time in weeks?", "ORR-FAN-2214-B", True),
    ("What is the unit price?", "$84.50", False),
    ("What is the lead time?", "14 weeks", False),
    ("What is the unit price?", "12V", False),
    ("How many fans does the TQ-40 have?", "4", False),
    ("What is the board power?", "50%", False),
    ("What is the cost?", "4.3.2", False),
    ("What is the cost?", "v6.9.1", False),
    ("What is the part number of the cost-reduced fan?", "ORR-FAN-2214-B", False),
    ("Which firmware version fixed the lead time bug?", "FW-1204", False),
    ("What is the auxiliary power connector?", "12V-2x6", False),
    ("What is the unit price?", "6-pin PCIe", False),
])
def test_type_guard_span_shapes(text, span, bad):
    import retrieve
    assert verify._type_mismatch(retrieve.analyze(text), span) is bad


def test_type_guard_only_with_join_flag(monkeypatch):
    recs = [("Parts!R1", "[sheet Parts] columns: part_number | product"),
            sheet_row("Parts", 2, part_number="TAM-BOM-6460-A", product="ZT-220"),
            sheet_row("Costs", 2, part_number="TAM-BOM-6460-A", qty_break=5000, unit_cost_usd=358)]
    recs[1] = sheet_row("Parts", 2, part_number="TAM-BOM-6460-A", description="Power supply", product="ZT-220")
    q = Q("What is the unit price of the power supply for the ZT-220 at 5,000 unit volume?", entities=["ZT-220"],
          numbers=["5,000"], attr_terms=["price"])

    def run():
        idx = ix.build([pf("b.xlsx", "xlsx", recs)])
        rid = next(r.rid for r in idx.records if r.src == "join")
        pkg = Package(rids=[rid], text="")
        return verify.verify(q, ReaderOut("TAM-BOM-6460-A", support=[rid]), pkg, idx), verify.verify(
            q, ReaderOut("358", support=[rid]), pkg, idx)

    monkeypatch.setattr(config, "XLSX_JOIN", True)
    bad, good = run()
    assert not bad.ok and bad.reason == "TYPE_MISMATCH"
    assert good.ok and good.answer_span == "358"
    # guard off with the flag: v5 verifier accepts the part number (nothing else rejects it)
    idx = ix.build([pf("b.xlsx", "xlsx", recs)])
    monkeypatch.setattr(config, "XLSX_JOIN", False)
    rid = next(r.rid for r in idx.records if r.loc == "Costs!R2")
    v = verify.verify(q, ReaderOut("TAM-BOM-6460-A", support=[rid]), Package(rids=[rid], text=""), idx)
    assert v.reason != "TYPE_MISMATCH"


# ---------------------------------------------------------------- H1 join eligibility

def _join_idx(monkeypatch, flag=True):
    monkeypatch.setattr(config, "XLSX_JOIN", flag)
    return ix.build([workbook(extra=[
        sheet_row("LeadTimes", 4, part_number="ORR-PSU-3301-A", lead_time_weeks=14, qty=500)])])


def _join_of(idx, loc_prefix):
    return next(r for r in idx.records if r.src == "join" and r.loc.startswith(loc_prefix))


def test_join_needs_question_word_in_s_part(monkeypatch):
    idx = _join_idx(monkeypatch)
    j = _join_of(idx, "LeadTimes!R2~")
    ok = Q("What is the lead time, in weeks, of the power supply unit for the TQ-40?", entities=["TQ-40"],
           attr_terms=["lead_time"])
    assert verify._check(ok, idx, j.rid, "14") is None
    # nothing but attribute words and the product: the S part cannot vouch for it
    bare = Q("What is the lead time of the TQ-40?", entities=["TQ-40"], attr_terms=["lead_time"])
    assert verify._check(bare, idx, j.rid, "14") == "ENTITY_MISMATCH"
    price = Q("What is the unit price of the TQ-40 at 5,000 unit volume?", entities=["TQ-40"],
              numbers=["5,000"], attr_terms=["price"])
    assert verify._check(price, idx, j.rid, "14") == "ENTITY_MISMATCH"
    # the word guard never touches a plain record, and is off without the flag
    plain = next(r for r in idx.records if r.loc == "Parts!R2")
    assert verify._check(bare, idx, plain.rid, "Power supply unit") is None
    monkeypatch.setattr(config, "XLSX_JOIN", False)
    assert verify._check(bare, idx, j.rid, "14") is None


def test_join_numbers_must_come_from_r_part(monkeypatch):
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    recs = [sheet_row("Parts", 2, part_number="ORR-PSU-3301-A", description="Power supply 1600 W", product="TQ-40"),
            sheet_row("LeadTimes", 2, part_number="ORR-PSU-3301-A", lead_time_weeks=14)]
    idx = ix.build([pf("b.xlsx", "xlsx", recs)])
    j = _join_of(idx, "LeadTimes!R2~")
    mk = lambda n: Q("What is the lead time of the power supply for the TQ-40 at %s?" % n, entities=["TQ-40"],
                     numbers=[n], attr_terms=["lead_time"])
    assert verify._check(mk("1600"), idx, j.rid, "14") == "COND_MISMATCH"      # 1600 only in the S part
    assert verify._check(mk("14"), idx, j.rid, "14") is None                   # 14 is in the R part


# ---------------------------------------------------------------- H8 date condition

@pytest.mark.parametrize("text,dates", [
    ("on 2026-09-23", {(2026, 9, 23)}), ("September 23, 2026", {(2026, 9, 23)}), ("23 Sep 2026", {(2026, 9, 23)}),
    ("logs/prod_20260923.log", {(2026, 9, 23)}), ("Sept. 3rd, 2026", {(2026, 9, 3)}),
    ("2026-13-45", set()), ("version 4.3.2", set()), ("on 2026-09-23T10:02:11Z", {(2026, 9, 23)}),
])
def test_dates_in(text, dates):
    assert verify._dates_in(text) == dates


def test_date_cond_flag(monkeypatch):
    files = [FileEntry(file=f, kind="log", status="ok") for f in ("logs/a_2026-09-02.log", "logs/b.log")]
    recs = [Record(0, files[0].file, "name", "logs a", meta=True, src="name"),
            Record(1, files[0].file, "L1", "ERROR rack door alarm code E7731"),
            Record(2, files[1].file, "name", "logs b", meta=True, src="name"),
            Record(3, files[1].file, "L1", "2026-09-23T10:00:00 ERROR rack door alarm code E7740"),
            Record(4, files[1].file, "L2", "2026-09-05T10:00:00 ERROR rack door alarm code E7750")]
    idx = ix.Index(files, recs)
    q = Q("What error code is logged when the rack door alarm occurs on 2026-09-23?")
    assert verify._check(q, idx, 1, "E7731") is None                       # off: v5 does not look at dates
    assert verify._check(q, idx, 4, "E7750") is None
    monkeypatch.setattr(config, "DATE_COND", True)
    assert verify._check(q, idx, 1, "E7731") == "DATE_MISMATCH"            # path has 09-02, not 09-23
    assert verify._check(q, idx, 4, "E7750") == "DATE_MISMATCH"
    assert verify._check(q, idx, 3, "E7740") is None                       # the record text carries the date
    # a date in the file name the question names is satisfied by the path
    qn = Q("What error code is in logs/a_2026-09-02.log?")
    assert verify._check(qn, idx, 1, "E7731") is None
    q2 = Q("What error code is logged on September 2, 2026?")
    assert verify._check(q2, idx, 1, "E7731") is None                      # other written form, same date
    out = verify.verify(q, ReaderOut("E7731", support=[1]), Package(rids=[1, 3], text=""), idx)
    assert not out.ok and out.reason in ("DATE_MISMATCH", "NO_SPAN")
    v = verify.verify(q, ReaderOut("E7740", support=[3]), Package(rids=[1, 3], text=""), idx)
    assert v.ok and v.support_rid == 3


def test_date_cond_default_off_and_knob():
    assert config.DATE_COND is False
    from eval import ab_suite
    assert ab_suite.resolve_env({"MC3_DATE_COND": "1"}) == {"DATE_COND": True}


# ---------------------------------------------------------------- review fixes

def test_formula_chain_across_sheets_multi_pass(monkeypatch, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["item", "total", "cycle", "dep"])
    ws.append(["all", "=Late!B2+1", "=Summary!C2+1", "=Missing!A1"])
    late = wb.create_sheet("Late")
    late.append(["item", "subtotal"])
    late.append(["x", "=SUM(Late!C2:C3)*2"])
    late["C2"] = 3
    late["C3"] = 4
    p = tmp_path / "chain.xlsx"
    wb.save(p)
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    t = {r["loc"]: r["text"] for r in parse._xlsx(p)}
    assert "total=15" in t["Summary!R2"]                 # Late!B2 = 14 is resolved first, then +1
    assert "cycle" not in t["Summary!R2"] and "dep" not in t["Summary!R2"]   # unresolved: no value, never 0
    # flag off keeps v5 (reads the unevaluated formula cell as 0)
    monkeypatch.setattr(config, "XLSX_JOIN", False)
    t = {r["loc"]: r["text"] for r in parse._xlsx(p)}
    assert "total=1" in t["Summary!R2"]


def test_eval_formula_unresolved_reference():
    grids = {"a": {}, "b": {(1, 1): 5}}
    assert parse._eval_formula("=B!A1+1", grids["a"], grids, "a", {("b", 1, 1)}) is None
    assert parse._eval_formula("=SUM(B!A1:A2)", grids["a"], grids, "a", {("b", 2, 1)}) is None
    assert parse._eval_formula("=B!A1+1", grids["a"], grids, "a", set()) == 6


def test_join_big_bucket_and_budget(monkeypatch):
    import time
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    recs = [sheet_row("Parts", 2, part_number="ORR-PSU-3301-A", product="TQ-40")]
    recs += [sheet_row("Log", i + 2, part_number="ORR-PSU-3301-A", seq=i) for i in range(5000)]
    t0 = time.monotonic()
    idx = ix.build([pf("big.xlsx", "xlsx", recs)])
    assert time.monotonic() - t0 < 3
    assert not derived(idx, "join")           # a key in more than JOIN_MAX_BUCKET rows is a category
    recs = [sheet_row("Parts", 2, part_number="ORR-PSU-3301-A", product="TQ-40")]
    recs += [sheet_row("Log", i + 2, part_number="ORR-PSU-3301-A", seq=i) for i in range(150)]
    assert len(derived(ix.build([pf("ok.xlsx", "xlsx", recs)]), "join")) == 150


def test_derived_fail_open(monkeypatch):
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    monkeypatch.setattr(config, "PDF_COLS", True)
    monkeypatch.setattr(ix, "_join_records", lambda recs, deadline=None: 1 / 0)
    monkeypatch.setattr(ix, "_col_records", lambda recs, deadline=None: 1 / 0)
    idx = ix.build([workbook(), pf("s.pdf", "pdf", PDF_TABLE)])
    assert {r.src for r in idx.records} == {"name", "text"}
    # a time overrun gives none as well
    monkeypatch.undo()
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    monkeypatch.setattr(ix, "DERIVED_BUDGET_S", -1.0)
    assert not derived(ix.build([workbook()]), "join")


def test_cols_limits(monkeypatch):
    monkeypatch.setattr(config, "PDF_COLS", True)
    wide = "Parameter  " + "  ".join(f"TQ-{i}" for i in range(20, 20 + ix.COL_MAX_PRODUCTS + 1))
    assert ix._col_header(wide) is None
    ok = "Parameter  " + "  ".join(f"TQ-{i}" for i in range(20, 20 + ix.COL_MAX_PRODUCTS))
    assert len(ix._col_header(ok)) == ix.COL_MAX_PRODUCTS
    # many tables: the per-file record cap
    t = []
    n = 0
    for _ in range(600):
        t.append((f"p1.L{n + 1}", "Parameter    TQ-20    TQ-40    TQ-60"))
        for j in range(COL_ROWS := 4):
            t.append((f"p1.L{n + 2 + j}", f"Item {j}    1 W    2 W    3 W"))
        n += 5
    # all on one page with consecutive line numbers
    cols = derived(ix.build([pf("s.pdf", "pdf", t)]), "col")
    assert len(cols) == ix.COL_MAX_RECORDS


@pytest.mark.parametrize("span", ["14d", "6wk", "3yr", "100GbE", "400G", "x16", "X8", "16x", "PCIe5", "DDR5-4800",
                                  "3V3", "1.5d"])
def test_type_guard_spares_quantity_shapes(span):
    import retrieve
    q = retrieve.analyze("What is the lead time of the power supply?")
    assert verify._type_mismatch(q, span) is False


def test_type_guard_join_records_only(monkeypatch):
    recs = [sheet_row("Parts", 2, part_number="TAM-BOM-6460-A", description="Power supply", product="ZT-220"),
            sheet_row("Costs", 2, part_number="TAM-BOM-6460-A", qty_break=5000, unit_cost_usd=358)]
    monkeypatch.setattr(config, "XLSX_JOIN", True)
    idx = ix.build([pf("b.xlsx", "xlsx", recs)])
    q = Q("What is the unit price of the power supply for the ZT-220 at 5,000 unit volume?", entities=["ZT-220"],
          numbers=["5,000"], attr_terms=["price"])
    plain = next(r for r in idx.records if r.loc == "Parts!R2")
    join = next(r for r in idx.records if r.src == "join")
    assert verify._check(q, idx, plain.rid, "TAM-BOM-6460-A") != "TYPE_MISMATCH"
    assert verify._check(q, idx, join.rid, "TAM-BOM-6460-A") == "TYPE_MISMATCH"
