"""Defensive corpus walker and per-format extractors.

Every file becomes one dict:
    {"path": rel, "kind": ..., "status": "ok|unreadable|encrypted|binary|empty|error",
     "records": [{"loc": str, "text": str, "stale": bool}],   # atomic evidence units
     "text": str,                                               # records joined, for search
     "images": [("file", rel) | ("zip", member) | ("pdfpage", n)],  # need the vision model
     "flags": {"withdrawn": bool, "supersedes": [...], "revision": str|None}}

A record is the smallest unit that can carry an answer on its own: a log line, a
CSV/XLSX row rendered as header=value pairs, a docx paragraph or table row, a PDF
line, a Python constant. Nothing here raises: a failure on one file is recorded
on that file and the walk continues. Walk order is sorted so local and graded
runs see the same order.
"""

from __future__ import annotations

import ast
import csv
import datetime as dt
import io
import os
import re
import signal
import stat
import threading
import time
import xml.etree.ElementTree as ET
import zipfile
from contextlib import contextmanager
from pathlib import Path

TEXT_EXT = {".txt", ".log", ".md", ".py", ".json", ".yaml", ".yml", ".ini", ".cfg", ".conf",
            ".toml", ".sh", ".c", ".h", ".cpp", ".js", ".ts", ".xml", ".html", ".rst", ".out"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}

# (?<![a-z]) instead of \b: "_" is a word char, so \b would miss "r1_WITHDRAWN"
WITHDRAWN_RE = re.compile(
    r"(?<![a-z])(withdrawn|superseded|obsolete|deprecated|do not use|retracted|void)(?![a-z])", re.I)
SUPERSEDES_RE = re.compile(r"supersedes\s+(?:revision\s+|rev\.?\s*|r(?=\d))?([\w.\-]+)", re.I)
STALE_MARK_RE = re.compile(
    r"\((?:legacy|obsolete|superseded|withdrawn|deprecated|retired|old)\)"
    r"|\b(?:status|state)\s*[:=]\s*(?:legacy|obsolete|superseded|withdrawn|deprecated|retired)\b",
    re.I)
PDF_MAX_PAGES = 500       # text pages read per PDF
PDF_VISION_PAGES = 60     # only the first pages are candidates for the VLM
MAX_RECORDS = 20_000      # records kept per file (the first ones)
XLSX_MAX_ROWS, XLSX_MAX_COLS = 20_000, 256   # cells read per sheet
XLSX_MAX_MERGE_CELLS = 10_000                # larger merged ranges are not propagated
FILE_TIMEOUT_S = 20       # wall time per file (main thread only)
DRAFT_RE = re.compile(r"(?<![a-z])draft(?![a-z])", re.I)
REV_RE = re.compile(r"(?:^|[_\-\s])(?:r|rev|v)(\d+[a-z]?)(?=[_\-\s.]|$)", re.I)
STALE_VALUES = {"withdrawn", "superseded", "obsolete", "deprecated", "retired", "cancelled",
                "canceled", "void", "inactive", "replaced", "eol", "discontinued"}
STATUS_COLS = {"status", "state", "lifecycle", "revision_status", "doc_status", "validity"}

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class _FileTimeout(BaseException):
    """Raised by the SIGALRM guard; BaseException so a parser's ``except Exception`` cannot swallow it."""


@contextmanager
def _time_guard(seconds: float):
    """Raise _FileTimeout in the main thread after ``seconds`` of wall time; restores the previous
    handler and timer. A no-op off the main thread or where SIGALRM/setitimer do not exist."""
    if not hasattr(signal, "setitimer") or threading.current_thread() is not threading.main_thread():
        yield
        return

    def on_alarm(_signum, _frame):
        raise _FileTimeout()

    t0 = time.monotonic()
    old_handler = signal.signal(signal.SIGALRM, on_alarm)
    old_delay, old_interval = signal.setitimer(signal.ITIMER_REAL, seconds)
    try:
        yield
    finally:
        while True:     # the one-shot alarm may fire inside this block: retry once it is spent
            try:
                signal.setitimer(signal.ITIMER_REAL, 0)
                signal.signal(signal.SIGALRM, old_handler)
                if old_delay > 0:
                    left = max(0.001, old_delay - (time.monotonic() - t0))
                    signal.setitimer(signal.ITIMER_REAL, left, old_interval)
                break
            except _FileTimeout:
                continue


def _rec(loc: str, text: str, stale: bool = False) -> dict:
    return {"loc": loc, "text": text.strip(), "stale": stale}


# ---------------------------------------------------------------- text decoding

def _decode(b: bytes) -> str | None:
    """Best-effort text decode; None if the bytes do not look like text."""
    if not b:
        return ""
    for bom, enc in ((b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe\x00\x00", "utf-32"),
                     (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16")):
        if b.startswith(bom):
            return b.decode(enc, errors="replace")
    head = b[:4096]
    if b"\x00" in head:
        even, odd = head[0::2].count(0), head[1::2].count(0)
        n = len(head) // 2 or 1
        if odd / n > 0.3 and even / n < 0.05:
            return b.decode("utf-16-le", errors="replace")
        if even / n > 0.3 and odd / n < 0.05:
            return b.decode("utf-16-be", errors="replace")
        return None
    try:
        return b.decode("utf-8")
    except UnicodeDecodeError:
        printable = sum(32 <= c < 127 or c in (9, 10, 13) or c >= 160 for c in head)
        if printable / len(head) < 0.85:
            return None
        return b.decode("cp1252", errors="replace")


def _sniff(head: bytes, path: Path) -> str | None:
    """Content type from magic bytes, so a misnamed file still gets parsed."""
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\x89PNG") or head.startswith(b"\xff\xd8\xff") or head[:4] in (
            b"GIF8", b"II*\x00", b"MM\x00*") or head[:2] == b"BM" or (
            head[:4] == b"RIFF" and head[8:12] == b"WEBP"):
        return "image"
    if head.startswith(b"PK\x03\x04"):
        try:
            names = set(zipfile.ZipFile(path).namelist())
        except Exception:
            return None
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if any(n.startswith("ppt/slides/") for n in names):
            return "pptx"
    return None


# ---------------------------------------------------------------- OOXML helpers

def _chart_records(z: zipfile.ZipFile, prefix: str) -> list[dict]:
    """Cached series values from embedded charts (word/charts, xl/charts)."""
    out = []
    for n in sorted(x for x in z.namelist() if x.startswith(prefix) and x.endswith(".xml")
                    and "/charts/chart" in x):
        try:
            root = ET.fromstring(z.read(n))
        except ET.ParseError:
            continue
        title = " ".join(t.text or "" for t in root.iter(f"{A}t")).strip()

        def vals_of(node):
            # explicit None check: Element truthiness is deprecated (len-based)
            return [] if node is None else [v.text or "" for v in node.iter(f"{C}v")]

        for ser in root.iter(f"{C}ser"):
            name = " ".join(vals_of(ser.find(f"{C}tx")))
            cats = vals_of(ser.find(f"{C}cat"))
            vals = vals_of(ser.find(f"{C}val"))
            pairs = "; ".join(f"{c}={v}" for c, v in zip(cats, vals)) or "; ".join(vals)
            out.append(_rec(f"{n.split('/')[-1]}", f"chart {title} series {name}: {pairs}"))
    return out


def _para_text(p: ET.Element) -> str:
    parts = []
    for el in p.iter():
        if el.tag == f"{W}t":
            parts.append(el.text or "")
        elif el.tag == f"{W}tab":
            parts.append("\t")
        elif el.tag in (f"{W}br", f"{W}cr"):
            parts.append(" ")
    return "".join(parts)


def _strip_deleted(root: ET.Element) -> None:
    """Tracked changes: the current text wins, so drop <w:del> subtrees."""
    for parent in root.iter():
        for child in list(parent):
            if child.tag in (f"{W}del", f"{W}moveFrom"):
                parent.remove(child)


def _docx_body(el: ET.Element, recs: list, label: str, counter: list) -> None:
    for child in el:
        tag = child.tag
        if tag == f"{W}p":
            t = _para_text(child).strip()
            if t:
                counter[0] += 1
                recs.append(_rec(f"{label}p{counter[0]}", t))
        elif tag == f"{W}tbl":
            counter[1] += 1
            rows = []
            for tr in child.iter(f"{W}tr"):
                cells = [" ".join(_para_text(p).strip() for p in tc.iter(f"{W}p")).strip()
                         for tc in tr.findall(f"{W}tc")]
                if any(cells):
                    rows.append(cells)
            if not rows:
                continue
            header = rows[0]
            recs.append(_rec(f"{label}table{counter[1]}.header", " | ".join(header)))
            for i, r in enumerate(rows[1:], 1):
                if len(header) == len(r) and len(header) > 1:
                    txt = "; ".join(f"{h}={v}" for h, v in zip(header, r) if v)
                else:
                    txt = " | ".join(r)
                recs.append(_rec(f"{label}table{counter[1]}.row{i}", txt))
        elif tag in (f"{W}sdt", f"{W}sdtContent", f"{W}customXml", f"{W}body"):
            _docx_body(child, recs, label, counter)


def _docx(p: Path) -> tuple[list, list]:
    z = zipfile.ZipFile(p)
    names = z.namelist()
    recs: list[dict] = []
    order = ["word/document.xml"] + sorted(n for n in names if re.match(
        r"word/(header|footer|footnotes|endnotes|comments)\d*\.xml$", n))
    for n in order:
        if n not in names:
            continue
        root = ET.fromstring(z.read(n))
        _strip_deleted(root)
        label = "" if n == "word/document.xml" else n.split("/")[-1][:-4] + "."
        body = root.find(f"{W}body")
        _docx_body(body if body is not None else root, recs, label, [0, 0])
    recs += _chart_records(z, "word/")
    media = [("zip", n) for n in names
             if n.startswith("word/media/") and Path(n).suffix.lower() in IMAGE_EXT]
    return recs, media


def _pptx(p: Path) -> tuple[list, list]:
    z = zipfile.ZipFile(p)
    recs = []
    slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                    key=lambda s: int(re.findall(r"\d+", s)[-1]))
    for n in slides:
        root = ET.fromstring(z.read(n))
        for i, para in enumerate(root.iter(f"{A}p")):
            t = "".join(x.text or "" for x in para.iter(f"{A}t")).strip()
            if t:
                recs.append(_rec(f"{n.split('/')[-1][:-4]}.p{i}", t))
    media = [("zip", n) for n in z.namelist()
             if n.startswith("ppt/media/") and Path(n).suffix.lower() in IMAGE_EXT]
    return recs + _chart_records(z, "ppt/"), media


# ---------------------------------------------------------------- spreadsheets

_REF = re.compile(r"(?:(?:'[^']+'|[A-Za-z_][\w. ]*)!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
_REF_SHEET = re.compile(r"(?:(?P<sheet>'(?:[^']|'')+'|[A-Za-z_][\w.]*)!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
_FUNCS = {"SUM": lambda *a: sum(_flat(a)), "MIN": lambda *a: min(_flat(a)),
          "MAX": lambda *a: max(_flat(a)), "AVERAGE": lambda *a: sum(_flat(a)) / len(_flat(a)),
          "ROUND": lambda x, n=0: round(x, int(n)), "ABS": abs}


def _flat(a):
    out = []
    for x in a:
        out.extend(_flat(x) if isinstance(x, (list, tuple)) else [x])
    return [x for x in out if isinstance(x, (int, float))]


def _eval_formula(formula: str, ws_values, sheets: dict | None = None, cur: str | None = None,
                  unresolved: set | None = None) -> float | None:
    """Tiny, whitelisted evaluator for arithmetic/SUM-style formulas whose cached
    value is missing (files written by openpyxl carry no cached results).

    ``sheets`` ({lower-cased title: grid}, set only with config.XLSX_JOIN) resolves ``Sheet!B3`` references
    against the referenced sheet; without it every reference is read from ``ws_values`` (the current sheet).
    ``unresolved`` ({(sheet, row, col)}, needs ``cur`` = the current sheet's lower-cased title) holds formula
    cells without a value yet: reading one makes the whole formula unresolved (None), never 0."""
    from openpyxl.utils import column_index_from_string as col

    def cell(c, r, grid=ws_values, name=cur):
        if unresolved and (name, int(r), col(c)) in unresolved:
            raise ValueError("references an unresolved formula")
        v = grid.get((int(r), col(c)))
        return v if isinstance(v, (int, float)) else 0

    def repl(m):
        c1, r1, c2, r2 = m.groups()[-4:]
        grid, name = ws_values, cur
        if sheets is not None and m.group("sheet") is not None:
            name = m.group("sheet")
            name = name[1:-1].replace("''", "'") if name.startswith("'") else name
            name = name.lower()
            grid = sheets.get(name)
            if grid is None:
                raise ValueError("unknown sheet")
        if c2:
            if (int(r2) - int(r1) + 1) * (col(c2) - col(c1) + 1) > XLSX_MAX_MERGE_CELLS:
                raise ValueError("range too large")
            return "[" + ",".join(repr(cell(get_col(ci), ri, grid, name))
                                  for ri in range(int(r1), int(r2) + 1)
                                  for ci in range(col(c1), col(c2) + 1)) + "]"
        return repr(cell(c1, r1, grid, name))

    from openpyxl.utils import get_column_letter as get_col
    try:
        expr = (_REF if sheets is None else _REF_SHEET).sub(repl, formula.lstrip("=")).replace("^", "**")
        tree = ast.parse(expr, mode="eval")
    except (SyntaxError, ValueError, MemoryError, RecursionError):
        return None
    ok = (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant, ast.Call, ast.Name, ast.Load,
          ast.List, ast.Add, ast.Sub, ast.Mult, ast.Div, ast.USub, ast.UAdd)   # no Pow: 9**9**9 hangs
    for node in ast.walk(tree):
        if not isinstance(node, ok) or (isinstance(node, ast.Name) and node.id not in _FUNCS):
            return None
    try:
        return eval(compile(tree, "<f>", "eval"), {"__builtins__": {}}, dict(_FUNCS))
    except Exception:
        return None


def _fmt_number(v, nf: str | None) -> str:
    if isinstance(v, bool):
        return str(v).upper()
    if isinstance(v, (dt.datetime, dt.date)):
        return v.date().isoformat() if isinstance(v, dt.datetime) and not (
            v.hour or v.minute or v.second) else v.isoformat()
    if not isinstance(v, (int, float)):
        return str(v)
    nf = (nf or "General").split(";")[0]
    if nf == "General" or not re.search(r"[0#]", nf):
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return format(v, ".15g")
    pct = "%" in nf
    x = v * 100 if pct else v
    m = re.search(r"[0#,]*\.?([0#]*)", nf[nf.find("0") if "0" in nf else 0:])
    decimals = len(m.group(1)) if m and "." in nf else 0
    thousands = "," in nf.split(".")[0]
    s = f"{x:,.{decimals}f}" if thousands else f"{x:.{decimals}f}"
    return s + ("%" if pct else "")


def _rows_of(ws):
    """Rows of a sheet clamped to XLSX_MAX_ROWS x XLSX_MAX_COLS and to the sheet's own extent
    (iter_rows(max_row=N) materialises empty cells up to N)."""
    return ws.iter_rows(max_row=min(ws.max_row or 1, XLSX_MAX_ROWS), max_col=min(ws.max_column or 1, XLSX_MAX_COLS))


def _fill_formulas(wb_f, grids: dict[str, dict]) -> None:
    """config.XLSX_JOIN: results for formula cells without a cached value, across sheets, in up to 3 passes.
    A formula that reads a formula cell still unresolved stays without a value (never 0)."""
    pending = []   # (sheet key, row, col, formula)
    for ws in wb_f.worksheets:
        key = ws.title.lower()
        grid = grids.get(key)
        if grid is None:
            continue
        for row in _rows_of(ws):
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and \
                        grid.get((c.row, c.column)) in (None, c.value):
                    pending.append((key, c.row, c.column, c.value))
    unresolved = {(k, r, c) for k, r, c, _ in pending}
    for _ in range(3):
        if not pending:
            break
        still = []
        for key, r, c, f in pending:
            val = _eval_formula(f, grids[key], grids, key, unresolved)
            if val is not None:
                grids[key][(r, c)] = val
                unresolved.discard((key, r, c))
            else:
                still.append((key, r, c, f))
        if len(still) == len(pending):
            break
        pending = still


def _xlsx(p: Path) -> list[dict]:
    import openpyxl

    import config
    wb_v = openpyxl.load_workbook(p, data_only=True)
    try:
        wb_f = openpyxl.load_workbook(p, data_only=False)
    except Exception:
        wb_f = None
    recs = []
    # config.XLSX_JOIN: grids of every sheet by lower-cased title, so a formula can read another sheet
    all_grids: dict[str, dict] | None = None
    if config.XLSX_JOIN:
        all_grids = {ws.title.lower(): {(c.row, c.column): c.value for row in _rows_of(ws) for c in row
                                        if c.value is not None} for ws in wb_v.worksheets}
        if wb_f is not None:
            _fill_formulas(wb_f, all_grids)
    for ws in wb_v.worksheets:
        state = getattr(ws, "sheet_state", "visible")
        hidden = state != "visible"
        wsf = wb_f[ws.title] if wb_f else None
        if all_grids is not None:
            grid = all_grids[ws.title.lower()]   # built and formula-filled before the loop
        else:
            grid = {}
            for row in _rows_of(ws):
                for c in row:
                    if c.value is not None:
                        grid[(c.row, c.column)] = c.value
        # fill missing cached formula results
        if wsf is not None and all_grids is None:
            for row in _rows_of(wsf):
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and \
                            grid.get((c.row, c.column)) in (None, c.value):
                        val = _eval_formula(c.value, grid, all_grids)
                        if val is not None:
                            grid[(c.row, c.column)] = val
        # merged ranges: propagate the top-left value
        for mr in ws.merged_cells.ranges:
            if (mr.max_row - mr.min_row + 1) * (mr.max_col - mr.min_col + 1) > XLSX_MAX_MERGE_CELLS:
                continue
            v = grid.get((mr.min_row, mr.min_col))
            for r in range(mr.min_row, mr.max_row + 1):
                for cc in range(mr.min_col, mr.max_col + 1):
                    grid.setdefault((r, cc), v)
        fmt = {(c.row, c.column): c.number_format
               for row in _rows_of(ws) for c in row}
        if not grid:
            continue
        rows_idx = sorted({r for r, _ in grid})
        cols_idx = sorted({c for _, c in grid})

        def cells(r):
            return [_fmt_number(grid[(r, c)], fmt.get((r, c))) if (r, c) in grid else ""
                    for c in cols_idx]

        # header = first row with >=2 non-empty cells, mostly text, not a merged title
        header_r = None
        for r in rows_idx[:12]:
            vals = [grid.get((r, c)) for c in cols_idx]
            filled = [v for v in vals if v not in (None, "")]
            distinct = {str(v) for v in filled}
            if len(distinct) >= 2 and sum(isinstance(v, str) for v in filled) >= 0.7 * len(filled):
                header_r = r
                break
        tag = f"[sheet {ws.title}{' (hidden)' if hidden else ''}]"
        header = cells(header_r) if header_r else None
        for r in rows_idx:
            vals = cells(r)
            if not any(vals):
                continue
            if header_r and r < header_r:
                recs.append(_rec(f"{ws.title}!R{r}", f"{tag} {' '.join(dict.fromkeys(v for v in vals if v))}", hidden))
            elif r == header_r:
                recs.append(_rec(f"{ws.title}!R{r}", f"{tag} columns: {' | '.join(vals)}", hidden))
            else:
                if header:
                    pairs = [(h or f"col{i + 1}", v) for i, (h, v) in enumerate(zip(header, vals)) if v]
                    txt = "; ".join(f"{h}={v}" for h, v in pairs)
                    stale = hidden or any(h.strip().lower() in STATUS_COLS and
                                          v.strip().lower() in STALE_VALUES for h, v in pairs)
                else:
                    txt, stale = " | ".join(vals), hidden
                recs.append(_rec(f"{ws.title}!R{r}", f"{tag} {txt}", stale))
        for row in _rows_of(ws):
            for c in row:
                if c.comment:
                    recs.append(_rec(f"{ws.title}!{c.coordinate}.comment",
                                     f"{tag} comment on {c.coordinate}: {c.comment.text}", hidden))
    try:
        recs += _chart_records(zipfile.ZipFile(p), "xl/")
    except Exception:
        pass
    return recs


def _csv(text: str) -> list[dict]:
    csv.field_size_limit(2**31 - 1)   # default 128 KiB raises "field larger than field limit" on one long cell
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(x.strip() for x in r)]
    if not rows:
        return []
    header = [h.strip() for h in rows[0]]
    recs = [_rec("header", "columns: " + " | ".join(header))]
    for i, r in enumerate(rows[1:], 2):
        pairs = [(h, v.strip()) for h, v in zip(header, r) if v.strip()]
        stale = any(h.lower() in STATUS_COLS and v.lower() in STALE_VALUES for h, v in pairs)
        recs.append(_rec(f"row{i}", "; ".join(f"{h}={v}" for h, v in pairs), stale))
    return recs


# ---------------------------------------------------------------- pdf / code / text

def _pdf_pages(p: Path) -> tuple[list[str], list[bool]]:
    """Per-page text and 'has image objects' flags; pypdfium2 if importable, else pypdf."""
    try:
        import pypdfium2
    except ImportError:
        import pypdf
        r = pypdf.PdfReader(str(p))
        if r.is_encrypted:
            r.decrypt("")
        pages = [(pg.extract_text() or "") for pg in r.pages[:PDF_MAX_PAGES]]
        return pages, [False] * len(pages)
    pdf = pypdfium2.PdfDocument(str(p))
    pages, has_img = [], []
    try:
        for i in range(min(len(pdf), PDF_MAX_PAGES)):
            page = pdf[i]
            pages.append(page.get_textpage().get_text_bounded() or "")
            try:
                objs = page.get_objects(filter=[pypdfium2.raw.FPDF_PAGEOBJ_IMAGE])
                has_img.append(next(iter(objs), None) is not None)
            except Exception:  # noqa: BLE001 - image count is only a hint
                has_img.append(False)
    finally:
        pdf.close()
    return pages, has_img


def _pdf_encrypted(p: Path) -> bool:
    """True if the PDF needs a password we do not have (the empty user password is tried)."""
    import pypdf
    try:
        r = pypdf.PdfReader(str(p))
        return bool(r.is_encrypted and not r.decrypt(""))
    except pypdf.errors.FileNotDecryptedError:   # a PdfReadError subclass: keep before the next clause
        return True
    except (pypdf.errors.PdfStreamError, pypdf.errors.PdfReadError):
        return False    # unparsable for pypdf: not provably encrypted, let the pdfium text path try
    except pypdf.errors.DependencyError:  # AES needs `cryptography`; let pdfium decide
        try:
            import pypdfium2
            pypdfium2.PdfDocument(str(p)).close()
            return False
        except Exception:  # noqa: BLE001 - password error, or pdfium missing: refuse
            return True


def _pdf(p: Path) -> tuple[str, list, list]:
    """Returns (status, records, pages_needing_vision); vision pages are 1-based."""
    if _pdf_encrypted(p):
        return "encrypted", [], []
    pages, has_img = _pdf_pages(p)
    vision = []
    for i, (t, img) in enumerate(zip(pages, has_img), 1):
        n = len(t.strip())
        if i <= PDF_VISION_PAGES and (n < 40 or (img and n < 200)):
            vision.append(("pdfpage", i))
    recs = []
    for i, t in enumerate(pages, 1):
        for j, line in enumerate(t.splitlines(), 1):
            if line.strip():
                recs.append(_rec(f"p{i}.L{j}", re.sub(r"\.{4,}", " ... ", line)))
    return "ok", recs, vision


def _py_records(text: str) -> list[dict]:
    """Module/class constants and argument defaults, values resolved through names.
    The literal in code is authoritative; comments are kept as separate records."""
    try:
        tree = ast.parse(text)
    except SyntaxError:
        return []
    consts: dict[str, str] = {}
    out = []

    def lit(node):
        try:
            return repr(ast.literal_eval(node))
        except Exception:
            if isinstance(node, ast.Name) and node.id in consts:
                return consts[node.id]
            try:
                return ast.unparse(node)
            except Exception:
                return None

    def visit(body, scope):
        for node in body:
            if isinstance(node, (ast.Assign, ast.AnnAssign)):
                targets = node.targets if isinstance(node, ast.Assign) else [node.target]
                if node.value is None:
                    continue
                v = lit(node.value)
                for t in targets:
                    if isinstance(t, ast.Name) and v is not None:
                        consts.setdefault(t.id, v)
                        out.append(_rec(f"L{node.lineno}.const", f"code constant {scope}{t.id} = {v}"))
            elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                a = node.args
                pos = a.posonlyargs + a.args
                for arg, d in zip(pos[len(pos) - len(a.defaults):], a.defaults):
                    out.append(_rec(f"L{node.lineno}.default",
                                    f"code default {scope}{node.name}({arg.arg}) = {lit(d)}"))
                for arg, d in zip(a.kwonlyargs, a.kw_defaults):
                    if d is not None:
                        out.append(_rec(f"L{node.lineno}.default",
                                        f"code default {scope}{node.name}({arg.arg}) = {lit(d)}"))
            elif isinstance(node, ast.ClassDef):
                visit(node.body, f"{scope}{node.name}.")

    visit(tree.body, "")
    return out


def _text_records(text: str, py: bool = False) -> list[dict]:
    """One record per non-blank line; in .py files comment-only lines get loc ``L{n}.comment``."""
    return [_rec(f"L{i}.comment" if py and line.strip().startswith("#") else f"L{i}", line)
            for i, line in enumerate(text.splitlines(), 1) if line.strip()]


# ---------------------------------------------------------------- driver

def _flags(rel: str, text: str) -> dict:
    head = text[:600]
    name = Path(rel).stem
    # Only the document's own title lines decide its status: a revision history
    # saying "Revision 1 is withdrawn" must not mark revision 2 as withdrawn.
    title = [ln for ln in text.splitlines() if ln.strip()][:3]
    withdrawn = bool(WITHDRAWN_RE.search(name)) or any(
        re.search(r"\b(withdrawn|superseded by|obsolete|do not (design|use))\b", ln, re.I)
        and not re.search(r"\bsupersedes\b", ln, re.I) for ln in title)
    rev = REV_RE.search(name)
    draft = bool(DRAFT_RE.search(Path(rel).name)) or any(DRAFT_RE.search(ln) for ln in title)
    return {"withdrawn": withdrawn, "draft": draft,
            "supersedes": SUPERSEDES_RE.findall(head),
            "revision": rev.group(1) if rev else None}


def _extract_body(rec: dict, path: Path, rel: str, ext: str) -> None:
    st = path.stat()
    if not stat.S_ISREG(st.st_mode):    # FIFO / socket / device: opening it could block forever
        rec["status"] = "error"
        rec["error"] = "not a regular file"
    elif st.st_size == 0:
        rec["status"] = "empty"
    else:
        with open(path, "rb") as fh:  # PermissionError on mode 000 without DAC_OVERRIDE
            head = fh.read(8192)
        kind = _sniff(head, path) or {".docx": "docx", ".xlsx": "xlsx", ".xlsm": "xlsx",
                                      ".pptx": "pptx", ".csv": "csv", ".tsv": "csv"}.get(ext)
        if kind is None and ext in IMAGE_EXT:
            kind = "image"
        if kind == "pdf":
            rec["kind"] = "pdf"
            rec["status"], rec["records"], rec["images"] = _pdf(path)
        elif kind == "docx":
            rec["kind"] = "docx"
            rec["records"], rec["images"] = _docx(path)
        elif kind == "pptx":
            rec["kind"] = "pptx"
            rec["records"], rec["images"] = _pptx(path)
        elif kind == "xlsx":
            rec["kind"] = "xlsx"
            rec["records"] = _xlsx(path)
        elif kind == "image":
            rec["kind"] = "image"
            rec["images"] = [("file", rel)]
        else:
            text = _decode(path.read_bytes())
            if text is None:
                rec["status"] = "binary"
            elif kind == "csv":
                rec["records"] = _csv(text)
            else:
                rec["kind"] = "py" if ext == ".py" else (ext.lstrip(".") or "txt")
                rec["records"] = _text_records(text, ext == ".py")
                if ext == ".py":
                    rec["records"] += _py_records(text)


def extract(root: Path, path: Path) -> dict:
    rel = path.relative_to(root).as_posix()
    ext = path.suffix.lower()
    rec = {"path": rel, "kind": ext.lstrip(".") or "none", "status": "ok",
           "records": [], "images": [], "flags": {}}
    try:
        with _time_guard(FILE_TIMEOUT_S):
            _extract_body(rec, path, rel, ext)
    except _FileTimeout:
        rec.update(status="error", error="timeout", records=[], images=[])
    except PermissionError:
        rec["status"] = "unreadable"
    except Exception as e:  # noqa: BLE001 - one file must never kill the walk
        rec["status"] = "error"
        rec["error"] = f"{type(e).__name__}: {e}"[:300]
    if len(rec["records"]) > MAX_RECORDS:
        rec["note"] = f"records truncated: kept the first {MAX_RECORDS} of {len(rec['records'])}"
        rec["records"] = rec["records"][:MAX_RECORDS]
    for r in rec["records"]:
        if STALE_MARK_RE.search(r["text"]):
            r["stale"] = True
    rec["text"] = "\n".join(r["text"] for r in rec["records"])
    rec["flags"] = _flags(rel, rec["text"])
    return rec


def walk(root: Path) -> list[dict]:
    recs = []
    for dirpath, dirnames, filenames in os.walk(root, onerror=lambda _e: None):
        dirnames.sort()
        for fn in sorted(filenames):
            recs.append(extract(root, Path(dirpath) / fn))
    return recs


if __name__ == "__main__":
    import json
    import sys
    for r in walk(Path(sys.argv[1])):
        print(json.dumps({"path": r["path"], "kind": r["kind"], "status": r["status"],
                          "flags": r["flags"], "images": r["images"],
                          "records": r["records"][:40]}, ensure_ascii=False))
